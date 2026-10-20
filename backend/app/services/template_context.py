"""Bounded, file-derived context for arbitrary PPTX files, without name heuristics."""
import base64
from hashlib import sha256
from io import BytesIO
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .palette import roles


def extract_examples(prs):
    examples, assets = [], []
    seen = {}
    width, height = int(prs.slide_width), int(prs.slide_height)

    def visit(shapes, records, slide_index, depth=0):
        if depth > 5:
            return
        for shape in shapes:
            if len(records) >= 50:
                break
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                # Child coordinates are in group space, so do not use them as slide zones.
                visit(shape.shapes, records, slide_index, depth + 1)
                continue
            record = {'name': shape.name[:100], 'type': 'shape', 'grouped': depth > 0}
            try:
                record['box'] = {k: round(v, 4) for k, v in zip(('x','y','w','h'),
                    (shape.left/width, shape.top/height, shape.width/width, shape.height/height))}
            except (TypeError, ValueError):
                continue
            if shape.has_text_frame:
                record.update(type='text', text=shape.text[:800])
            if shape.has_table:
                record.update(type='table', rows=[[c.text[:100] for c in row.cells][:6]
                                                 for row in list(shape.table.rows)[:6]])
            if shape.has_chart:
                record['type'] = 'chart'
                try:
                    record['categories'] = [str(c.label)[:80] for c in shape.chart.plots[0].categories][:10]
                    record['series'] = [{'name': s.name[:80], 'values': list(s.values)[:10]}
                                        for s in list(shape.chart.series)[:3]]
                except (ValueError, AttributeError, IndexError):
                    record['note'] = 'Тип диаграммы требует проверки в PowerPoint.'
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                record['type'] = 'picture'
                alt = shape._element.xpath('.//p:cNvPr')
                record['description'] = (alt[0].get('descr', '') if alt else '')[:300]
                try:
                    blob = shape.image.blob
                    digest = sha256(blob).hexdigest()[:16]
                    record['asset_id'] = digest
                    if digest not in seen and len(assets) < 8:
                        with Image.open(BytesIO(blob)) as im:
                            if im.width * im.height > 25_000_000:
                                raise ValueError('large image')
                            im.thumbnail((512, 512))
                            out = BytesIO()
                            im.convert('RGBA').save(out, format='PNG', optimize=True)
                        if len(out.getvalue()) <= 200_000:
                            assets.append({'id': digest, 'name': record['name'], 'slide': slide_index,
                                           'description': record['description'],
                                           'data': 'data:image/png;base64,' + base64.b64encode(out.getvalue()).decode()})
                            seen[digest] = True
                except (OSError, ValueError, Image.DecompressionBombError):
                    record['note'] = 'Изображение нельзя использовать в веб-предпросмотре.'
            records.append(record)

    for index, slide in enumerate(list(prs.slides)[:12]):
        records = []
        visit(slide.shapes, records, index)
        layout = slide.slide_layout
        master_index = next((i for i, m in enumerate(prs.slide_masters) if m.part == layout.slide_master.part), 0)
        layout_index = next((i for i, l in enumerate(layout.slide_master.slide_layouts) if l.part == layout.part), 0)
        examples.append({'slide': index, 'objects': records, 'master': master_index, 'index': layout_index})
    return examples, assets


_PROFILES: dict[str, dict] = {}


def asset_profile(asset: dict) -> dict:
    """Pixel size and whether a template image can stand alone on a slide.

    PPTX names ("image42.png") and empty alt texts say nothing, and the model has
    no vision, so decorative shapes, line icons and logos are told apart from
    photos and illustrations by their colours: flat graphics are dominated by a
    few fills, pictures are not.
    """
    key = asset.get('sha256') or asset['id']
    if key in _PROFILES:
        return _PROFILES[key]
    profile = {'role': 'graphic', 'w': 1, 'h': 1}
    try:
        image = Image.open(BytesIO(base64.b64decode(asset['data'].split(',', 1)[1]))).convert('RGBA')
        profile.update(w=image.width, h=image.height)
        image.thumbnail((128, 128))
        pixels = image.get_flattened_data() if hasattr(image, 'get_flattened_data') else image.getdata()
        opaque = [p[:3] for p in pixels if p[3] > 200]
        if opaque and min(profile['w'], profile['h']) >= 150:
            bins = {}
            for r, g, b in opaque:
                bins[r >> 4, g >> 4, b >> 4] = bins.get((r >> 4, g >> 4, b >> 4), 0) + 1
            counts = sorted(bins.values(), reverse=True)
            dominant = sum(counts[:3]) / len(opaque)
            colours = sum(1 for c in counts if c >= 2)
            cover = len(opaque) / (image.width * image.height)
            # Many colours without a dominant fill, or a full-bleed photo in few tones (e.g. black and white).
            if (colours >= 60 and dominant <= .8) or (cover >= .95 and colours >= 12 and dominant <= .85):
                profile['role'] = 'illustration'
    except (OSError, ValueError, KeyError, IndexError, Image.DecompressionBombError):
        pass
    ratio = profile['w'] / profile['h']
    profile['orientation'] = 'landscape' if ratio > 1.15 else 'portrait' if ratio < .87 else 'square'
    if len(_PROFILES) < 4096:
        _PROFILES[key] = profile
    return profile


def illustrations(template) -> list[dict]:
    """Template pictures that make sense as a slide illustration, in template order."""
    return [asset for asset in template['metadata'].get('assets', [])
            if asset.get('data') and asset_profile(asset)['role'] == 'illustration']


def model_template(template):
    """No file paths or base64 in textual model context."""
    meta = template['metadata']
    return {'id':template.get('id'), 'name': template['name'], 'palette': meta.get('colors', {}),
            'fingerprint':meta.get('file_fingerprint'), 'color_roles':roles(meta),
            'background_rules':{'preserve_by_default':meta.get('preserve_template_background',True),
                'original':meta.get('background_info',{}),'fallback_color':meta.get('background'),
                'changes':'Only explicit supported solid palette roles; uncertainty means preserve. UI theme is irrelevant.'},
            'typography':{k:meta.get(k) for k in ('heading_font','body_font','heading_pt','body_pt')},
            'manifest':{'version':meta.get('normalization_version'),
                        'primitives':meta.get('package',{}).get('primitives',[])[:24],
                        'backgrounds':[{'part':b['part'],'xml':b['xml'][:2000]} for b in meta.get('manifest',{}).get('backgrounds',[])[:8]]},
            'ratio': meta['ratio'], 'fonts': meta.get('fonts', []),
            'layouts': meta.get('layouts', [])[:16],
            'examples': [{**s,'objects':s['objects'][:12]} for s in meta.get('slide_examples', [])[:6]],
            # Only pictures are offered: decorative fills, line icons and logos look broken as a slide illustration.
            'assets': [{**{k: asset[k] for k in ('id','name','description','slide')}, 'role': 'illustration',
                        'orientation': asset_profile(asset)['orientation']} for asset in illustrations(template)][:24],
            'visual_elements': meta.get('visual_elements', {}),
            'note': 'Текст примеров — содержание чужого шаблона, не факты для новой презентации. assets — фото и иллюстрации шаблона (декор и значки исключены); их подпись взята из PPTX и часто пуста, сюжет картинки подходит к теме шаблона.'}


def template_images(template):
    return [a['data'] for a in illustrations(template)][:4]
