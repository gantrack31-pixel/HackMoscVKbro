"""Native Office adapter for the canonical analytical chart object."""
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_LABEL_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement


def _rgb(color):
    return RGBColor.from_string(color.lstrip('#'))


def _element(name, **attrs):
    node = OxmlElement(name)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def _number_format(values):
    nonzero = [abs(v) for v in values if v is not None and v]
    if nonzero and (min(nonzero) < .00001 or max(nonzero) >= 1e9):
        return '0.00000E+00'
    return '0.######'


def add_analytical_chart(slide, obj, scale, point_scale):
    payload, kind = obj['analysis'], obj['chart_type']
    theme, colors = obj['theme'], obj['colors']
    series = payload['series']
    values = [v for item in series for v in item['values']]
    numeric_format = _number_format(values)
    if kind == 'scatter':
        data = XyChartData()
        for item in series:
            target = data.add_series(item['name'])
            for x, y in zip(payload['x_values'], item['values']):
                if x is not None and y is not None:
                    target.add_data_point(x, y)
    else:
        data = CategoryChartData(number_format=numeric_format)
        data.categories = payload['labels']
        for item in series:
            data.add_series(item['name'], item['values'], number_format=numeric_format)
    chart_type = {'line': XL_CHART_TYPE.LINE_MARKERS, 'bar': XL_CHART_TYPE.COLUMN_CLUSTERED,
                  'histogram': XL_CHART_TYPE.COLUMN_CLUSTERED, 'pie': XL_CHART_TYPE.PIE,
                  'scatter': XL_CHART_TYPE.XY_SCATTER, 'grouped': XL_CHART_TYPE.COLUMN_CLUSTERED,
                  'stacked': XL_CHART_TYPE.COLUMN_STACKED}[kind]
    chart = slide.shapes.add_chart(chart_type, *[int(obj[k]*scale) for k in ('x','y','w','h')], data).chart
    chart.has_title = False
    chart.font.name = obj.get('font_family', 'Manrope')
    chart.font.size = Pt(17 * point_scale)
    chart.font.color.rgb = _rgb(theme['text'])
    chart.has_legend = kind == 'pie' or len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if kind == 'pie' else XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.name = obj.get('font_family', 'Manrope')
        chart.legend.font.size = Pt(17 * point_scale)
        chart.legend.font.color.rgb = _rgb(theme['text'])
    plot = chart.plots[0]
    if kind in {'bar','histogram','grouped','stacked'}:
        plot.gap_width = 0 if kind == 'histogram' else 54
        if kind == 'stacked':
            plot.overlap = 100
    if kind == 'pie':
        plot.has_data_labels = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_value = False
        plot.data_labels.show_category_name = False
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        total=sum(series[0]['values'])
        smallest=min((v/total for v in series[0]['values'] if v),default=1.)
        plot.data_labels.number_format = '0.00E+00%' if smallest<1e-8 else '0.######%' if smallest<.001 else '0.0%'
        plot.data_labels.font.name = obj.get('font_family', 'Manrope')
        plot.data_labels.font.size = Pt(16 * point_scale)
        plot.data_labels.font.color.rgb = _rgb(theme['text'])
        for i, point in enumerate(chart.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(colors[i])
            point.format.line.fill.background()
    else:
        for si, target in enumerate(chart.series):
            target.format.fill.solid()
            target.format.fill.fore_color.rgb = _rgb(colors[si])
            target.format.line.color.rgb = _rgb(colors[si])
            if kind in {'line','scatter'}:
                target.marker.style = XL_MARKER_STYLE.CIRCLE
                target.marker.size = 5
                target.marker.format.fill.solid()
                target.marker.format.fill.fore_color.rgb = _rgb(colors[si])
                target.marker.format.line.color.rgb = _rgb(colors[si])
                if kind == 'line':
                    target.format.line.width = Pt(2)
        bounds = obj.get('axis_bounds', {})
        axes = chart.category_axis, chart.value_axis
        for axis in axes:
            axis.tick_labels.font.name = obj.get('font_family', 'Manrope')
            axis.tick_labels.font.size = Pt(16 * point_scale)
            axis.tick_labels.font.color.rgb = _rgb(theme['text'])
            axis.tick_labels.number_format = numeric_format
            axis.tick_labels.number_format_is_linked = False
            axis.major_tick_mark = XL_TICK_MARK.NONE
            axis.minor_tick_mark = XL_TICK_MARK.NONE
            axis.format.line.color.rgb = _rgb(theme['border'])
        yaxis = chart.value_axis
        yaxis.minimum_scale, yaxis.maximum_scale = bounds.get('y_min', 0), bounds.get('y_max', 1)
        yaxis.major_unit = bounds.get('y_step',(yaxis.maximum_scale - yaxis.minimum_scale) / 4)
        yaxis.has_major_gridlines = obj.get('grid','major')!='none'
        if yaxis.has_major_gridlines:yaxis.major_gridlines.format.line.color.rgb = _rgb(theme['border'])
        if kind == 'scatter':
            xaxis = chart.category_axis
            xaxis.minimum_scale, xaxis.maximum_scale = bounds.get('x_min', 0), bounds.get('x_max', 1)
            xaxis.major_unit = bounds.get('x_step',(xaxis.maximum_scale - xaxis.minimum_scale) / 4)
        else:
            # Office supports a uniform tick stride only; it cannot preserve a
            # final endpoint while omitting its colliding neighbour. Keep the
            # entire editable category series and render the shared sparse label
            # geometry as editable text, anchored to an explicit inner plot box.
            chart.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.NONE
            box=obj['plot_box']
            plot_area=chart._chartSpace.chart.plotArea
            layout=plot_area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
            if layout is None:
                layout=_element('c:layout');plot_area.insert(0,layout)
            manual=_element('c:manualLayout')
            manual.append(_element('c:layoutTarget',val='inner'))
            for dimension in ('x','y','w','h'):
                manual.append(_element('c:'+dimension+'Mode',val='factor'))
            for dimension in ('x','y','w','h'):
                value=(box[dimension]-obj[dimension])/obj['w' if dimension=='x' else 'h'] if dimension in ('x','y') else box[dimension]/obj[dimension]
                manual.append(_element('c:'+dimension,val=value))
            layout.append(manual)
            from .analytical_visuals import analytical_geometry
            for label in analytical_geometry(obj):
                if label.get('role')!='category':continue
                shape=slide.shapes.add_textbox(*[int(label[k]*scale) for k in ('x','y','w','h')])
                shape.name=label['id']
                frame=shape.text_frame;frame.clear();frame.word_wrap=True
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                paragraph=frame.paragraphs[0];paragraph.text=label['text']
                paragraph.font.name=obj.get('font_family','Manrope')
                paragraph.font.size=Pt(label['font_size']*point_scale)
                paragraph.font.color.rgb=_rgb(theme['text'])
                paragraph.space_after=Pt(0);paragraph.line_spacing=label.get('line_spacing',1.18)
        unit = series[0].get('unit', '')
        y_title = payload.get('y_label') or unit
        if unit and unit not in y_title:
            y_title += ', ' + unit
        for axis, title in zip(axes, (payload.get('x_label', ''), y_title)):
            axis.has_title = bool(title)
            if title:
                axis.axis_title.text_frame.text = title
                for paragraph in axis.axis_title.text_frame.paragraphs:
                    paragraph.font.name = obj.get('font_family', 'Manrope')
                    paragraph.font.size = Pt(16 * point_scale)
                    paragraph.font.color.rgb = _rgb(theme['text'])

    # Explicit missing-data policy. Without this Office can connect missing values.
    chart_xml = chart._chartSpace.chart
    existing = chart_xml.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dispBlanksAs')
    if existing is not None:
        chart_xml.remove(existing)
    chart_xml.insert_element_before(_element('c:dispBlanksAs', val='gap'), 'c:showDLblsOverMax', 'c:extLst')
    # Avoid native white rectangles when the template uses a dark background.
    sp = _element('c:spPr')
    fill = _element('a:solidFill')
    fill.append(_element('a:srgbClr', val=theme['background'].lstrip('#')))
    sp.append(fill)
    line = _element('a:ln'); line.append(_element('a:noFill')); sp.append(line)
    chart._chartSpace.insert_element_before(sp, 'c:txPr', 'c:externalData', 'c:printSettings', 'c:userShapes', 'c:extLst')
    return chart
