"""Читаем PPTX как набор правил; исходный файл никогда не изменяем."""
from collections import Counter
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import re
from lxml import etree
from pptx import Presentation
from .. import database as db
from .theme_catalog import THEMES, theme_metadata

NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
PARSER = etree.XMLParser(resolve_entities=False, no_network=True)

def validate_pptx(path: Path):
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 12000 or sum(i.file_size for i in entries) > 350 * 1024**2:
                raise ValueError('Распакованный PPTX слишком велик для локальной версии.')
            if 'ppt/presentation.xml' not in archive.namelist():
                raise ValueError('В файле нет презентации PowerPoint.')
            if any('vbaProject' in i.filename for i in entries):
                raise ValueError('Поддерживаются PPTX без макросов.')
    except BadZipFile as exc:
        raise ValueError('Файл повреждён или не является PPTX.') from exc

def analyze_template(path: Path) -> dict:
    validate_pptx(path)
    prs = Presentation(str(path))
    width, height = int(prs.slide_width), int(prs.slide_height)
    if width <= 0 or height <= 0: raise ValueError('Некорректный размер слайда.')
    fonts, sizes = Counter(), Counter()
    colors, theme_font = {}, 'Arial'
    with ZipFile(path) as archive:
        theme_names = sorted(n for n in archive.namelist() if re.fullmatch(r'ppt/theme/theme\d+\.xml',n))
        if theme_names:
            root = etree.fromstring(archive.read(theme_names[0]), PARSER)
            scheme = root.find('.//a:clrScheme', NS)
            if scheme is not None:
                for color in scheme:
                    if len(color):
                        value = color[0].get('val') or color[0].get('lastClr')
                        if not value or not re.fullmatch(r'[0-9a-fA-F]{6}', value): value=color[0].get('lastClr')
                        if value and re.fullmatch(r'[0-9a-fA-F]{6}',value): colors[etree.QName(color).localname]='#'+value
            latin = root.find('.//a:majorFont/a:latin', NS)
            if latin is not None and latin.get('typeface'): theme_font = latin.get('typeface')
        for name in archive.namelist():
            if re.fullmatch(r'ppt/slides/slide\d+\.xml',name):
                root=etree.fromstring(archive.read(name), PARSER)
                for node in root.findall('.//a:rPr',NS):
                    latin=node.find('a:latin',NS)
                    if latin is not None and latin.get('typeface') and not latin.get('typeface').startswith('+'):
                        fonts[latin.get('typeface')] += 1
                    size=node.get('sz')
                    if size and size.isdigit(): sizes[round(int(size)/100,1)]+=1
    layouts=[]
    for master_index, master in enumerate(prs.slide_masters):
        for layout_index, layout in enumerate(master.slide_layouts):
            boxes=[]
            for shape in layout.placeholders:
                try:
                    box={'x':round(shape.left/width,4),'y':round(shape.top/height,4),
                         'w':round(shape.width/width,4),'h':round(shape.height/height,4),
                         'type':str(shape.placeholder_format.type),'index':shape.placeholder_format.idx}
                    if box['w']>0 and box['h']>0: boxes.append(box)
                except (TypeError,ValueError): continue
            layouts.append({'master':master_index,'index':layout_index,'name':layout.name,'boxes':boxes,
                            'static_shapes':sum(not s.is_placeholder for s in layout.shapes)})
    headings=[(size,count) for size,count in sizes.items() if 25<=size<=60]
    bodies=[(size,count) for size,count in sizes.items() if 16<=size<=25]
    return {'count':len(prs.slides),'ratio':round(width/height,5),'width_emu':width,'height_emu':height,
            'colors':colors,'accent':colors.get('accent1','#0077FF'),'background':colors.get('lt1','#FFFFFF'),
            'font':fonts.most_common(1)[0][0] if fonts else theme_font,
            'theme_font':theme_font,'fonts':[name for name,_ in fonts.most_common(6)],
            'heading_pt':max(headings,key=lambda x:x[1])[0] if headings else 32,
            'body_pt':max(bodies,key=lambda x:x[1])[0] if bodies else 18,
            'layouts':layouts,'master_count':len(prs.slide_masters),'source':'pptx',
            'warnings':['Предпросмотр показывает содержимое и размещение. Графика мастера сохраняется в PPTX; проверьте итоговый файл в PowerPoint.']}

BUILTINS=[
    ('tech','VK Tech','Корпоративные','#0077FF',54),('workspace','VK WorkSpace','Корпоративные','#0077FF',29),
    ('education','VK Education','Образование','#73A929',55),('pitch','Питч: большая идея','Стартапы','#7851BA',12),
    ('report','Квартал в цифрах','Отчёты','#3479BD',10),('portfolio','Портфолио','Креативные','#B95644',10),
    ('minimal','Чистый лист','Минимализм','#54636F',10),('startup','Продуктовый запуск','Стартапы','#29745F',12),
    ('research','Исследование','Образование','#7754C4',12)]
BUILTIN_IDS={item[0] for item in BUILTINS} | {item[0] for item in THEMES}

def seed_templates(storage: Path):
    folder=storage/'templates';folder.mkdir(parents=True,exist_ok=True)
    for tid,name,category,accent,count in BUILTINS:
        if db.template_get(tid): continue
        source=folder/f'{tid}.pptx'
        metadata=analyze_template(source) if source.exists() else {'count':count,'ratio':16/9,'font':'Manrope','heading_pt':32,
            'body_pt':18,'accent':accent,'background':'#FFFFFF','colors':{'accent1':accent},'layouts':[],
            'source':'starter','warnings':['Авторская стартовая тема; исходный корпоративный PPTX не загружен.']}
        metadata.update(category=category,cover_title={'tech':'Технологии, которые работают на вас','workspace':'Всё для команды. В одном пространстве.',
            'education':'Знания, которые меняют будущее','pitch':'Большая идея. Чёткий план.','report':'Результаты, которые говорят',
            'portfolio':'Работы говорят за нас.','minimal':'Меньше деталей. Больше смысла.','startup':'Следующий шаг начинается здесь',
            'research':'Вопросы. Данные. Открытия.'}[tid])
        db.template_save(tid,name,str(source) if source.exists() else None,metadata)
    for theme in THEMES:
        if not db.template_get(theme[0]):
            db.template_save(theme[0],theme[1],None,theme_metadata(theme))
