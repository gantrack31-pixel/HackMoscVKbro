"""Экспорт на Python. Тексты, таблицы и диаграммы PPTX остаются редактируемыми."""
from io import BytesIO
from html import escape
import base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from ..models import DeckContent
from .layout import build_scene, selected_layout, FONTS, geometry_nodes

def rgb(value): return RGBColor.from_string(value.lstrip('#'))

def scene_svg(scene):
    parts=[]
    for obj in scene['objects']:
        for n in geometry_nodes(obj):
            if n['type']=='rect': parts.append(f'<rect x="{n["x"]}" y="{n["y"]}" width="{n["w"]}" height="{n["h"]}" fill="{n["fill"]}"/>')
            else:
                spans=''.join(f'<tspan x="{n["x"]}" y="{n["y"]+n["font_size"]+i*n["font_size"]*1.28}">{escape(line)}</tspan>' for i,line in enumerate(n['lines']))
                parts.append(f'<text font-family="Manrope, Arial" font-size="{n["font_size"]}" font-weight="{750 if n["bold"] else 400}" fill="{n["color"]}">{spans}</text>')
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {scene["width"]} {scene["height"]}" role="img">'+''.join(parts)+'</svg>'

def export_html(content,template,variant):
    font=base64.b64encode((FONTS/'Manrope-Regular.ttf').read_bytes()).decode()
    slides=''.join('<section>'+scene_svg(build_scene(slide,template['metadata'],variant,i))+'</section>' for i,slide in enumerate(content.slides))
    html=f'<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{escape(content.title)}</title><style>@font-face{{font-family:Manrope;src:url(data:font/ttf;base64,{font})}}body{{background:#e2e8f0;margin:0;padding:24px}}section{{max-width:1280px;margin:0 auto 24px;break-after:page}}svg{{display:block;width:100%;background:white}}@media print{{body{{padding:0}}section{{margin:0}}@page{{size:landscape;margin:0}}}}</style>{slides}</html>'
    return html.encode('utf-8')

def export_pdf(content,template,variant):
    stream=BytesIO();metadata=template['metadata'];height=1280/metadata['ratio']
    pdf=canvas.Canvas(stream,pagesize=(1280,height));pdf.setTitle(content.title)
    for i,slide in enumerate(content.slides):
        scene=build_scene(slide,metadata,variant,i)
        for obj in scene['objects']:
            for n in geometry_nodes(obj):
                if n['type']=='rect':
                    pdf.setFillColor(HexColor(n['fill']));pdf.rect(n['x'],height-n['y']-n['h'],n['w'],n['h'],fill=1,stroke=0)
                else:
                    pdf.setFont('DecklyBold' if n['bold'] else 'Deckly',n['font_size']);pdf.setFillColor(HexColor(n['color']))
                    for j,line in enumerate(n['lines']):pdf.drawString(n['x'],height-n['y']-n['font_size']-j*n['font_size']*1.28,line)
        pdf.showPage()
    pdf.save();return stream.getvalue()

def export_pptx(content: DeckContent, template: dict, variant: str):
    metadata=template['metadata'];prs=Presentation(template['path']) if template.get('path') else Presentation()
    if template.get('path'):
        # В python-pptx пока нет публичного remove_slide. Удаляем только слайды копии,
        # сохраняя мастера, макеты, тему и связанные с ними ресурсы.
        for slide_id in list(prs.slides._sldIdLst):
            prs.part.drop_rel(slide_id.rId);prs.slides._sldIdLst.remove(slide_id)
    else:
        prs.slide_width=Inches(13.333333);prs.slide_height=Inches(13.333333/metadata['ratio'])
    chosen=selected_layout(metadata)
    layout=prs.slide_masters[chosen['master']].slide_layouts[chosen['index']] if chosen else min(prs.slide_layouts,key=lambda l:len(l.placeholders))
    scale=prs.slide_width/1280;point_scale=scale/12700
    for index,slide_content in enumerate(content.slides):
        slide=prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            element=ph._element;element.getparent().remove(element)
        scene=build_scene(slide_content,metadata,variant,index)
        for obj in scene['objects']:
            x,y,w,h=[int(obj[key]*scale) for key in ['x','y','w','h']]
            if obj['id']=='background':
                slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(obj['fill']);continue
            if obj['type']=='rect':
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
                shape.fill.solid();shape.fill.fore_color.rgb=rgb(obj['fill']);shape.line.fill.background()
            elif obj['type']=='text':
                shape=slide.shapes.add_textbox(x,y,w,h);frame=shape.text_frame;frame.clear()
                frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                frame.word_wrap=True
                for j,line in enumerate(obj['lines']):
                    paragraph=frame.paragraphs[0] if j==0 else frame.add_paragraph()
                    paragraph.text=line;paragraph.font.name=metadata.get('font','Manrope')
                    paragraph.font.size=Pt(obj['font_size']*point_scale);paragraph.font.bold=obj['bold']
                    paragraph.font.color.rgb=rgb(obj['color']);paragraph.space_after=Pt(0);paragraph.line_spacing=1.28
            elif obj['type']=='chart':
                data=CategoryChartData();data.categories=obj['labels'];data.add_series(obj['unit'] or 'Значение',obj['values'])
                chart=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,x,y,w,h,data).chart
                chart.has_legend=False;chart.value_axis.has_title=True;chart.value_axis.axis_title.text_frame.text=obj['unit'] or 'Значение'
                chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=rgb(obj['fill'])
            elif obj['type']=='table':
                rows=obj['rows'];table=slide.shapes.add_table(len(rows),len(rows[0]),x,y,w,h).table
                for ri,row in enumerate(rows):
                    for ci,value in enumerate(row):
                        cell=table.cell(ri,ci);cell.text=value
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.size=Pt(obj['font_size']*point_scale);paragraph.font.name=metadata.get('font','Manrope')
                            paragraph.font.bold=ri==0;paragraph.font.color.rgb=rgb(obj['header_color'] if ri==0 else '#0F172A')
                        cell.fill.solid();cell.fill.fore_color.rgb=rgb(obj['fill'] if ri==0 else '#F1F5F9' if ri%2==0 else '#FFFFFF')
        slide.notes_slide.notes_text_frame.text=slide_content.notes+('\nИсточник: '+slide_content.source_quote if slide_content.source_quote else '')
    stream=BytesIO();prs.save(stream);return stream.getvalue()
