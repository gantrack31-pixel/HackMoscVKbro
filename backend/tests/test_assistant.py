from io import BytesIO
from uuid import uuid4
import asyncio
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from app import database as db
from app.config import settings
from app.main import app
from app.models import DeckContent, Slide
from app.services import llm
from app.services.template_context import model_template
from app.services.layout import build_scene
from app.services.export import export_pptx


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'database', tmp_path/'assistant.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path/'data')
    monkeypatch.setattr(settings, 'mode', 'live')
    monkeypatch.setattr(settings, 'base_url', 'http://127.0.0.1:8001/v1')
    monkeypatch.setattr(settings, 'email_verification_required', False)
    with TestClient(app) as c:
        auth = c.post('/api/auth/register', json={'email':'assistant@example.test', 'password':'Assistant-Test-2037!',
                    'first_name':'Test','last_name':'Editor'}).json()
        c.headers['X-CSRF-Token'] = auth['csrf']
        content = DeckContent(title='Пилот',slides=[Slide(title='TODO',body='Много текста. '*80),
                                                   Slide(title='Итог',body='30 участников.')])
        db.project_create('ai-test','tech',content.model_dump(),'Пилот охватил 30 участников.',auth['user']['id'])
        yield c


def request_for(client, **kwargs):
    p=client.get('/api/projects/ai-test').json()
    return {'content':p['content'],'variant':'a','base_updated_at':p['updated_at'],
            'instruction':'Сократи текст','slide':0,**kwargs}


def test_edit_sees_unsaved_text_and_undo_restores_it(client,monkeypatch):
    body=request_for(client)
    body['content']['slides'][0]['body']='Моя новая ручная правка — 30 участников.'
    async def respond(system,payload):
        assert payload['presentation']['slides'][0]['body']==body['content']['slides'][0]['body']
        assert payload['allowed_slides']==[0]
        assert 'geometry' in payload and 'template' in payload
        assert 'image_data' not in payload['presentation']['slides'][0]
        return {'summary':'Сократил текст.', 'changes':[{'slide':0,'title':'Пилот','body':'30 участников.'}]}
    monkeypatch.setattr(llm,'complete_json',respond)
    response=client.post('/api/projects/ai-test/assistant',json=body)
    assert response.status_code==200,response.text
    p=response.json()['project']
    assert p['content']['slides'][0]['body']=='30 участников.'
    assert p['content']['slides'][1]==body['content']['slides'][1]
    assert p['can_undo'] and response.json()['changed_slides']==[0]
    restored=client.post('/api/projects/ai-test/undo').json()
    assert restored['content']==body['content']


@pytest.mark.parametrize('changes', [
    [{'slide':1,'body':'Не разрешено'}],
    [{'slide':0,'body':'Один'},{'slide':0,'body':'Два'}],
    [{'slide':0,'template_asset_id':'foreign-asset','kind':'image'}],
    [{'slide':0,'kind':'chart'}],
    [{'slide':0,'image_data':'data:bad'}],
])
def test_invalid_or_out_of_scope_patches_never_save(client,monkeypatch,changes):
    body=request_for(client)
    async def respond(*args): return {'summary':'Правка','changes':changes}
    monkeypatch.setattr(llm,'complete_json',respond)
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==502
    p=client.get('/api/projects/ai-test').json()
    assert p['content']==body['content'] and not p['can_undo']


def test_concurrent_save_wins_and_ai_does_not_overwrite(client,monkeypatch):
    body=request_for(client)
    async def respond(*args):
        latest=body['content'].copy();latest['title']='Сохранено другой вкладкой'
        db.project_update('ai-test',latest,'b')
        return {'summary':'Правка','changes':[{'slide':0,'title':'AI title'}]}
    monkeypatch.setattr(llm,'complete_json',respond)
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==409
    assert client.get('/api/projects/ai-test').json()['title']=='Сохранено другой вкладкой'


def test_noop_preserves_unsaved_draft_and_provider_failure(client,monkeypatch):
    body=request_for(client);body['content']['slides'][0]['body']='Несохранённый текст'
    async def noop(*args): return {'summary':'Изменения не нужны','changes':[]}
    # edit-mode can apply deterministic fixes, so use a clean draft.
    body['content']['slides'][0]['title']='Пилот'
    monkeypatch.setattr(llm,'complete_json',noop)
    response=client.post('/api/projects/ai-test/assistant',json=body)
    assert response.status_code==200 and response.json()['project'] is None
    async def fail(*args): raise llm.LLMError('Провайдер недоступен')
    monkeypatch.setattr(llm,'complete_json',fail)
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==502
    assert not client.get('/api/projects/ai-test').json()['can_undo']


def test_repair_reaudits_and_access_is_owner_scoped(client,monkeypatch):
    body=request_for(client, action='repair', slide=None)
    async def respond(system,payload):
        assert any(i['code']=='placeholder' for i in payload['audit_issues'])
        return {'summary':'Убрал заглушку и лишний текст.', 'changes':[{'slide':0,'title':'Пилот','body':'30 участников.'}]}
    monkeypatch.setattr(llm,'complete_json',respond)
    response=client.post('/api/projects/ai-test/assistant',json=body)
    assert response.status_code==200
    assert not any(i['code']=='placeholder' for i in response.json()['issues'])
    assert client.post('/api/projects/ai-test/assistant',json=body,headers={'X-CSRF-Token':'bad'}).status_code==403
    auth=client.post('/api/auth/register',json={'email':'other-ai@example.test','password':'Assistant-Test-2037!',
                   'first_name':'Other','last_name':'User'}).json()
    client.headers['X-CSRF-Token']=auth['csrf']
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==404


def test_demo_and_stale_version_are_explicit(client,monkeypatch):
    body=request_for(client)
    monkeypatch.setattr(settings,'mode','demo')
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==409
    body['base_updated_at']='old'
    assert client.post('/api/projects/ai-test/assistant',json=body).status_code==409


def test_unseen_pptx_extracts_free_boxes_tables_charts_and_reusable_picture(client):
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(.5),Inches(7),Inches(.6)).text='Неизвестный макет'
    slide.shapes.add_textbox(Inches(1),Inches(2),Inches(6),Inches(2)).text='Содержимое из произвольного PPTX'
    slide.shapes.add_table(2,2,Inches(1),Inches(4),Inches(4),Inches(1)).table.cell(0,0).text='План'
    chart=CategoryChartData();chart.categories=['A','B'];chart.add_series('Пилот',[12,8])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6),Inches(2), Inches(3),Inches(2),chart)
    picture=BytesIO();Image.new('RGB',(48,48),'#317bee').save(picture,format='PNG');picture.seek(0)
    slide.shapes.add_picture(picture, Inches(8), Inches(.5), width=Inches(1))
    binary=BytesIO();prs.save(binary)
    response=client.post('/api/templates',files={'file':(str(uuid4())+'.pptx',binary.getvalue(),'application/octet-stream')})
    assert response.status_code==201,response.text
    t=response.json();meta=t['metadata']
    assert {o['type'] for o in meta['slide_examples'][0]['objects']}=={'text','table','chart','picture'}
    assert any(l['name']=='Композиция исходного слайда 1' for l in meta['layouts'])
    assert len(meta['assets'])==1
    context=model_template(t)
    assert 'data' not in context['assets'][0] and 'path' not in context
    content=DeckContent(title='Reuse',slides=[Slide(title='Картинка',kind='image',template_asset_id=meta['assets'][0]['id'])])
    scene=build_scene(content.slides[0],meta,'a',0)
    assert any(o['type']=='image' for o in scene['objects'])
    exported=Presentation(BytesIO(export_pptx(content,db.template_get(t['id']),'a')))
    assert any(s.shape_type==13 for s in exported.slides[0].shapes)
    packet=client.post('/api/materials/import',files={'files':('material.pptx',binary.getvalue(),'application/octet-stream')})
    assert packet.status_code==200,packet.text
    assert 'Неизвестный макет' in packet.json()['text'] and 'Пилот: 12.0, 8.0' in packet.json()['text']
