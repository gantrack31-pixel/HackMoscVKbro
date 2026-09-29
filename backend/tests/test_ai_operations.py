"""Cancellation, publication races and the optimized generation contract."""
import asyncio
from copy import deepcopy
from io import BytesIO
from uuid import uuid4
import httpx
import pytest
from pptx import Presentation
from app import main, database as db
from app.config import settings
from app.models import DeckContent, Slide, GenerateRequest
from app.services import images, llm, operations, provider_http, workflow
from test_assistant import client, request_for


@pytest.mark.parametrize('kind', ['outline','assistant','image','audit'])
def test_direct_stop_cancels_backend_and_releases_slot(client, monkeypatch, kind):
    original=db.project_get('ai-test')
    body=request_for(client)
    routes={
        'outline':('/api/outline', {'template_id':'tech','prompt':'A real outline','count':3}, 'make_outline'),
        'assistant':('/api/projects/ai-test/assistant', body, 'edit_presentation'),
        'image':('/api/images/generate', {'prompt':'An original illustration'}, 'generate_image'),
        'audit':('/api/projects/ai-test/audit/content', {}, 'review_content'),
    }
    async def check():
        started=asyncio.Event(); stopped=asyncio.Event()
        async def slow(*args, **kwargs):
            started.set()
            try: await asyncio.Future()
            finally: stopped.set()
        path,payload,symbol=routes[kind]
        monkeypatch.setattr(main,symbol,slow)
        monkeypatch.setattr(main,'generation_slots',asyncio.Semaphore(1))
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='http://testserver',
                                     cookies=client.cookies,headers=dict(client.headers)) as api:
            jid=(await api.post('/api/ai/operations')).json()['job_id']
            task=asyncio.create_task(api.post(path,json=payload,headers={'X-AI-Operation-ID':jid}))
            await asyncio.wait_for(started.wait(),3)
            response=await api.post(f'/api/jobs/{jid}/cancel')
            assert response.json()['state']=='cancelled'
            result=await asyncio.wait_for(task,3)
            assert result.status_code==409, result.text
            assert result.json()['detail']['code']=='AI_CANCELLED'
            assert stopped.is_set() and main.generation_slots._value==1
            assert (await api.get(f'/api/jobs/{jid}')).json()['state']=='cancelled'
    asyncio.run(check())
    assert db.project_get('ai-test')==original


@pytest.mark.parametrize('regenerate', [False,True])
def test_cancel_full_work_stops_pipeline_without_publication(client,monkeypatch,regenerate):
    original=db.project_get('ai-test');uid=original['user_id'];jid=str(uuid4())
    db.job_create(jid,uid,'regenerate' if regenerate else 'generate',{})
    async def check():
        started=asyncio.Event();stopped=asyncio.Event()
        async def design(*args):
            started.set()
            try: await asyncio.Future()
            finally: stopped.set()
        monkeypatch.setattr(main,'create_design_variants',design)
        monkeypatch.setattr(main,'generation_slots',asyncio.Semaphore(1))
        work=main.regenerate_job(jid,original,'New layout',uid,'never-published') if regenerate else main.generate_job(
            jid,GenerateRequest(template_id='tech',content=DeckContent(**original['content'])),uid,'never-published')
        task=asyncio.create_task(work)
        await asyncio.wait_for(started.wait(),3)
        assert db.job_cancel(jid,uid)=='cancelled'
        await asyncio.wait_for(task,3)
        assert stopped.is_set() and main.generation_slots._value==1
    asyncio.run(check())
    assert db.job_get(jid)['state']=='cancelled'
    assert db.project_get('never-published') is None
    assert db.project_get('ai-test')==original


def test_cancel_ownership_queued_terminal_and_atomic_publish_guards(client):
    original=db.project_get('ai-test');uid=original['user_id'];jid=str(uuid4())
    db.job_create(jid,uid,'ai',{})
    assert db.job_cancel(jid,'someone-else') is None
    db.job_create('foreign','someone-else','ai',{})
    assert client.post('/api/jobs/foreign/cancel').status_code==404
    assert client.post(f'/api/jobs/{jid}/cancel').json()['state']=='cancelled'
    assert not db.job_start(jid,uid)
    db.job_set(jid,'failed','failed',error='late exception')
    assert db.job_get(jid)['state']=='cancelled'
    assert not db.project_ai_update('ai-test',uid,original['updated_at'],original['content'],
                                    original['content'],'a',jid)
    with pytest.raises(ValueError):
        db.job_complete_with_project(jid,'never','tech',original['content'],'',uid)
    assert db.project_get('never') is None
    assert client.post('/api/outline',headers={'X-AI-Operation-ID':jid},
                       json={'template_id':'tech','prompt':'Plan','count':3}).json()['detail']['code']=='AI_CANCELLED'
    completed=str(uuid4());db.job_create(completed,uid,'ai',{});db.job_start(completed,uid)
    assert db.project_ai_update('ai-test',uid,original['updated_at'],original['content'],
                                original['content'],'a',completed)
    assert db.job_cancel(completed,uid)=='complete'
    db.job_set(completed,'failed','failed')
    assert db.job_get(completed)['state']=='complete'


def test_business_deadlines_only_full_builds(client,monkeypatch):
    limits=[]
    async def spy(jid,factory,*,request=None,business_limit=None):
        limits.append(business_limit)
        return {'ok':True}
    monkeypatch.setattr(operations,'run_operation',spy)
    asyncio.run(main.generate_job('id',None,'owner'))
    asyncio.run(main.regenerate_job('id',None,'instruction','owner'))
    client.post('/api/outline',json={'template_id':'tech','prompt':'Plan','count':3})
    client.post('/api/projects/ai-test/assistant',json=request_for(client))
    client.post('/api/images/generate',json={'prompt':'Original image'})
    client.post('/api/projects/ai-test/audit/content')
    assert limits==[300,300,None,None,None,None]
    assert workflow.OUTLINE_BUDGET_SECONDS is None


def test_deadline_really_cancels_and_prevents_publish(client):
    uid=db.project_get('ai-test')['user_id'];jid=str(uuid4());db.job_create(jid,uid,'generate',{})
    async def check():
        stopped=asyncio.Event()
        async def slow():
            db.job_start(jid,uid)
            try: await asyncio.Future()
            finally: stopped.set()
        await operations.run_operation(jid,slow,business_limit=.02)
        assert stopped.is_set()
    asyncio.run(check())
    assert db.job_get(jid)['state']=='failed'
    assert db.job_get(jid)['project_id'] is None


def test_expired_deadline_and_late_provider_cannot_publish(client):
    original=db.project_get('ai-test');uid=original['user_id']
    jid=str(uuid4());db.job_create(jid,uid,'generate',{});db.job_start(jid,uid)
    token=operations.operation_deadline.set(0)
    try:
        with pytest.raises(TimeoutError):
            db.job_complete_with_project(jid,'late-result','tech',original['content'],'',uid)
    finally: operations.operation_deadline.reset(token)
    assert db.project_get('late-result') is None
    async def check():
        async def late():
            db.job_cancel(jid,uid)  # cancellation wins just as a provider returns
            return {'late':'response'}
        with pytest.raises(main.HTTPException) as exc:
            await operations.run_operation(jid,late,request=object())
        assert exc.value.detail['code']=='AI_CANCELLED'
    asyncio.run(check())
    assert db.job_get(jid)['state']=='cancelled'


def test_images_bounded_and_children_cancelled(monkeypatch):
    monkeypatch.setattr(settings,'image_concurrency',2)
    async def check(cancel):
        active=peak=finished=0;started=asyncio.Event();release=asyncio.Event()
        async def generate(prompt):
            nonlocal active,peak,finished
            active+=1;peak=max(peak,active)
            if active==2: started.set()
            try:
                await release.wait()
                return {'image_data':'test-data'}
            finally: active-=1;finished+=1
        monkeypatch.setattr(images,'generate_image',generate)
        deck=DeckContent(title='Images',slides=[Slide(title=str(i),kind='image',image_prompt='original') for i in range(6)])
        task=asyncio.create_task(images.fill_missing_images(deck))
        await asyncio.wait_for(started.wait(),2)
        if cancel:
            task.cancel()
            with pytest.raises(asyncio.CancelledError): await task
            assert finished==2 and not any(s.image_data for s in deck.slides)
        else:
            release.set();await task
            assert finished==6 and all(s.image_data for s in deck.slides)
        assert peak==2 and active==0
    asyncio.run(check(False));asyncio.run(check(True))


def test_full_pipeline_audits_timings_no_discarded_exports(client,monkeypatch):
    monkeypatch.setattr(settings,'mode','demo')
    original=db.project_get('ai-test');content=deepcopy(original['content']);audited=[]
    real_audit=main.audit_deck;real_export=main.export_pptx
    def audit(*args,**kwargs):
        assert kwargs.get('scenes')
        audited.append(args[2]);return real_audit(*args,**kwargs)
    def unused_export(*args,**kwargs): raise AssertionError('PPTX should only be exported on demand')
    monkeypatch.setattr(main,'audit_deck',audit);monkeypatch.setattr(main,'export_pptx',unused_export)
    response=client.post('/api/generate',json={'template_id':'tech','content':content})
    jid=response.json()['job_id'];job=client.get(f'/api/jobs/{jid}').json()
    assert job['state']=='complete',job
    pid=job['project_id'];project=client.get(f'/api/projects/{pid}').json()
    for source,result in zip(content['slides'],project['content']['slides']):
        for key in ['title','body','bullets','chart','table','icon_names']:
            assert source[key]==result[key]
    provenance=project['provenance'];times=provenance['timings']
    assert set(times)=={'template','design','images','layout','audit','persist','total'}
    assert all(v>=0 for v in times.values())
    assert abs(provenance['duration_seconds']-times['total'])<.011
    assert set(provenance['audit'])==set('abc') and sorted(audited)==list('abc')
    audited.clear()
    response=client.post(f'/api/projects/{pid}/regenerate',json={'instruction':'More whitespace'})
    regen=client.get('/api/jobs/'+response.json()['job_id']).json()
    assert regen['state']=='complete',regen
    assert sorted(audited)==list('abc')
    assert client.get(f'/api/projects/{pid}').json()==project
    monkeypatch.setattr(main,'export_pptx',real_export)
    output=client.get(f'/api/projects/{regen["project_id"]}/export/pptx')
    assert output.status_code==200
    assert len(Presentation(BytesIO(output.content)).slides)==len(content['slides'])


def test_provider_pool_and_task_tokens_preserve_custom_client(monkeypatch):
    monkeypatch.setattr(settings,'base_url','http://127.0.0.1:8001/v1')
    seen=[]
    async def response(request):
        import json
        body=json.loads(request.content);seen.append(body['max_tokens'])
        assert request.extensions['timeout']['connect']==15
        assert request.extensions['timeout']['read']==settings.timeout
        return httpx.Response(200,json={'choices':[{'message':{'content':'{}'}}]})
    async def check():
        await provider_http.start();pool=provider_http.get()
        assert pool is provider_http.get() and not pool.is_closed
        async with httpx.AsyncClient(transport=httpx.MockTransport(response)) as custom:
            for schema in ['DeckContent','ThreeDesignPlans','DesignPlan','AssistantPlan']:
                await llm.complete_json('test',{'schema':{'title':schema}},client=custom)
            await llm.complete_json('test',{'source_chunks':[]},client=custom)
            assert not custom.is_closed
        await provider_http.close()
        assert pool.is_closed and provider_http.get() is None
    asyncio.run(check())
    assert seen==[settings.outline_max_tokens,settings.design_max_tokens,settings.design_max_tokens,
                  settings.assistant_max_tokens,settings.audit_max_tokens]
