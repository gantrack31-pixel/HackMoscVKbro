"""Startup migrates stored metadata before a catalog request or generation."""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import database as db
from app.config import settings
from app.main import app
from app.services import templates
from test_workflow import register


@pytest.mark.parametrize('stale', ['version', 'assets', 'fingerprint'])
def test_startup_refreshes_private_template_assets_and_keeps_catalog_identity(tmp_path, monkeypatch, stale):
    monkeypatch.setattr(settings, 'database', tmp_path / 'catalog.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path / 'storage')
    monkeypatch.setattr(settings, 'email_verification_required', False)
    monkeypatch.setattr(templates, 'seed_pack', lambda: None)
    path = tmp_path / 'source.pptx'
    picture = BytesIO()
    Image.new('RGB', (80, 60), '#315AAA').save(picture, format='PNG')
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6]).shapes.add_picture(picture, Inches(1), Inches(1))
    presentation.save(path)
    fresh = templates.analyze_template(path)
    original = {**fresh, 'category': 'Мои шаблоны', 'cover_title': 'Моя обложка'}
    if stale == 'version': original.update(normalization_version='old', assets=[])
    elif stale == 'assets': original.pop('assets')
    else: original.update(file_fingerprint='outdated', assets=[])

    with TestClient(app) as client:
        owner = register(client)['id']
        db.template_save('private-refresh', 'Мой шаблон', str(path), original, owner)
        with db.connection() as connection:
            connection.execute("UPDATE templates SET created_at='2000-01-01T00:00:00Z' WHERE id='private-refresh'")

    with TestClient(app) as client:
        stored = db.template_get('private-refresh')
        assert stored['metadata']['assets'] == fresh['assets']
        assert stored['metadata']['file_fingerprint'] == fresh['file_fingerprint']
        assert stored['metadata']['normalization_version'] == fresh['normalization_version']
        assert stored['user_id'] == owner and stored['path'] == str(path)
        assert stored['created_at'] == '2000-01-01T00:00:00Z'
        assert stored['metadata']['cover_title'] == 'Моя обложка'
        assert stored['metadata']['category'] == 'Мои шаблоны'
        monkeypatch.setattr('app.main.analyze_template', lambda *_: pytest.fail('Catalog reparsed a PPTX'))
        login = client.post('/api/auth/login', json={'email': 'anna@example.test', 'password': 'Example-Password-1947'})
        assert login.status_code == 200
        response = client.get('/api/templates?summary=true')
        summary = next(item for item in response.json() if item['id'] == 'private-refresh')
        assert summary['metadata']['assets'][0]['data'].startswith('/api/templates/private-refresh/assets/')


@pytest.mark.parametrize('source_exists', [False, True], ids=['missing', 'corrupt'])
def test_startup_retains_metadata_when_source_is_unavailable(tmp_path, monkeypatch, source_exists):
    monkeypatch.setattr(settings, 'database', tmp_path / 'catalog.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path / 'storage')
    monkeypatch.setattr(templates, 'seed_pack', lambda: None)
    db.initialize()
    if source_exists: (tmp_path / 'missing.pptx').write_bytes(b'broken archive')
    metadata = {'normalization_version': 'old', 'assets': [{'id': 'saved-image', 'data': ''}]}
    db.template_save('missing', 'Missing', str(tmp_path / 'missing.pptx'), metadata, 'owner')
    with TestClient(app):
        assert db.template_get('missing')['metadata'] == metadata


def test_startup_does_not_reanalyze_current_template_metadata(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'database', tmp_path / 'catalog.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path / 'storage')
    monkeypatch.setattr(templates, 'seed_pack', lambda: None)
    path = tmp_path / 'current.pptx'
    Presentation().save(path)
    metadata = templates.analyze_template(path)
    db.initialize()
    db.template_save('current', 'Current', str(path), metadata, 'owner')
    monkeypatch.setattr('app.main.analyze_template', lambda *_: pytest.fail('Unchanged source was reanalyzed'))

    with TestClient(app):
        assert db.template_get('current')['metadata'] == metadata
