"""Regression evidence for all nine analytical visualizations and three adapters."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pytest
from pypdf import PdfReader
from pptx import Presentation
from app.models import Slide, DeckContent
from app.services.layout import build_scene
from app.services.export import export_pptx, export_html, export_pdf
from app.services.audit import audit_deck
from app.services.analytical_visuals import number_label

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
KINDS = ['table','line','bar','histogram','pie','scatter','grouped','stacked','kpi']


def analytical_slide(kind):
    visual = {'chart_type':kind,'labels':['Январь','Февраль','Март','Апрель'],
              'series':[{'name':'Выручка','values':[10,15,12,20],'unit':'млн руб.'}],
              'x_label':'Месяц','y_label':'Выручка','source':'Синтетический набор для проверки, 2026',
              'question':'Как меняется выручка?','takeaway':'Выручка выросла с 10 до 20 млн руб.',
              'dataset_fingerprint':'a'*64,'source_fingerprint':'b'*64,
              'why':'Месячные наблюдения одной величины позволяют сравнить динамику.'}
    if kind in {'grouped','stacked'}:
        visual['series'].append({'name':'Подписки','values':[4,5,6,8],'unit':'млн руб.'})
    elif kind == 'histogram':
        visual.update(labels=['[0; 10)','[10; 20)','[20; 30)','[30; 40]'],
                      series=[{'name':'Наблюдения','values':[8,14,10,4],'unit':'шт.'}],
                      x_label='Интервал значения',y_label='Частота')
    elif kind == 'scatter':
        visual.update(x_values=[2,3,5,7],x_label='Расходы, млн руб.')
    elif kind == 'table':
        visual.update(columns=['Месяц','Выручка, млн руб.'],rows=[['Январь','10'],['Февраль','15'],['Март','12'],['Апрель','20']])
    elif kind == 'kpi':
        visual['kpis']=[{'label':'Выручка','value':57,'unit':'млн руб.','aggregation':'Сумма','period':'Январь–апрель 2026'},
                        {'label':'Среднее за месяц','value':14.25,'unit':'млн руб.','aggregation':'Среднее','period':'Январь–апрель 2026'}]
    return Slide(title=visual['takeaway'],kind='analysis',analytics=visual)


def template(background='#FFFFFF',ratio=16/9):
    return {'metadata':{'ratio':ratio,'font':'Manrope','background':background,'accent':'#006AD8',
                        'colors':{'accent1':'#006AD8','accent2':'#0B8C72','dk1':'#142A45','lt1':'#FFFFFF'},'layouts':[]}}


@pytest.mark.parametrize('kind',KINDS)
@pytest.mark.parametrize('variant',list('abc'))
@pytest.mark.parametrize('background',['#FFFFFF','#111827'])
def test_all_analytical_scenes_have_auditable_labels_and_no_clipping(kind,variant,background):
    slide=analytical_slide(kind);original=slide.model_dump();t=template(background)
    content=DeckContent(title='Проверка',slides=[slide])
    scene=build_scene(slide,t['metadata'],variant,0)
    assert slide.model_dump()==original
    assert len({o['id'] for o in scene['render_objects']})==len(scene['render_objects'])
    issues=audit_deck(content,t,variant,'',scenes=[scene])
    assert not [i for i in issues if i.severity=='error'],[(i.code,i.object_id) for i in issues]
    assert not [i for i in issues if i.code.startswith('contrast')]
    assert any(o['id']=='analysis-source' for o in scene['render_objects'])
    if kind not in {'table','kpi','pie'}:
        assert sum(o['role']=='axis-tick' for o in scene['render_objects'] if 'role' in o)>=3


def test_compositions_change_layout_but_keep_complete_content_and_data():
    slide=analytical_slide('grouped');t=template()
    scenes=[build_scene(slide,t['metadata'],v,0) for v in 'abc']
    charts=[next(o for o in s['objects'] if o.get('analytical')) for s in scenes]
    assert len({tuple(c[k] for k in ('x','y','w','h')) for c in charts})==3
    assert charts[0]['analysis']==charts[1]['analysis']==charts[2]['analysis']


@pytest.mark.parametrize('kind',KINDS)
def test_export_roundtrip_contains_native_editable_data_and_vector_pdf(kind):
    slide=analytical_slide(kind);content=DeckContent(title='Анализ',slides=[slide]);t=template()
    output=export_pptx(content,t,'a')
    prs=Presentation(BytesIO(output))
    assert len(prs.slides)==1
    shapes=prs.slides[0].shapes
    with ZipFile(BytesIO(output)) as package:
        assert not any(name.startswith('ppt/media/') for name in package.namelist())
        charts=[name for name in package.namelist() if name.startswith('ppt/charts/chart') and name.endswith('.xml')]
        if kind=='table':
            native=next(s.table for s in shapes if s.has_table)
            assert len(native.rows)==5 and native.cell(4,1).text=='20'
        elif kind=='kpi':
            assert any('14,25' in s.text for s in shapes if s.has_text_frame)
        else:
            assert len(charts)==1
            xml=ET.fromstring(package.read(charts[0]))
            tag={'line':'lineChart','pie':'pieChart','scatter':'scatterChart'}.get(kind,'barChart')
            assert xml.find('.//c:'+tag,NS) is not None
            assert xml.find('.//c:dispBlanksAs',NS).attrib['val']=='gap'
            assert any(name.startswith('ppt/embeddings/') for name in package.namelist())
            if kind=='histogram':assert xml.find('.//c:gapWidth',NS).attrib['val']=='0'
            if kind=='stacked':assert xml.find('.//c:grouping',NS).attrib['val']=='stacked'
            if kind in {'grouped','stacked'}:assert len(xml.findall('.//c:ser',NS))==2
    pdf=PdfReader(BytesIO(export_pdf(content,t,'a')))
    assert len(pdf.pages)==1
    assert 'Выручка' in pdf.pages[0].extract_text()
    assert not list(pdf.pages[0].images)
    html=export_html(content,t,'a').decode()
    assert '<svg ' in html and 'Синтетический набор' in html and 'data:font/ttf;base64,' in html
    assert 'http://' not in html.replace('http://www.w3.org/2000/svg','')


def test_line_gaps_tiny_values_and_negative_baseline_are_preserved():
    slide=analytical_slide('line')
    slide.analytics.series[0].values=[.00000012,None,-.00000024,.00000036]
    scene=build_scene(slide,template()['metadata'],'a',0)
    chart=next(o for o in scene['objects'] if o['type']=='chart')
    assert chart['axis_bounds']['y_min']<=-.00000024
    assert chart['axis_bounds']['y_max']>=.00000036
    assert chart['axis_bounds']['y_step']==pytest.approx(.0000002)
    segments=[o for o in scene['render_objects'] if o['id'].startswith('analysis-segment')]
    assert len(segments)==1 and segments[0]['id'].endswith('-3')
    assert number_label(.00000012)!=number_label(.00000024)!='0'
    pptx=export_pptx(DeckContent(title='Малые числа',slides=[slide]),template(),'a')
    with ZipFile(BytesIO(pptx)) as z:
        xml=ET.fromstring(z.read('ppt/charts/chart1.xml'))
        pts=xml.findall('.//c:val/c:numRef/c:numCache/c:pt',NS)
        assert [p.attrib['idx'] for p in pts]==['0','2','3']
        assert float(pts[0].find('c:v',NS).text)==.00000012


def test_dense_series_keeps_every_observation_with_explicit_label_stride():
    slide=analytical_slide('line')
    slide.analytics.labels=[f'Месяц {i}' for i in range(120)]
    slide.analytics.series[0].values=list(range(120))
    scene=build_scene(slide,template()['metadata'],'a',0)
    assert len([o for o in scene['render_objects'] if o['id'].startswith('analysis-point-')])==120
    chart=next(o for o in scene['objects'] if o['type']=='chart')
    assert chart['label_stride']>1


def test_audit_catches_expanded_chart_label_and_table_cell_overflow():
    slide=analytical_slide('line');t=template();content=DeckContent(title='Проверка',slides=[slide])
    scene=build_scene(slide,t['metadata'],'a',0)
    tick=next(o for o in scene['render_objects'] if o['id']=='analysis-y-tick-0')
    tick['x']=-1;tick['used_height']=999
    issues=audit_deck(content,t,'a','',scenes=[scene])
    assert {'bounds','overflow'} <= {i.code for i in issues if i.object_id==tick['id']}
    slide=analytical_slide('table');slide.analytics.rows[0][0]='Длинный текст '*100
    issues=audit_deck(DeckContent(title='Проверка',slides=[slide]),t,'a','')
    assert any(i.code=='overflow' and i.object_id.startswith('analysis-table-') for i in issues)


def test_html_escapes_user_content_and_preserves_aspect_ratio_and_font_warning():
    slide=analytical_slide('bar');slide.analytics.source='<script>alert("x")</script>'
    t=template(ratio=4/3);t['metadata']['font']='Unavailable Font'
    scene=build_scene(slide,t['metadata'],'a',0)
    assert scene['height']==960 and scene['font_warnings']
    html=export_html(DeckContent(title='X',slides=[slide]),t,'a').decode()
    assert '<script>' not in html and '&lt;script&gt;' in html


def test_scatter_retains_only_complete_pairs_in_native_chart():
    slide=analytical_slide('scatter')
    slide.analytics.x_values=[1,None,3,4]
    slide.analytics.series[0].values=[2,5,None,8]
    output=export_pptx(DeckContent(title='Пары',slides=[slide]),template(),'a')
    native=next(s.chart for s in Presentation(BytesIO(output)).slides[0].shapes if s.has_chart)
    assert list(native.series[0].values)==[2,8]


def test_blank_or_edited_evidence_is_not_clean():
    slide=analytical_slide('kpi');slide.title='Новый вывод';slide.analytics.source=''
    slide.analytics.kpis[0].period=''
    issues=audit_deck(DeckContent(title='Проверка',slides=[slide]),template(),'a','')
    assert {'analysis_source_missing','analysis_headline_changed','analysis_kpi_context'} <= {i.code for i in issues}


def test_inconsistent_numeric_claim_is_an_error_linked_to_chart():
    slide=analytical_slide('line')
    slide.analytics.calculations=[__import__('app.analytics_models',fromlist=['CalculationRecord']).CalculationRecord(
        id='growth',operation='relative_change',baseline=10,comparison=20,result=999)]
    issues=audit_deck(DeckContent(title='Проверка',slides=[slide]),template(),'a','')
    assert any(i.code.startswith('analysis_invalid_') and i.severity=='error' and i.object_id=='analysis-chart' for i in issues)


def test_pdf_bold_is_not_aliased_to_regular_and_long_heading_wraps():
    from app.services.layout import pdfmetrics
    text='Больше вложений связано с большим числом заявок'
    assert pdfmetrics.getFont('DecklyBold') is not pdfmetrics.getFont('Deckly')
    assert pdfmetrics.stringWidth(text,'DecklyBold',46)>pdfmetrics.stringWidth(text,'Deckly',46)
    slide=analytical_slide('scatter');slide.title=text
    metadata=template()['metadata'];metadata['heading_pt']=34
    scene=build_scene(slide,metadata,'a',0)
    title=next(o for o in scene['objects'] if o['id']=='title')
    assert len(title['lines'])>=2
    native=Presentation(BytesIO(export_pptx(DeckContent(title='Текст',slides=[slide]),{'metadata':metadata},'a')))
    assert any(s.has_text_frame and s.text==text for s in native.slides[0].shapes)


@pytest.mark.parametrize('variant',['a','b','c'])
def test_date_axis_keeps_final_endpoint_without_colliding_or_dropping_data(variant):
    slide=analytical_slide('line')
    slide.analytics.labels=[f'2025-{month:02}-01' for month in range(1,13)]
    slide.analytics.series[0].values=[100+2*month for month in range(12)]
    scene=build_scene(slide,template()['metadata'],variant,0)
    labels=[o for o in scene['render_objects'] if o.get('role')=='category']
    assert labels[-1]['text']=='2025-12-01'
    assert all(left['x']+left['w']+7.99<=right['x'] for left,right in zip(labels,labels[1:]))
    assert len([o for o in scene['render_objects'] if o['id'].startswith('analysis-point-')])==12
    output=export_pptx(DeckContent(title='Даты',slides=[slide]),template(),variant)
    native=Presentation(BytesIO(output))
    chart=next(s.chart for s in native.slides[0].shapes if s.has_chart)
    assert [c.label for c in chart.plots[0].categories]==slide.analytics.labels
    assert list(chart.series[0].values)==slide.analytics.series[0].values
    native_labels=[s.text for s in native.slides[0].shapes if s.name.startswith('analysis-category-')]
    assert native_labels==[o['text'] for o in labels]
    with ZipFile(BytesIO(output)) as archive:
        xml=ET.fromstring(archive.read('ppt/charts/chart1.xml'))
        assert xml.find('.//c:catAx/c:tickLblPos',NS).attrib['val']=='none'
        assert xml.find('.//c:plotArea/c:layout/c:manualLayout/c:layoutTarget',NS).attrib['val']=='inner'


def test_actual_chart_mark_contrast_is_checked_without_flagging_soft_grid():
    slide=analytical_slide('bar');t=template();t['metadata']['chart_palette']=['#FFFFFF']
    issues=audit_deck(DeckContent(title='Проверка',slides=[slide]),t,'a','')
    assert any(i.code=='analysis_mark_contrast' and i.object_id.startswith('analysis-bar-') for i in issues)
    assert not any(i.object_id.startswith('analysis-grid-') for i in issues)
