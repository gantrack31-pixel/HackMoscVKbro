"""Regression cases from QA: preserve claims while fitting editable visual labels."""
from io import BytesIO
from copy import deepcopy
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
import pytest
from pptx import Presentation
from reportlab.pdfbase.ttfonts import TTFontFile

from app.models import DeckContent, Slide
from app.services.audit import audit_deck
from app.services.export import export_pptx
from app.services.layout import build_scene
from pptx.oxml.ns import qn
from test_workflow import client, register


CASES = json.loads((Path(__file__).parent / 'fixtures/diagram_cards.json').read_text())


@pytest.mark.parametrize('case', CASES, ids=lambda case: case['name'])
@pytest.mark.parametrize('variant', 'abc')
def test_qa_diagrams_and_icon_cards_fit_without_losing_text(case, variant):
    content = DeckContent(title='QA labels', slides=case['slides'])
    for index, slide in enumerate(content.slides):
        scene = build_scene(slide, case['metadata'], variant, index)
        labels = [obj for obj in scene['objects'] if obj['id'].startswith(('node-label-', 'icon-label-'))]
        assert [obj['text'] for obj in labels] == slide.bullets
        for label in labels:
            assert label['font_size'] >= 16
            assert label['used_height'] <= label['h'] + .01, label
            card_id = label['id'].replace('node-label-', 'node-bg-').replace('icon-label-', 'icon-card-')
            card = next(obj for obj in scene['objects'] if obj['id'] == card_id)
            assert label['y'] >= card['y']
            assert label['y'] + label['h'] <= card['y'] + card['h'] + .01
        issues = audit_deck(DeckContent(title='QA', slides=[slide]), {'metadata': case['metadata']}, variant, '', scenes=[scene])
        assert not [issue for issue in issues if issue.code == 'overflow'], issues


def test_unfittable_generated_labels_do_not_publish_a_project(client):
    register(client)
    slide = Slide(title='Схема с избыточным содержанием', kind='diagram',
                  bullets=['Очень длинная подпись узла ' * 80 for _ in range(6)])
    response = client.post('/api/generate', json={
        'template_id': 'tech', 'content': DeckContent(title='Не публиковать переполнение', slides=[slide]).model_dump(),
    })
    assert response.status_code == 202, response.text
    job = client.get('/api/jobs/' + response.json()['job_id']).json()
    assert job['state'] == 'failed', job
    assert 'не помещается' in job['error']
    assert not job['project_id']
    assert client.get('/api/projects').json() == []


def test_fitted_qa_visuals_publish_successfully(client):
    register(client)
    content = DeckContent(title='Исправленные схемы и карточки', slides=CASES[0]['slides'])
    response = client.post('/api/generate', json={'template_id': 'tech', 'content': content.model_dump()})
    assert response.status_code == 202, response.text
    job = client.get('/api/jobs/' + response.json()['job_id']).json()
    assert job['state'] == 'complete', job
    project = client.get('/api/projects/' + job['project_id']).json()
    assert [s['bullets'] for s in project['content']['slides']] == [s.bullets for s in content.slides]
    for variant in 'abc':
        issues = client.get(f'/api/projects/{job["project_id"]}/audit?variant={variant}').json()['issues']
        assert not [issue for issue in issues if issue['code'] == 'overflow']


def test_pptx_embeds_both_editable_manrope_faces_once():
    template = {'metadata': {'ratio': 16 / 9}}
    deck = DeckContent(title='Portable font', slides=[Slide(title='Заголовок', body='Основной текст')])
    output = export_pptx(deck, template, 'a')
    from app.services.pptx_fonts import embed_manrope
    presentation = Presentation(BytesIO(output))
    embed_manrope(presentation)
    saved = BytesIO()
    presentation.save(saved)
    output = saved.getvalue()
    namespaces = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
                  'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with ZipFile(BytesIO(output)) as archive:
        presentation = etree.fromstring(archive.read('ppt/presentation.xml'))
        embedded = presentation.xpath('p:embeddedFontLst/p:embeddedFont[p:font/@typeface="Manrope"]', namespaces=namespaces)
        assert len(embedded) == 1
        assert presentation.get('embedTrueTypeFonts') == '1'
        assert presentation.get('saveSubsetFonts') == '0'
        relations = etree.fromstring(archive.read('ppt/_rels/presentation.xml.rels'))
        targets = {rel.get('Id'): rel.get('Target') for rel in relations}
        for face in ('regular', 'bold'):
            rid = embedded[0].find('p:' + face, namespaces).get('{' + namespaces['r'] + '}id')
            font = archive.read('ppt/' + targets[rid])
            assert int.from_bytes(font[:4], 'little') == len(font)
            assert int.from_bytes(font[8:12], 'little') == 0x00020001
            assert font[34:36] == b'LP'  # EOT magic, not a raw TTF under another extension.
            font_size = int.from_bytes(font[4:8], 'little')
            face_data = TTFontFile(BytesIO(font[-font_size:]))
            assert face_data.familyName == b'Manrope'
            assert face_data.styleName == (b'Regular' if face == 'regular' else b'Bold')
        assert len([name for name in archive.namelist() if name.startswith('ppt/fonts/')]) == 2


def test_embedding_keeps_font_relationships_shared_with_a_template_face():
    from app.services.pptx_fonts import embed_manrope
    presentation = Presentation()
    embed_manrope(presentation)
    fonts = presentation._element.find(qn('p:embeddedFontLst'))
    original = fonts[0]
    template_face = deepcopy(original)
    template_face.find(qn('p:font')).set('typeface', 'TemplateAlias')
    fonts.append(template_face)
    relationships = [node.get(qn('r:id')) for node in template_face if node.tag != qn('p:font')]
    original_parts = {rid: presentation.part.rels[rid].target_part for rid in relationships}
    embed_manrope(presentation)
    assert all(presentation.part.rels[rid].target_part is part for rid, part in original_parts.items())
