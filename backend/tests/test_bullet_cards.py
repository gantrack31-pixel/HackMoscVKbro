"""Text slides with 2–6 bullets render as cards; same words, every format, no audit regressions."""
from io import BytesIO
import pytest
from pptx import Presentation
from app.models import DeckContent, ObjectStyle, Slide, SlideDesign
from app.services.audit import audit_deck
from app.services.export import export_pptx
from app.services.layout import build_scene

BULLETS=['Единое понимание целей снижает переделки','Чёткие этапы помогают распределять ресурсы',
         'Стратегия — основа для решений при изменениях']


def template(background='#FFFFFF'):
    return {'metadata':{'ratio':16/9,'font':'Manrope','background':background,'accent':'#006AD8',
                        'colors':{'accent1':'#006AD8','accent2':'#0B8C72','dk1':'#142A45','lt1':'#FFFFFF'},'layouts':[]}}


def slide(**update):
    return Slide(**{'title':'Почему запуску нужна стратегия','body':'Без общей рамки команда тянет в разные стороны.',
                    'bullets':BULLETS,**update})


def by_id(scene):
    return {o['id']:o for o in scene['objects']}


def inside(inner,outer):
    return (inner['x']>=outer['x']-.5 and inner['y']>=outer['y']-.5
            and inner['x']+inner['w']<=outer['x']+outer['w']+.5 and inner['y']+inner['h']<=outer['y']+outer['h']+.5)


@pytest.mark.parametrize('background',['#FFFFFF','#101418'])
@pytest.mark.parametrize('variant',list('abc'))
def test_bullets_become_cards_that_fit_and_pass_audit(variant,background):
    scene=build_scene(slide(),template(background)['metadata'],variant,1)
    objects=by_id(scene)
    texts=[objects[f'bullet-{i}'] for i in range(3)]
    assert [t['text'] for t in texts]==BULLETS and objects['body']['text'].startswith('Без общей рамки')
    assert len({t['font_size'] for t in texts})==1
    for i,text in enumerate(texts):
        frame=objects.get(f'bullet-card-{i}') or {**objects[f'bullet-rule-{i}'],'h':text['y']+text['h']-objects[f'bullet-rule-{i}']['y']}
        assert inside(text,frame) and text['used_height']<=text['h']+1
        assert inside(frame,{'x':0,'y':0,'w':scene['width'],'h':scene['height']})
    content=DeckContent(title='Тест',slides=[slide()])
    problems=[i for i in audit_deck(content,template(background),variant,'')
              if i.object_id.startswith('bullet') and i.category in {'layout','template'}]
    assert problems==[]


def test_variants_draw_different_card_styles():
    meta=template()['metadata']
    ids=[set(by_id(build_scene(slide(),meta,v,1))) for v in 'abc']
    assert 'bullet-stripe-0' in ids[0] and 'bullet-badge-0' in ids[1] and 'bullet-rule-0' in ids[2]


@pytest.mark.parametrize('case',['seven_bullets','one_bullet','canvas_edit','other_kind'])
def test_plain_text_block_is_kept_when_cards_do_not_apply(case):
    update={'seven_bullets':{'bullets':[f'Пункт {i}' for i in range(7)]},'one_bullet':{'bullets':['Один пункт']},
            'canvas_edit':{'design':SlideDesign(objects={'body':ObjectStyle(x=100,y=300)})},
            'other_kind':{'kind':'steps'}}[case]
    scene=build_scene(slide(**update),template()['metadata'],'a',1)
    assert not any(o['id'].startswith('bullet-') for o in scene['objects'])


def test_overlong_bullets_fall_back_to_plain_text():
    long_bullets=['Очень длинный тезис, '+'который подробно раскрывает мысль '*8 for _ in range(6)]
    scene=build_scene(slide(bullets=long_bullets),template()['metadata'],'a',1)
    assert not any(o['id'].startswith('bullet-') for o in scene['objects'])


def test_pptx_contains_card_shapes_and_editable_bullet_text():
    output=export_pptx(DeckContent(title='Тест',slides=[slide()]),template(),'b')
    shapes=Presentation(BytesIO(output)).slides[0].shapes
    texts=[s.text_frame.text for s in shapes if s.has_text_frame]
    assert all(b in texts for b in BULLETS)
    assert sum(1 for s in shapes if s.shape_type==1)>=3  # auto shapes: cards and badges
