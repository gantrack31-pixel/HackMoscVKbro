"""Real demo generation + editable export using only a temporary database/storage."""
import io
import json
import zipfile
from copy import deepcopy
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader
from app.main import app
from app.config import settings


def register(client,email):
    response=client.post('/api/auth/register',json={'email':email,'password':'Analytics-Example-1947','first_name':'Тест','last_name':'Аналитики'})
    assert response.status_code==201,response.text
    client.headers['X-CSRF-Token']=response.json()['csrf']


def upload_choice(client,text,choice):
    response=client.post('/api/analysis/upload',files={'file':('synthetic.csv',text.encode(),'text/csv')},data={'intent':'Показать проверенный вывод; синтетические данные'})
    assert response.status_code==200,response.text
    id=response.json()['id']
    response=client.put(f'/api/analysis/{id}/selection',json=choice)
    assert response.status_code==200,response.text
    response=client.post(f'/api/analysis/{id}/plan',json={})
    assert response.status_code==200,response.text
    return id,response.json()['slides']


def test_authenticated_analysis_real_jobs_variants_exports_and_owner_validation(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'database',tmp_path/'integrated.sqlite')
    monkeypatch.setattr(settings,'storage',tmp_path/'data')
    monkeypatch.setattr(settings,'mode','demo')
    monkeypatch.setattr(settings,'app_env','development')
    monkeypatch.setattr(settings,'email_verification_required',False)
    with TestClient(app) as client:
        register(client,'analyst@example.test')
        id,slides=upload_choice(client,'Дата,Revenue\n2025-01-01,100\n2025-12-01,123',{'chart_type':'line','dimension':'Дата','measures':['Revenue']})
        _,bars=upload_choice(client,'Region,Revenue\nСевер,100\nЮг,123\nЗапад,77',{'chart_type':'bar','dimension':'Region','measures':['Revenue']})
        _,pie=upload_choice(client,'Channel,Revenue\nПрямые,50\nПартнёры,30\nОнлайн,20',{'chart_type':'pie','dimension':'Channel','measures':['Revenue'],'aggregation':'sum'})
        plain=client.post(f'/api/analysis/{id}/plan',json={'template_id':'tech'}).json()
        other_template=next(t['id'] for t in client.get('/api/templates').json() if t['id']!='tech')
        alternate=client.post(f'/api/analysis/{id}/plan',json={'template_id':other_template}).json()
        assert plain['slides']==alternate['slides'] and plain['preview']['template_id']!=alternate['preview']['template_id']
        tokens=json.dumps({'name':'Analytical brand','colors':{'accent':'#008877','background':'#FFFFFF','text':'#182A44'}}).encode()
        uploaded=client.post('/api/brands/upload',files={'file':('tokens.json',tokens,'application/json')})
        assert uploaded.status_code==201,uploaded.text
        brand_id=uploaded.json()['id']
        branded=client.post(f'/api/analysis/{id}/plan',json={'template_id':'tech','brand_style_id':brand_id}).json()
        assert branded['slides']==plain['slides']
        assert '#008877' in json.dumps(branded['preview']['scene']).lower()
        assert branded['preview']['brand_style_id']==brand_id
        assert client.post(f'/api/analysis/{id}/plan',json={'template_id':'tech','brand_style_id':'not-owned'}).status_code==404
        slides+=bars+pie
        request={'template_id':'tech','source_text':'Синтетический пример: выручка 100 → 123; причинные выводы не делаются.',
                 'content':{'title':'Проверяемая аналитика','slides':slides}}
        changed=deepcopy(request);changed['content']['slides'][0]['analytics']['series'][0]['values'][1]=999
        assert client.post('/api/generate',json=changed).status_code==409
        response=client.post('/api/generate',json=request)
        assert response.status_code==202,response.text
        jid=response.json()['job_id']
        job=client.get('/api/jobs/'+jid).json()
        assert job['state']=='complete',job
        pid=job['project_id']
        project=client.get('/api/projects/'+pid).json()
        assert len(project['content']['slides'])==3
        assert [s['analytics'] for s in project['content']['slides']]==[s['analytics'] for s in slides]
        assert '23%' in project['content']['slides'][0]['title']
        # All variants retain the same evidence; their compositions are selected by renderer.
        for variant in ('a','b','c'):
            response=client.get(f'/api/projects/{pid}/export/pptx?variant={variant}')
            assert response.status_code==200,response.text[:500]
            deck=Presentation(io.BytesIO(response.content))
            assert len(deck.slides)==3
            charts=[shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
            assert len(charts)==3
            assert list(charts[0].series[0].values)==[100.0,123.0]
            assert list(charts[1].series[0].values)==[100.0,123.0,77.0]
            assert list(charts[2].series[0].values)==[50.0,30.0,20.0]
        pdf=client.get(f'/api/projects/{pid}/export/pdf?variant=b')
        assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
        assert len(PdfReader(io.BytesIO(pdf.content)).pages)==3
        html=client.get(f'/api/projects/{pid}/export/html?variant=c')
        assert html.status_code==200 and '23%' in html.text and '2025-01-01' in html.text
        assert 'OPENROUTER_API_KEY' not in html.text and 'file:///' not in html.text
        assert client.get(f'/api/analysis/{id}/rows?limit=1').json()['next_offset']==1
        assert client.get(f'/api/analysis/{id}').json()['selection']['chart_type']=='line'
        register(client,'another-analyst@example.test')
        assert client.get(f'/api/projects/{pid}').status_code==404
        assert client.get(f'/api/projects/{pid}/export/pptx').status_code==404
        assert client.post('/api/generate',json=request).status_code==404
        assert client.get(f'/api/analysis/{id}/source').status_code==404
