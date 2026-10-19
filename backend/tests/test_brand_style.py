"""Brand inputs, precedence, ownership, bounded parsing and native export."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import pytest
from PIL import Image
from pptx import Presentation
from reportlab.pdfgen import canvas
from app import database as db
from app.config import settings
from app.models import Slide,DeckContent,SlideDesign
from app.services.brand_parser import parse_brand,BrandError,MAX_BYTES
from app.services.brand_style import parse_isolated,resolve_brand,upload,get,combine,apply_brand,remove
from app.services.export import export_pptx,export_html
from app.services.layout import build_scene
from app.services.audit import audit_deck
from test_analytical_rendering import analytical_slide,template
from test_workflow import client,register


def tokens(**extra):
    return json.dumps({'name':'Проверочный бренд','colors':{'accent':'#008877','background':'#FFFFFF','text':'#182A44'},
                       'typography':{'heading':'MissingFont','body':'Manrope','heading_pt':34,'body_pt':18},
                       'spacing':{'slide_margin':72,'grid_gap':32},'chart':{'series_palette':['#008877','#315AAA'],'grid':'none'},
                       **extra},ensure_ascii=False).encode()


def logo(ink=(0,0,0,255)):
    output=BytesIO();Image.new('RGBA',(300,60),ink).save(output,format='PNG');return output.getvalue()


def package(items):
    output=BytesIO()
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for name,value in items:z.writestr(name,value)
    return output.getvalue()


def parsed(data=None,name='brand.json'):
    result=parse_brand(data or tokens(),name);result['id']='brand-test';return result


def test_token_json_and_figma_export_extract_actual_content():
    value=parsed()
    assert value['tokens']['colors']['accent']=='#008877'
    assert value['tokens']['typography']['heading_font']=='MissingFont'
    figma={'document':{'type':'DOCUMENT','children':[{'type':'FRAME','children':[
        {'type':'RECTANGLE','name':'Brand primary','fills':[{'type':'SOLID','color':{'r':.2,'g':.4,'b':.6}}]},
        {'type':'TEXT','style':{'fontFamily':'Manrope','fontSize':32}}]}]}}
    result=parse_brand(json.dumps(figma).encode(),'figma.json')
    assert result['kind']=='figma-json' and result['tokens']['colors']['accent']=='#336699'
    assert result['tokens']['typography']['heading_font']=='Manrope'
    w3c={'color':{'primary':{'$type':'color','$value':'#123456'}},'font':{'heading':{'$type':'fontFamily','$value':'Manrope'}},
         'spacing':{'margin':{'$type':'dimension','$value':{'value':64,'unit':'px'}}}}
    result=parse_brand(json.dumps(w3c).encode(),'tokens.json')
    assert result['tokens']['colors']['accent']=='#123456' and result['tokens']['spacing']['slide_margin']==64


def test_pdf_extracts_explicit_tokens_and_keeps_unresolved_rules():
    output=BytesIO();pdf=canvas.Canvas(output)
    for i,text in enumerate(['Primary: #008877','Background: #FFFFFF','Heading font: Manrope','Slide margin: 72','Keep logos away from dense imagery.']):
        pdf.drawString(40,760-i*24,text)
    pdf.save()
    result=parse_isolated(output.getvalue(),'guide.pdf')
    assert result['tokens']['colors']['accent']=='#008877'
    assert result['tokens']['typography']['heading_font']=='Manrope'
    assert result['tokens']['spacing']['slide_margin']==72
    assert result['unresolved_rules']==['Keep logos away from dense imagery.']
    assert 'exact-pdf-text' in {e['method'] for e in result['evidence']}


def test_zip_token_precedence_and_logo_original_aspect_preserved():
    early=json.dumps({'colors':{'accent':'#FF8800'}}).encode()
    source=package([('a.json',early),('z.json',tokens()),('logo.png',logo()),('sample-slide.png',logo((10,70,120,255)))])
    result=parsed(source,'brand.zip')
    assert result['tokens']['colors']['accent']=='#008877'
    assert len(result['logos'])==1 and result['logos'][0]['width']/result['logos'][0]['height']==5
    assert result['conflicts'] and any(i['kind']=='sample-image' for i in result['inputs'])
    assert any(e['method']=='sampled-image-approximation' for e in result['evidence'])


@pytest.mark.parametrize('member',['../secret.json','/root.json','C:/brand.json','x\\secret.json','nested.zip','run.exe'])
def test_unsafe_archives_rejected(member):
    with pytest.raises(BrandError):parse_brand(package([(member,b'{}')]),'brand.zip')


def test_bad_signatures_encryption_scale_and_fig_are_actionable():
    with pytest.raises(BrandError,match='PDF'):parse_brand(b'not a pdf','guide.pdf')
    with pytest.raises(BrandError,match='JSON'):parse_brand(b'fig-kiwi','file.fig')
    with pytest.raises(BrandError):parse_brand(b'x'*(MAX_BYTES+1),'brand.json')
    with pytest.raises(BrandError):parse_brand(b'{"x":NaN}','brand.json')
    value={};current=value
    for i in range(40):current['x']={};current=current['x']
    with pytest.raises(BrandError,match='глубокая'):parse_brand(json.dumps(value).encode(),'brand.json')


def test_worker_does_not_inherit_provider_secrets_or_ambient_import_paths(monkeypatch):
    monkeypatch.setenv('OPENROUTER_API_KEY','sentinel-not-real-secret')
    monkeypatch.setenv('PYTHONPATH','untrusted-path')
    original=subprocess.run;captured={}
    def invoke(*args,**kwargs):
        captured.update(kwargs)
        return original(*args,**kwargs)
    monkeypatch.setattr(subprocess,'run',invoke)
    result=parse_isolated(tokens(),'brand.json')
    assert result['kind']=='tokens-json'
    assert 'OPENROUTER_API_KEY' not in captured['env'] and 'PYTHONPATH' not in captured['env']
    assert captured['timeout']==12
    def timeout(*args,**kwargs):raise subprocess.TimeoutExpired('worker',12)
    monkeypatch.setattr(subprocess,'run',timeout)
    with pytest.raises(BrandError,match='12 секунд'):parse_isolated(tokens(),'brand.json')


def test_brand_wins_without_mutating_template_and_explicit_template_override_works():
    original=template();snapshot=deepcopy(original)
    result=resolve_brand(original,parsed())
    assert original==snapshot and result['metadata']['accent']=='#008877'
    receipt=result['metadata']['brand_style']
    assert receipt['conflicts'] and receipt['font_mapping'][0]['effective']=='Manrope'
    assert result['metadata']['heading_font']=='Manrope'
    override=resolve_brand(original,parsed(),prefer_template=True)
    assert override['metadata']['accent']==original['metadata']['accent']
    assert override['metadata']['brand_style']['precedence']=='template'


def test_declared_margin_is_used_in_actual_analytical_geometry():
    t=resolve_brand(template(),parsed(tokens(spacing={'slide_margin':96})))
    scene=build_scene(analytical_slide('bar'),t['metadata'],'a',0)
    assert next(o for o in scene['objects'] if o['id']=='title')['x']==96


@pytest.mark.parametrize('variant',list('abc'))
def test_brand_palette_grid_logo_and_source_match_preview_and_native_export(variant):
    brand=parsed(package([('style.json',tokens()),('logo.png',logo())]),'brand.zip')
    t=resolve_brand(template(),brand)
    slide=analytical_slide('grouped');content=DeckContent(title='Бренд',slides=[slide])
    scene=build_scene(slide,t['metadata'],variant,0)
    chart=next(o for o in scene['objects'] if o['type']=='chart')
    assert chart['colors'][:2]==['#008877','#315AAA']
    assert not any(o['id'].startswith('analysis-grid') for o in scene['render_objects'])
    image=next(o for o in scene['objects'] if o['id']=='brand-logo')
    assert image['w']/image['h']==pytest.approx(5)
    issues=audit_deck(content,t,variant,'',scenes=[scene])
    assert not [i for i in issues if i.severity=='error'],[(i.code,i.object_id) for i in issues]
    output=export_pptx(content,t,variant)
    native=Presentation(BytesIO(output))
    assert sum(s.shape_type==13 for s in native.slides[0].shapes)==1
    native_chart=next(s.chart for s in native.slides[0].shapes if s.has_chart)
    assert str(native_chart.series[0].format.fill.fore_color.rgb)=='008877'
    assert not native_chart.value_axis.has_major_gridlines
    assert 'data:image/png;base64,' in export_html(content,t,variant).decode()


def test_brand_background_overrides_source_layout_background_and_logo_contrast_is_reported():
    source=template('#112233');source['metadata']['layouts']=[{'source_slide':1,'master':0,'index':0,'name':'X','static_shapes':0,
        'background_info':{'kind':'solid','color':'#112233'},'boxes':[{'type':'TITLE','x':.1,'y':.1,'w':.8,'h':.2},{'type':'BODY','x':.1,'y':.4,'w':.8,'h':.4}]}]
    t=resolve_brand(source,parsed(package([('style.json',tokens()),('logo.png',logo((255,255,255,255)))]),'brand.zip'))
    for slide in (Slide(title='Обычный слайд',body='Проверка'),analytical_slide('line')):
        slide.design=SlideDesign(composition='grid',background_role='dark')
        scene=build_scene(slide,t['metadata'],'a',0)
        assert scene['objects'][0]['fill']=='#FFFFFF'
        issues=audit_deck(DeckContent(title='Бренд',slides=[slide]),t,'a','')
        assert any(i.code=='brand_logo_contrast' for i in issues)


def test_owner_isolation_combine_and_original_preservation(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'database',tmp_path/'brands.sqlite3');db.initialize()
    first=upload('owner-a',tokens(),'brand.json')
    second=upload('owner-a',json.dumps({'colors':{'accent':'#5522BB'}}).encode(),'second.json')
    assert get('owner-a',first['id'])['fingerprint']==sha256(tokens()).hexdigest()
    with pytest.raises(KeyError):get('owner-b',first['id'])
    with pytest.raises(KeyError):apply_brand(template(),first['id'],'owner-b')
    joined=combine('owner-a',[first['id'],second['id']],'Combined')
    assert apply_brand(template(),joined['id'],'owner-a')['metadata']['accent']=='#5522BB'
    with pytest.raises(KeyError):remove('owner-b',first['id'])
    remove('owner-a',first['id'])
    assert get('owner-a',joined['id'])['tokens']['colors']['accent']=='#5522BB'


def test_brand_routes_preserve_original_offer_effective_preview_and_enforce_owner(client):
    assert client.get('/api/brands').status_code==401
    register(client,'brand-a@example.test')
    result=client.post('/api/brands/upload',files={'file':('brand.json',tokens(),'application/json')})
    assert result.status_code==201,result.text
    brand_id=result.json()['id']
    preview=client.post(f'/api/brands/{brand_id}/preview',json={'template_id':'tech','prefer_template':False})
    assert preview.status_code==200,preview.text
    assert preview.json()['effective']['accent']=='#008877'
    assert preview.json()['brand_style']['font_mapping'][0]['effective']=='Manrope'
    assert client.get(f'/api/brands/{brand_id}/source').content==tokens()
    client.post('/api/auth/logout')
    register(client,'brand-b@example.test')
    assert client.get(f'/api/brands/{brand_id}').status_code==404
    assert client.get(f'/api/brands/{brand_id}/source').status_code==404
    assert client.post(f'/api/brands/{brand_id}/preview',json={'template_id':'tech'}).status_code==404
    assert client.delete(f'/api/brands/{brand_id}').status_code==404
    assert client.post('/api/brands/upload',files={'file':('file.fig',b'fig-kiwi','application/octet-stream')}).status_code==422


def test_api_data_brand_generation_snapshot_draft_preview_audit_export_and_revision(client):
    register(client,'brand-flow@example.test')
    branded=client.post('/api/brands/upload',files={'file':('brand.json',tokens(),'application/json')}).json()
    upload_data=client.post('/api/analysis/upload',files={'file':('monthly.csv','Месяц,Выручка млн руб.\n2026-01-01,10\n2026-02-01,15\n2026-03-01,20\n'.encode(),'text/csv')})
    assert upload_data.status_code==200,upload_data.text
    dataset_id=upload_data.json()['id']
    selection={'table_id':'t1','chart_type':'line','dimension':'Месяц','measures':['Выручка млн руб.'],'aggregation':'none'}
    selected=client.put(f'/api/analysis/{dataset_id}/selection',json=selection)
    assert selected.status_code==200,selected.text
    planned=client.post(f'/api/analysis/{dataset_id}/plan',json={})
    assert planned.status_code==200,planned.text
    content={'title':'Проверка брендбука','slides':planned.json()['slides'],'brand_style_id':branded['id']}
    generated=client.post('/api/generate',json={'template_id':'tech','content':content,'source_text':'Месячная синтетическая выручка'})
    assert generated.status_code==202,generated.text
    job=client.get('/api/jobs/'+generated.json()['job_id']).json()
    assert job['state']=='complete',job
    pid=job['project_id'];project=client.get('/api/projects/'+pid).json()
    assert project['template']['metadata']['accent']=='#008877'
    assert project['provenance']['effective_style']['brand_style_id']==branded['id']
    assert all(project['variants'][v][0]['objects'][0]['fill']=='#FFFFFF' for v in 'abc')
    original_version=project['updated_at']
    output=client.get(f'/api/projects/{pid}/export/pptx',params={'variant':'b','version':original_version})
    assert output.status_code==200
    native=Presentation(BytesIO(output.content))
    chart=next(shape.chart for shape in native.slides[0].shapes if shape.has_chart)
    assert list(chart.series[0].values)==[10,15,20]
    assert str(chart.series[0].format.fill.fore_color.rgb)=='008877'
    assert client.get(f'/api/projects/{pid}/export/pdf').content.startswith(b'%PDF')
    assert b'<svg' in client.get(f'/api/projects/{pid}/export/html').content
    audit=client.get(f'/api/projects/{pid}/audit').json()
    assert not [i for i in audit['issues'] if i['severity']=='error'],audit
    other=client.post('/api/brands/upload',files={'file':('new.json',json.dumps({'colors':{'accent':'#5522BB'}}).encode(),'application/json')}).json()
    draft=deepcopy(project['content']);draft['brand_style_id']=other['id']
    preview=client.post(f'/api/projects/{pid}/preview',json={'content':draft,'variant':'a','base_updated_at':original_version})
    assert preview.status_code==200,preview.text
    assert next(o for o in preview.json()['scenes'][0]['objects'] if o['type']=='chart')['colors'][0]=='#5522BB'
    saved=client.put('/api/projects/'+pid,json={'content':draft,'variant':'a','base_updated_at':original_version})
    assert saved.status_code==200,saved.text
    assert saved.json()['template']['metadata']['accent']=='#5522BB'
    assert client.get(f'/api/projects/{pid}/export/pptx',params={'version':original_version}).status_code==409
    # Source removal outside the API simulates missing original storage. The
    # generation snapshot must keep the already generated project readable.
    # Use a second untouched generated project so the snapshot matches its style.
    repeated=client.post('/api/generate',json={'template_id':'tech','content':content,'source_text':'Синтетический набор'})
    repeated_job=client.get('/api/jobs/'+repeated.json()['job_id']).json()
    assert repeated_job['state']=='complete',repeated_job
    with db.connection() as conn:conn.execute('DELETE FROM brand_styles WHERE id=?',(branded['id'],))
    frozen=client.get('/api/projects/'+repeated_job['project_id'])
    assert frozen.status_code==200 and frozen.json()['template']['metadata']['accent']=='#008877'
    assert client.get('/api/projects/'+repeated_job['project_id']+'/export/pptx').status_code==200
