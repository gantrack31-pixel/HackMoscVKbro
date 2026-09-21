"""Проверки реальных границ: аккаунты, изоляция, экспорт, OAuth, контракт LLM."""
import asyncio
from io import BytesIO
import json
from pathlib import Path
from urllib.parse import urlparse,parse_qs
from zipfile import ZipFile
import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from app.main import app
from app.config import settings,BASE
from app import database as db
from app.models import DeckContent,Slide,ChartData
from app.services.templates import analyze_template
from app.services.export import export_pptx,export_pdf,export_html,geometry_nodes
from app.services.audit import audit_deck,source_status
from app.services.llm import complete_json,LLMError,make_outline
from app.models import OutlineRequest

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'database',tmp_path/'test.sqlite3')
    monkeypatch.setattr(settings,'storage',tmp_path/'data')
    monkeypatch.setattr(settings,'mode','demo')
    with TestClient(app) as client:yield client

def register(client,email='anna@example.test'):
    response=client.post('/api/auth/register',json={'email':email,'password':'Example-Password-1947','first_name':'Анна','last_name':'Тестовая'})
    assert response.status_code==201,response.text
    client.headers['X-CSRF-Token']=response.json()['csrf']
    return response.json()['user']

def make_project(client):
    content=DeckContent(title='Проверка',slides=[Slide(title='Проверенный заголовок',body='Текст',source_quote='Текст')]).model_dump()
    job=client.post('/api/generate',json={'template_id':'tech','content':content,'source_text':'Текст'})
    assert job.status_code==202,job.text
    status=client.get('/api/jobs/'+job.json()['job_id']).json()
    assert status['state']=='complete',status
    return client.get('/api/projects/'+status['project_id']).json(),job.json()['job_id']

def test_registration_password_session_csrf_and_logout(client):
    assert client.get('/api/projects').status_code==401
    user=register(client)
    cookie=client.cookies.get('deckly_session')
    with db.connection() as connection:
        row=connection.execute('SELECT * FROM users').fetchone()
        assert row['password_hash'].startswith('scrypt$') and 'Example-Password' not in row['password_hash']
        assert cookie not in connection.execute('SELECT token_hash FROM sessions').fetchone()[0]
    assert client.get('/api/auth/me').json()['user']['id']==user['id']
    assert client.post('/api/auth/logout',headers={'X-CSRF-Token':'wrong'}).status_code==403
    assert client.post('/api/auth/logout',headers={'Origin':'https://evil.example'}).status_code==403
    assert client.post('/api/auth/logout').status_code==204
    assert client.get('/api/auth/me').json()['user'] is None
    bad=client.post('/api/auth/login',json={'email':user['email'],'password':'incorrect'})
    assert bad.status_code==401
    good=client.post('/api/auth/login',json={'email':user['email'].upper(),'password':'Example-Password-1947'})
    assert good.status_code==200 and good.json()['user']['id']==user['id']
    assert 'httponly' in good.headers['set-cookie'].lower()
    assert 'password_hash' not in good.text

def test_account_isolation_jobs_and_exports(client):
    register(client);project,jid=make_project(client);pid=project['id']
    for extension in ['pptx','pdf','html']:
        response=client.get(f'/api/projects/{pid}/export/{extension}')
        assert response.status_code==200 and len(response.content)>500
    first_cookies=dict(client.cookies);first_csrf=client.headers['X-CSRF-Token']
    register(client,'second@example.test')
    assert client.get('/api/projects').json()==[]
    for path in [f'/api/projects/{pid}',f'/api/projects/{pid}/audit',f'/api/projects/{pid}/export/pptx',f'/api/jobs/{jid}']:
        assert client.get(path).status_code==404
    assert client.put(f'/api/projects/{pid}',json={'content':project['content'],'variant':'a'}).status_code==404
    assert client.post(f'/api/projects/{pid}/undo').status_code==404
    assert client.delete(f'/api/projects/{pid}').status_code==404
    assert client.post(f'/api/projects/{pid}/fix',json={'issue_ids':['x'],'variant':'a'}).status_code==404

def test_favorites_persist_and_are_private(client):
    register(client)
    assert client.put('/api/favorites/tech').status_code==204
    assert client.get('/api/favorites').json()==['tech']
    register(client,'favorites-other@example.test')
    assert client.get('/api/favorites').json()==[]
    assert client.put('/api/favorites/nonexistent').status_code==404

def test_outline_audit_selected_fix_undo(client):
    register(client)
    template=db.template_get('tech');template['metadata']['accent']='#BBDDEE'
    db.template_save('tech',template['name'],template['path'],template['metadata'])
    outline=client.post('/api/outline',json={'template_id':'tech','prompt':'Тема\nПодтверждённый факт.','count':3,'mode':'text'})
    assert outline.status_code==200 and len(outline.json()['content']['slides'])==3
    project,_=make_project(client);pid=project['id']
    issues=client.get(f'/api/projects/{pid}/audit').json()['issues']
    fixable=[i['id'] for i in issues if i['fixable']]
    assert fixable
    assert client.post(f'/api/projects/{pid}/fix',json={'issue_ids':['forged'],'variant':'a'}).status_code==422
    fixed=client.post(f'/api/projects/{pid}/fix',json={'issue_ids':[fixable[0]],'variant':'a'})
    assert fixed.status_code==200 and fixed.json()['can_undo']
    after=client.get(f'/api/projects/{pid}/audit').json()['issues']
    assert fixable[0] not in [i['id'] for i in after]
    assert client.post(f'/api/projects/{pid}/undo').json()['content']==project['content']
    assert client.post(f'/api/projects/{pid}/audit/content').status_code==409

def test_unseen_template_ratio_private_and_bad_file(client):
    register(client)
    assert client.post('/api/templates',files={'file':('bad.pptx',b'not zip')}).status_code==422
    prs=Presentation();prs.slide_width=Inches(10);prs.slide_height=Inches(7.5);prs.slides.add_slide(prs.slide_layouts[0]);stream=BytesIO();prs.save(stream)
    response=client.post('/api/templates',files={'file':('four-three.pptx',stream.getvalue())})
    assert response.status_code==201,response.text
    template=response.json();assert template['metadata']['ratio']==pytest.approx(4/3,abs=1e-5)
    register(client,'other@example.test')
    assert template['id'] not in [t['id'] for t in client.get('/api/templates').json()]
    assert client.post('/api/outline',json={'template_id':template['id'],'prompt':'Текст','count':3}).status_code==404

@pytest.mark.parametrize('tid',['tech','workspace','education'])
def test_original_templates_three_editable_variants(tid):
    path=BASE/'data/templates'/f'{tid}.pptx'
    if not path.exists():pytest.skip('Исходные шаблоны не установлены')
    metadata=analyze_template(path);template={'path':str(path),'metadata':metadata}
    content=DeckContent(title='Контрольный пример',slides=[Slide(title='Содержание остаётся общим',body='Редактируемый текст.'),
        Slide(title='Динамика',kind='chart',chart=ChartData(labels=['Период А','Период Б'],values=[-12,18],unit='%')),
        Slide(title='План',kind='table',table=[['Этап','Статус'],['Дизайн','Готово']])])
    original=Presentation(path)
    for variant in ['a','b','c']:
        pptx=Presentation(BytesIO(export_pptx(content,template,variant)))
        assert len(pptx.slides)==3
        assert len(pptx.slide_masters)==len(original.slide_masters)
        assert any(s.has_chart for s in pptx.slides[1].shapes)
        assert any(s.has_table for s in pptx.slides[2].shapes)
        assert any(s.has_text_frame and 'Содержание остаётся общим' in s.text for s in pptx.slides[0].shapes)
        chart=next(s.chart for s in pptx.slides[1].shapes if s.has_chart)
        assert tuple(chart.series[0].values)==(-12.,18.)
    pdf=PdfReader(BytesIO(export_pdf(content,template,'a')))
    assert len(pdf.pages)==3 and 'Содержание' in pdf.pages[0].extract_text()
    assert 'Редактируемый' in export_html(content,template,'a').decode()

def test_source_and_negative_chart():
    assert source_status('точная\nцитата','Это точная цитата из текста')=='matched'
    assert source_status('чужие цифры','Исходник')=='not_found'
    assert source_status('','Исходник')=='missing'
    nodes=geometry_nodes({'type':'chart','x':50,'y':50,'w':800,'h':300,'labels':['А','Б'],'values':[-10,10],'fill':'#0077FF','unit':'%'})
    bars=[n for n in nodes if n['type']=='rect' and n['fill']=='#0077FF']
    assert len(bars)==2
    assert bars[0]['x']<bars[1]['x'] and bars[0]['x']+bars[0]['w']==pytest.approx(bars[1]['x'])
    with pytest.raises(ValueError):ChartData(labels=['А'],values=[float('nan')])

def test_llm_mock_success_and_errors(monkeypatch):
    monkeypatch.setattr(settings,'api_key','fake-secret-for-test')
    monkeypatch.setattr(settings,'base_url','https://model.example/v1')
    async def run():
        def success(request):
            assert request.headers['authorization']=='Bearer fake-secret-for-test'
            assert json.loads(request.content)['response_format']['type']=='json_object'
            return httpx.Response(200,json={'choices':[{'message':{'content':'```json\n{"ok":true}\n```'}}]})
        async with httpx.AsyncClient(transport=httpx.MockTransport(success)) as client:
            assert await complete_json('Инструкция',{'text':'Материал'},client)=={'ok':True}
        for status,data in [(401,{'error':'fake-secret-for-test'}),(200,{'choices':[{'message':{'content':'не JSON'}}]})]:
            async with httpx.AsyncClient(transport=httpx.MockTransport(lambda request:httpx.Response(status,json=data))) as client:
                with pytest.raises(LLMError) as exc:await complete_json('x',{},client)
                assert 'fake-secret-for-test' not in str(exc.value)
    asyncio.run(run())

def test_yandex_pkce_state_and_account(client,monkeypatch):
    monkeypatch.setattr(settings,'yandex_id','test-client')
    monkeypatch.setattr(settings,'yandex_secret','test-secret')
    start=client.get('/api/auth/yandex/start',follow_redirects=False)
    params=parse_qs(urlparse(start.headers['location']).query)
    assert params['code_challenge_method']==['S256'] and len(params['code_challenge'][0])==43
    state=params['state'][0]
    original=httpx.AsyncClient
    def handler(request):
        if request.url.host=='oauth.yandex.ru':
            body=parse_qs(request.content.decode());assert body['client_secret']==['test-secret'];assert body['code_verifier']
            return httpx.Response(200,json={'access_token':'provider-token'})
        assert request.headers['authorization']=='OAuth provider-token'
        return httpx.Response(200,json={'id':'yandex-user-001','default_email':'oauth@example.test','first_name':'Иван','last_name':'Тестовый'})
    monkeypatch.setattr('app.auth.httpx.AsyncClient',lambda **kwargs:original(transport=httpx.MockTransport(handler),**kwargs))
    response=client.get('/api/auth/yandex/callback',params={'state':state,'code':'one-time-code'},follow_redirects=False)
    assert response.status_code==302 and response.headers['location'].endswith('#home')
    assert client.get('/api/auth/me').json()['user']['email']=='oauth@example.test'
    repeated=client.get('/api/auth/yandex/callback',params={'state':state,'code':'one-time-code'},follow_redirects=False)
    assert 'auth-error=state' in repeated.headers['location']
    client.cookies.clear()
    start=client.get('/api/auth/yandex/start',follow_redirects=False);state=parse_qs(urlparse(start.headers['location']).query)['state'][0]
    client.cookies.clear()
    rejected=client.get('/api/auth/yandex/callback',params={'state':state,'code':'x'},follow_redirects=False)
    assert 'auth-error=state' in rejected.headers['location']
