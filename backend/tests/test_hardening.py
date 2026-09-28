import asyncio
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from app.models import DeckContent, Slide
from app.services import llm
from app.services.sources import locate_quote, source_chunks, slide_binding, validate_bindings
from app.services.layout import build_scene
from app.services.templates import analyze_template
from app.services.export import export_pptx
from test_assistant import client


@pytest.mark.parametrize('response', [
    '{"ok":true}', '```JSON\n{"ok":true}\n``', 'Ответ:\n{"ok":true}\n```',
    '<think>internal {bad}</think>\n```json\n{"ok":true}\n```',
    '````json\n<reasoning>internal</reasoning>{"ok":true}'])
def test_wrapped_json(response):
    assert llm.extract_json(response) == {'ok':True}


@pytest.mark.parametrize('response', [
    '<think>not closed {"ok":true}', '{"a":1}{"b":2}', '{"a":1,"a":2}',
    '[{"a":1}]', '{"a":NaN}', '{"a":', '{"a":1,}', '{bad} {"a":1}'])
def test_ambiguous_json_rejected(response):
    with pytest.raises(ValueError): llm.extract_json(response)


def test_json_strings_are_not_stripped():
    assert llm.extract_json('{"text":"<think>literal</think> ``` { }"}')['text'].startswith('<think>')


def test_redesign_retries_format_and_lossless_shorthand(client, monkeypatch):
    calls=[]
    async def response(system,payload):
        calls.append(payload.copy())
        if len(calls)==1: raise llm.LLMFormatError('truncated JSON')
        return {v:[{'slide':i,'composition':c,'density':'balanced','layout_shift':0}
                   for i in range(2)] for v,c in zip('abc',('editorial','split','grid'))}
    monkeypatch.setattr(llm,'complete_json',response)
    original=client.get('/api/projects/ai-test').json()['content']
    request=client.post('/api/projects/ai-test/regenerate',json={'instruction':'Обнови оформление'})
    job=client.get('/api/jobs/'+request.json()['job_id']).json()
    assert job['state']=='complete', job
    assert len(calls)==2 and 'format_correction' in calls[1]
    after=client.get('/api/projects/'+job['project_id']).json()['content']
    for before, current in zip(original['slides'],after['slides']):
        assert before['body']==current['body'] and before['source_quote']==current['source_quote']


def test_source_offsets_unicode_whitespace_and_cross_chunk():
    source='🙂 '+('я'*1590)+' Начало\n\t30 участников.'
    refs=locate_quote('Начало 30 участников.',source)
    assert len(refs)==2
    for ref in refs: assert source[ref['start']:ref['end']]==ref['quote']
    assert validate_bindings(refs,source)==refs
    assert validate_bindings([{**refs[0],'start':-1},{'chunk_id':{}}],source)==[]
    bound=slide_binding(Slide(title='Пилот',body='Начало 30 участников.'),source)
    assert bound['status']=='matched' and not bound['semantic_verified']
    assert slide_binding(Slide(title='Итоги',source_quote='Выдуманный факт'),source)['status']=='not_found'


def test_semantic_audit_safe_fallback_and_no_binary_data(client,monkeypatch):
    async def bad(system,payload):
        assert 'image_data' not in payload['presentation']['slides'][0]
        assert payload['source_chunks'][0]['id']==source_chunks('Пилот охватил 30 участников.')[0]['id']
        return {'issues':[{'slide':0,'title':'Проверить','detail':'Число спорно',
                           'source_bindings':[{'chunk_id':{},'start':0,'end':1,'quote':'x'}]},None]}
    monkeypatch.setattr(llm,'complete_json',bad)
    result=client.post('/api/projects/ai-test/audit/content')
    assert result.status_code==200
    assert result.json()['issues'][0]['grounding']=='unverified'
    async def unavailable(*args): raise llm.LLMError('provider unavailable')
    monkeypatch.setattr(llm,'complete_json',unavailable)
    assert client.post('/api/projects/ai-test/audit/content').json()['issues'][0]['grounding']=='unavailable'


def test_undo_redo_branch_and_noop(client):
    url='/api/projects/ai-test'
    initial=client.get(url).json()
    assert not initial['can_undo'] and not initial['can_redo']
    assert client.post(url+'/redo').status_code==409
    content=initial['content']; content['title']='Правка'
    saved=client.put(url,json={'content':content,'variant':'b'}).json()
    restored=client.post(url+'/undo').json()
    assert restored['can_redo'] and not restored['can_undo']
    client.put(url,json={'content':restored['content'],'variant':restored['variant']})
    assert client.get(url).json()['can_redo']  # no-op must not erase future
    assert client.post(url+'/redo').json()['content']==saved['content']
    client.post(url+'/undo')
    content['title']='Другая ветка'
    assert not client.put(url,json={'content':content,'variant':'a'}).json()['can_redo']
    assert client.post(url+'/redo').status_code==409
    assert client.post(url+'/undo',headers={'X-CSRF-Token':'wrong'}).status_code==403
    assert 'future' not in client.get(url).json()


def test_typography_autofit_without_design_or_fixed_flags():
    meta={'ratio':16/9,'heading_pt':32,'body_pt':24,'heading_font':'Cambria','body_font':'Calibri'}
    short=build_scene(Slide(title='Вывод',body='Коротко.'),meta,'a',0)
    long=build_scene(Slide(title='Вывод',body='Много слов. '*90),meta,'a',0)
    objects={o['id']:o for o in long['objects']}
    short_body=next(o for o in short['objects'] if o['id']=='body')
    assert objects['body']['font_size']<short_body['font_size']
    assert objects['title']['font_size']>objects['body']['font_size']>objects['footer']['font_size']
    assert objects['title']['font_family']=='Cambria' and objects['body']['font_family']=='Calibri'
    assert objects['body']['used_height']==len(objects['body']['lines'])*objects['body']['font_size']*objects['body']['line_spacing']


def test_ingestion_inventories_every_part_and_late_image(client,tmp_path):
    prs=Presentation()
    for i in range(13): prs.slides.add_slide(prs.slide_layouts[6])
    im=BytesIO(); Image.new('RGB',(1200,700),'purple').save(im,format='PNG')
    original_image=im.getvalue()
    prs.slides[-1].shapes.add_picture(BytesIO(original_image),Inches(1),Inches(1),Inches(4))
    prs.slide_masters[0].background.fill.gradient()
    text=prs.slides[-1].shapes.add_textbox(0,0,Inches(5),Inches(1)).text_frame.paragraphs[0]
    text.text='Текст'; text.font.name='Custom Corporate Font'
    out=BytesIO(); prs.save(out); original=out.getvalue()
    upload=client.post('/api/templates',files={'file':('unseen.pptx',original)})
    assert upload.status_code==201,upload.text
    t=upload.json(); meta=t['metadata']
    assert 'Custom Corporate Font' in meta['fonts']
    with ZipFile(BytesIO(original)) as archive:
        assert {p['name']:p['sha256'] for p in meta['package']['parts']}=={n:sha256(archive.read(n)).hexdigest() for n in archive.namelist()}
    assert any('gradFill' in p['counts'] for p in meta['package']['primitives'])
    assert len(meta['assets'])==1 and meta['assets'][0]['data']
    assert client.get(f"/api/templates/{t['id']}/original").content==original
    asset=meta['assets'][0]
    path=tmp_path/'source.pptx'; path.write_bytes(original)
    content=DeckContent(title='Пример',slides=[Slide(title='Фото',kind='image',template_asset_id=asset['id'])])
    exported=export_pptx(content,{'path':str(path),'metadata':meta},'a')
    with ZipFile(BytesIO(exported)) as archive:
        assert original_image in [archive.read(n) for n in archive.namelist() if n.startswith('ppt/media/')]
        assert b'gradFill' in archive.read('ppt/slideMasters/slideMaster1.xml')
