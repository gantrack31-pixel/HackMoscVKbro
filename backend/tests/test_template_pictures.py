"""Template pictures: telling them from decor, offering them to the model, keeping them on slides."""
import base64
import random
from io import BytesIO
from PIL import Image, ImageDraw
from pptx import Presentation
from app.models import DeckContent, OutlineRequest, Slide
from app.services.export import export_pptx
from app.services.layout import build_scene
from app.services.llm import ensure_template_pictures
from app.services.template_context import asset_profile, model_template, template_images


def data_uri(image):
    buffer = BytesIO(); image.save(buffer, 'PNG')
    return 'data:image/png;base64,' + base64.b64encode(buffer.getvalue()).decode()


def picture(width=512, height=288, seed=1):
    """A photo-like image: many colours, no dominant fill."""
    rng = random.Random(seed)
    image = Image.new('RGB', (width, height))
    image.putdata([(rng.randrange(256), rng.randrange(256), rng.randrange(256)) for _ in range(width * height)])
    return image


def paw():
    """Decor typical of templates: one flat fill on transparency."""
    image = Image.new('RGBA', (512, 512), (0, 0, 0, 0))
    ImageDraw.Draw(image).ellipse((100, 100, 400, 400), fill=(150, 100, 60, 255))
    return image


def line_icon():
    image = Image.new('RGBA', (512, 512), (0, 0, 0, 0))
    ImageDraw.Draw(image).rounded_rectangle((40, 40, 470, 470), 60, outline=(255, 255, 255, 255), width=30)
    return image


def template():
    assets = [{'id': 'decor', 'name': 'image3.png', 'description': '', 'slide': -1, 'data': data_uri(paw())},
              {'id': 'photo1', 'name': 'image2.jpg', 'description': '', 'slide': -1, 'data': data_uri(picture(seed=1))},
              {'id': 'icon', 'name': 'image47.png', 'description': '', 'slide': -1, 'data': data_uri(line_icon())},
              {'id': 'photo2', 'name': 'image44.jpg', 'description': '', 'slide': 2, 'data': data_uri(picture(300, 450, seed=2))},
              {'id': 'tiny', 'name': 'image1.png', 'description': '', 'slide': -1, 'data': data_uri(picture(79, 79, seed=3))},
              {'id': 'photo3', 'name': 'image38.jpg', 'description': '', 'slide': -1, 'data': data_uri(picture(seed=4))}]
    return {'id': 'pics', 'name': 'Приют', 'metadata': {'ratio': 16 / 9, 'font': 'Manrope', 'background': '#FFFFFF',
            'accent': '#006AD8', 'colors': {'accent1': '#006AD8', 'dk1': '#142A45', 'lt1': '#FFFFFF'}, 'layouts': [], 'assets': assets}}


def deck(count=9):
    slides = [Slide(title='Приют для животных', kind='title', body='Как мы помогаем')]
    slides += [Slide(title=f'Направление {chr(1040 + i)}', kind='text', body='Описание направления.',
                     bullets=['Первый тезис', 'Второй тезис']) for i in range(count - 1)]
    return DeckContent(title='Приют', slides=slides)


def test_profile_separates_pictures_from_decor_icons_and_thumbnails():
    roles = {a['id']: asset_profile(a)['role'] for a in template()['metadata']['assets']}
    assert roles == {'decor': 'graphic', 'photo1': 'illustration', 'icon': 'graphic',
                     'photo2': 'illustration', 'tiny': 'graphic', 'photo3': 'illustration'}
    assert asset_profile(template()['metadata']['assets'][3])['orientation'] == 'portrait'


def test_model_sees_only_pictures_with_orientation():
    offered = model_template(template())['assets']
    assert [a['id'] for a in offered] == ['photo1', 'photo2', 'photo3']
    assert all(a['role'] == 'illustration' and 'data' not in a for a in offered)
    assert len(template_images(template())) == 3


def request(prompt, count=9):
    return OutlineRequest(template_id='pics', prompt=prompt, count=count)


def test_explicit_request_for_photos_adds_template_pictures_without_changing_text():
    content = deck()
    before = [(s.title, s.body, s.bullets) for s in content.slides]
    ensure_template_pictures(content, request('Расскажи о приюте, добавь фото из шаблона'), template())
    pictured = [s for s in content.slides if s.kind == 'image']
    assert len(pictured) == 3
    assert {s.template_asset_id for s in pictured} == {'photo1', 'photo2', 'photo3'}
    assert content.slides[0].kind == 'title'
    assert [(s.title, s.body, s.bullets) for s in content.slides] == before


def test_model_choice_is_kept_without_request_or_when_already_illustrated():
    content = deck()
    ensure_template_pictures(content, request('Расскажи о приюте'), template())
    assert not any(s.kind == 'image' for s in content.slides)
    content = deck()
    for index, asset_id in [(2, 'photo1'), (4, 'photo2'), (6, 'photo3')]:
        content.slides[index].kind = 'image'; content.slides[index].template_asset_id = asset_id
    ensure_template_pictures(content, request('Нужны изображения'), template())
    assert sum(s.kind == 'image' for s in content.slides) == 3


def test_short_deck_gets_at_most_a_third_illustrated():
    content = deck(3)
    ensure_template_pictures(content, request('Добавь картинки', 3), template())
    assert sum(s.kind == 'image' for s in content.slides) == 1


def test_template_without_pictures_is_left_alone():
    bare = template(); bare['metadata']['assets'] = bare['metadata']['assets'][:1]
    content = deck()
    ensure_template_pictures(content, request('Добавь фото'), bare)
    assert not any(s.kind == 'image' for s in content.slides)


def test_image_slide_keeps_bullets_and_picture_proportions():
    t = template()
    slide = Slide(title='Истории питомцев', kind='image', body='Кого мы приютили.',
                  bullets=['Собаки из отлова', 'Кошки от волонтёров'], template_asset_id='photo2')
    scene = build_scene(slide, t['metadata'], 'a', 1)
    body = next(o for o in scene['objects'] if o['id'] == 'body')
    assert 'Собаки из отлова' in body['text'] and 'Кошки от волонтёров' in body['text']
    image = next(o for o in scene['objects'] if o['id'] == 'illustration')
    assert abs(image['w'] / image['h'] - 300 / 450) < .01
    assert image['x'] >= body['x'] + body['w'] and image['x'] + image['w'] <= scene['width'] + .01
    assert image['y'] >= 0 and image['y'] + image['h'] <= scene['height'] + .01
    content = DeckContent(title='Приют', slides=[slide])
    pptx = Presentation(BytesIO(export_pptx(content, t, 'a')))
    pictures = [s for s in pptx.slides[0].shapes if s.shape_type == 13]
    assert len(pictures) == 1 and abs(pictures[0].width / pictures[0].height - 300 / 450) < .02
