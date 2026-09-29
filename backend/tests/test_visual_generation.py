import asyncio
import json
from io import BytesIO
import httpx
import pytest
from pptx import Presentation
from app.config import settings
from app.models import DeckContent, Slide, SlideDesign, ThreeDesignPlans
from app.services import llm
from app.services.visual_grounding import validate_visual
from app.services.export import export_pptx
from app.services.palette import design_metadata
from app.services.layout import build_scene
from test_workflow import client

def chart(**kwargs):
    return Slide(title='Результаты',kind='chart',chart={
        'labels':['Команда А','Команда Б'],'values':[10,20],'unit':'%',**kwargs})

def test_schema_preserves_title_fields_and_removes_only_legacy_geometry():
    schema=llm.compact_schema(DeckContent)
    assert 'title' in schema['properties'] and 'title' in schema['$defs']['Slide']['properties']
    assert schema['title']=='DeckContent'
    assert 'smartart' not in schema['$defs']['SlideDesign']['properties']
    assert 'SmartArtNode' not in schema['$defs']
    for model in (DeckContent,ThreeDesignPlans):
        schema=llm.compact_schema(model)
        def check(value):
            if isinstance(value,dict):
                if '$ref' in value: assert value['$ref'].split('/')[-1] in schema['$defs']
                for child in value.values(): check(child)
            elif isinstance(value,list):
                for child in value: check(child)
        check(schema)

def test_gemma_folds_instructions_into_user_turn(monkeypatch):
    monkeypatch.setattr(settings,'model','google/gemma-3-27b-it')
    monkeypatch.setattr(settings,'base_url','https://provider.example/v1')
    monkeypatch.setattr(settings,'api_key','test-key')
    def respond(request):
        body=json.loads(request.content)
        assert [m['role'] for m in body['messages']]==['user']
        assert 'Never invent' in body['messages'][0]['content']
        return httpx.Response(200,json={'choices':[{'message':{'content':'```json\n{"ok":true}\n```'}}]})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as transport:
            assert await llm.complete_json('Never invent',{'schema':llm.compact_schema(DeckContent)},client=transport)=={'ok':True}
    asyncio.run(run())

def test_chart_grounding_pairs_values_and_units():
    source='Команда А: 10 %; Команда Б: 20 %'
    validate_visual(chart(),source)
    validate_visual(chart(),'Команда А: 10%; Команда Б: 20%')
    validate_visual(chart(),'Команда;Доля %\nКоманда А;10\nКоманда Б;20')
    for data in (chart(values=[11,20]),chart(values=[20,10]),chart(unit='руб'),chart(unit='')):
        with pytest.raises(ValueError): validate_visual(data,source)
    with pytest.raises(ValueError): validate_visual(chart(values=[20,10]),'Команда А: 10 %, Команда Б: 20 %')

def test_tables_require_sourced_cells_and_bounded_size():
    table=Slide(title='Сравнение',kind='table',table=[['Группа','Число'],['А','19']])
    validate_visual(table,'А: 19')
    with pytest.raises(ValueError): validate_visual(table.model_copy(update={'table':[['Группа','Число'],['А','9']]}),'А: 19')
    with pytest.raises(ValueError): validate_visual(table.model_copy(update={'table':[['Группа','Число']]+[['А','19']]*7}),'А: 19')
    # Manual/legacy oversized data can still be read and preserved unchanged.
    legacy=table.model_copy(update={'table':[['Группа','Число']]+[['А','19']]*7})
    validate_visual(legacy,'',previous=legacy)


def test_semantic_matrix_capacity_is_validated():
    slide=Slide(title='Матрица',kind='diagram',bullets=['А','Б','В','Г','Д'],design=SlideDesign(composition='grid',diagram_style='matrix'))
    with pytest.raises(ValueError,match='Матрица'): validate_visual(slide,'А Б В Г Д')

def test_visual_pptx_objects_are_native_and_design_uses_template_tokens(client):
    from app import database as db
    template=db.template_get('tech')
    slides=[chart(),Slide(title='Таблица',kind='table',table=[['Группа','Число'],['А','10']]),
            Slide(title='Процесс',kind='diagram',bullets=['Подготовить','Проверить','Запустить'])]
    content=DeckContent(title='Визуалы',slides=slides)
    prs=Presentation(BytesIO(export_pptx(content,template,'a')))
    assert any(s.has_chart for s in prs.slides[0].shapes)
    assert any(s.has_table for s in prs.slides[1].shapes)
    assert not any(s.shape_type==13 for s in prs.slides[2].shapes)
    assert len(prs.slides[2].shapes)>=4
    design=SlideDesign(composition='split',background_role='dark',heading_role='body',heading_weight='normal')
    metadata=design_metadata(template['metadata'],design)
    assert metadata['background'] in template['metadata']['colors'].values()
    slide=Slide(title='Фон',body='Текст без изменений',design=design)
    scene=build_scene(slide,template['metadata'],'a',0)
    assert scene['objects'][0]['override_template']

def test_outline_autonomously_selects_native_visuals_and_retries_bad_values(client,monkeypatch):
    from app import database as db
    from app.models import OutlineRequest
    monkeypatch.setattr(settings,'mode','live')
    calls=[]
    async def respond(system,payload):
        calls.append(payload)
        return DeckContent(title='Данные',slides=[chart(values=[999,20] if len(calls)==1 else [10,20]),
            Slide(title='Таблица',kind='table',table=[['Команда','Доля %'],['Команда А','10'],['Команда Б','20']]),
            Slide(title='Шаги',kind='diagram',bullets=['Подготовить','Запустить'])]).model_dump()
    monkeypatch.setattr(llm,'complete_json',respond)
    result=asyncio.run(llm.make_outline(OutlineRequest(template_id='tech',count=3,prompt='Команда А: 10%; Команда Б: 20%\nПодготовить, запустить.'),db.template_get('tech')))
    assert [s.kind for s in result.slides]==['chart','table','diagram'] and len(calls)==2
    assert result.slides[0].chart.values==[10,20] and 'format_correction' in calls[-1]

@pytest.mark.parametrize('kind',['chart','table','diagram','icons','text'])
def test_assistant_converts_without_new_numbers(client,monkeypatch,kind):
    from app import database as db
    from app.models import AssistantRequest
    from app.services.assistant import edit_presentation
    monkeypatch.setattr(settings,'mode','live')
    deck=DeckContent(title='Исходник',slides=[Slide(title='Команды',body='Команда А: 10%; Команда Б: 20%')])
    patch={'slide':0,'kind':kind}
    if kind=='chart': patch['chart']=chart().chart.model_dump()
    if kind=='table': patch['table']=[['Команда','Доля %'],['Команда А','10'],['Команда Б','20']]
    if kind in {'diagram','icons'}: patch['bullets']=['Команда А','Команда Б']
    async def respond(*args,**kwargs): return {'summary':'Представление изменено','changes':[patch]}
    monkeypatch.setattr(llm,'complete_json',respond)
    result,_=asyncio.run(edit_presentation(AssistantRequest(content=deck,variant='a',instruction='Измени представление',slide=0,base_updated_at='test'),db.template_get('tech'),''))
    assert result.slides[0].kind==kind and result.slides[0].body==deck.slides[0].body
