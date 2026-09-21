"""Реальные API/экспорты; облачная интеграция отдельно требует тестовый PostgreSQL."""
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import os
import pytest
import psycopg
from fastapi.testclient import TestClient
from pptx import Presentation
from app import database as db
from app.config import settings
from app.main import app
from app.models import DeckContent, Slide, ChartData
from app.services import cloud
from app.services.layout import build_scene
from app.services.export import export_pptx, export_pdf, export_html


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'database', tmp_path/'test.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path/'data')
    monkeypatch.setattr(settings, 'cloud_database_url', '')
    with TestClient(app) as client:
        response=client.post('/api/auth/register',json={'email':'experience@example.test','password':'Test-Experience-4761','first_name':'Тест','last_name':'Интерфейса'})
        assert response.status_code==201
        client.headers['X-CSRF-Token']=response.json()['csrf']
        owner=response.json()['user']['id']
        db.project_create('experience','tech',DeckContent(title='Проверка интерфейса',slides=[Slide(title='TODO',body='Факт',source_quote='Факт')]).model_dump(),'Факт',owner)
        yield client


def test_preview_is_authenticated_and_never_changes_saved_content(client):
    project=client.get('/api/projects/experience').json()
    draft=project['content'];draft['slides'][0].update(kind='table',title='Таблица',table=[['Этап','Статус'],['Дизайн','Готово']])
    response=client.post('/api/projects/experience/preview',json={'content':draft,'variant':'a'})
    assert response.status_code==200,response.text
    assert any(o['type']=='table' for o in response.json()['scenes'][0]['objects'])
    saved=client.get('/api/projects/experience').json()
    assert saved['content']['slides'][0]['title']=='TODO' and not saved['can_undo']
    assert client.post('/api/projects/experience/preview',json={'content':draft},headers={'X-CSRF-Token':'wrong'}).status_code==403
    client.post('/api/auth/logout')
    assert client.post('/api/projects/experience/preview',json={'content':draft}).status_code==401


def test_manual_audit_edit_rechecks_and_undo_restores_issue(client):
    original=client.get('/api/projects/experience').json()
    assert any(i['code']=='placeholder' for i in client.get('/api/projects/experience/audit').json()['issues'])
    original['content']['slides'][0]['title']='Результаты исследования'
    response=client.put('/api/projects/experience',json={'content':original['content'],'variant':'a'})
    assert response.status_code==200
    assert not any(i['code']=='placeholder' for i in client.get('/api/projects/experience/audit').json()['issues'])
    assert client.post('/api/projects/experience/undo').status_code==200
    assert any(i['code']=='placeholder' for i in client.get('/api/projects/experience/audit').json()['issues'])


def test_missing_cloud_is_not_reported_as_saved(client):
    assert client.get('/api/cloud/status').json()=={'configured':False,'ready':False,'templates':0,'projects':0}
    response=client.post('/api/projects/experience/cloud')
    assert response.status_code==503 and 'revision' not in response.json()
    assert client.post('/api/projects/missing/cloud').status_code==404
    assert client.get('/api/projects/experience').status_code==200


def test_cloud_failure_does_not_expose_credentials(client, monkeypatch):
    monkeypatch.setattr(settings,'cloud_database_url','postgresql://hidden:private-secret@example.invalid/db')
    def unavailable(*args, **kwargs): raise psycopg.OperationalError('private-secret')
    monkeypatch.setattr(cloud.psycopg,'connect',unavailable)
    response=client.get('/api/cloud/status')
    assert response.status_code==503 and 'private-secret' not in response.text and 'example.invalid' not in response.text


def test_cloud_save_is_owner_scoped(client, monkeypatch):
    monkeypatch.setattr(cloud,'save_project',lambda *args: pytest.fail('Чужой проект не должен доходить до облака'))
    result=client.post('/api/auth/register',json={'email':'other-cloud@example.test','password':'Test-Experience-4761','first_name':'Другой','last_name':'Пользователь'})
    client.headers['X-CSRF-Token']=result.json()['csrf']
    assert client.post('/api/projects/experience/cloud').status_code==404


def test_cloud_restore_reconstructs_source_pptx_as_a_separate_copy(client, monkeypatch):
    project=db.project_get('experience');template=db.template_get('tech')
    binary=export_pptx(DeckContent.model_validate(project['content']),template,'a')
    def load(pid, owner):
        assert pid=='remote-copy' and owner==project['user_id']
        return {'payload':{**project,'variant':'c'},'metadata':template['metadata'],'name':template['name'],'pptx':binary}
    monkeypatch.setattr(cloud,'load_project',load)
    response=client.post('/api/cloud/projects/remote-copy/restore')
    assert response.status_code==201,response.text
    restored=response.json()
    assert restored['id']!='experience' and restored['variant']=='c'
    assert restored['content']==project['content']
    local_template=db.template_get(restored['template_id'])
    assert local_template['user_id']==project['user_id']
    assert Path(local_template['path']).read_bytes()==binary
    assert db.project_get('experience')['variant']=='a'


def test_visual_types_have_distinct_geometry_and_editable_exports(client):
    template=db.template_get('tech')
    slides=[Slide(title='История',kind='title',body='Начало'),
            Slide(title='Сравнение',kind='table',table=[['Этап','Результат'],['Дизайн','Готово']]),
            Slide(title='Динамика',kind='chart',chart=ChartData(labels=['А','Б'],values=[-10,20],unit='%')),
            Slide(title='Дальше',kind='steps',bullets=['Проверить','Сохранить','Показать'])]
    scenes=[build_scene(s,template['metadata'],'a',i) for i,s in enumerate(slides)]
    assert any(o['id']=='title-underline' for o in scenes[0]['objects'])
    assert any(o['type']=='table' for o in scenes[1]['objects'])
    assert any(o['type']=='chart' for o in scenes[2]['objects'])
    assert len([o for o in scenes[3]['objects'] if o['id'].startswith('step-card-')])==3
    assert all(n['type'] in {'text','rect'} for scene in scenes for n in scene['render_objects'])
    assert all(n['x']>=0 and n['x']+n['w']<=scene['width']+1 for scene in scenes for n in scene['render_objects'])
    content=DeckContent(title='Типы слайдов',slides=slides)
    deck=Presentation(BytesIO(export_pptx(content,template,'a')))
    assert any(s.has_table for s in deck.slides[1].shapes)
    assert any(s.has_chart for s in deck.slides[2].shapes)
    assert export_pdf(content,template,'a').startswith(b'%PDF')
    assert '<svg' in export_html(content,template,'a').decode()


@pytest.mark.skipif(not os.getenv('TEST_CLOUD_DATABASE_URL'), reason='Нужна отдельная тестовая БД PostgreSQL')
def test_postgres_roundtrip_isolation_idempotency_and_rollback(client, monkeypatch):
    monkeypatch.setattr(settings,'cloud_database_url',os.environ['TEST_CLOUD_DATABASE_URL'])
    project=db.project_get('experience');owner=project['user_id'];template=db.template_get('tech')
    template_path=settings.storage/'cloud-test-template.pptx'
    template_path.write_bytes(export_pptx(DeckContent.model_validate(project['content']),template,'a'))
    template['path']=str(template_path)
    try:
        receipt=cloud.save_project(project,[template],owner)
        assert receipt['revision']==1
        assert cloud.save_project(project,[template],owner)['revision']==1
        saved=cloud.load_project('experience',owner)
        assert saved['payload']['content']==project['content']
        assert bytes(saved['pptx'])==open(template['path'],'rb').read()
        assert cloud.load_project('experience',str(uuid4())) is None
        project['content']['title']='Новая версия';project['updated_at']=db.now()
        assert cloud.save_project(project,[template],owner)['revision']==2
        with pytest.raises(cloud.CloudError):
            with cloud.connection() as conn:
                conn.execute('DELETE FROM deckly_private.presentations WHERE owner_id=%s',(owner,))
                conn.execute('SELECT 1/0')
        assert cloud.load_project('experience',owner) is not None
    finally:
        with cloud.connection() as conn:
            conn.execute('DELETE FROM deckly_private.presentations WHERE owner_id=%s',(owner,))
            conn.execute('DELETE FROM deckly_private.templates WHERE owner_id=%s',(owner,))
