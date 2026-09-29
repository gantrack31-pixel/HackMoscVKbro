import base64
import asyncio
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import get_args
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader
import pytest
from pydantic import ValidationError
from app.icon_catalog import ICONS, IconName, model_icons
from app.models import Slide, SlidePatch, DeckContent, OutlineRequest
from app.services.icons import ASSETS, ATTRIBUTION, icon_data
from app.services.layout import build_scene
from app.services.export import export_pptx, export_html, export_pdf
from app.services.palette import contrast, roles


@pytest.mark.parametrize('name', list(ICONS))
def test_bundled_icons_have_originals_transparency_and_safe_theme_color(name):
    assert set(get_args(IconName)) == set(ICONS)
    assert (ASSETS / f'{name}.svg').stat().st_size > 100
    with Image.open(ASSETS / f'{name}.png') as source:
        alpha = source.convert('RGBA').getchannel('A')
        assert source.size == (192, 192)
        assert alpha.getextrema() == (0, 255)
    image = Image.open(BytesIO(base64.b64decode(icon_data(name, '#457891').split(',')[1])))
    assert image.getchannel('A').tobytes() == alpha.tobytes()
    assert all(pixel[:3] == (69, 120, 145) for pixel in image.get_flattened_data() if pixel[3])
    assert model_icons()['icons'][name]


def test_icon_allowlist_and_assistant_patch():
    assert SlidePatch(slide=0, icon_names=['heart','education']).icon_names == ['heart','education']
    for unsafe in ('../../config', 'https://example.com/icon.svg', '<svg/>'):
        with pytest.raises(ValidationError): Slide(title='Title', icon_names=[unsafe])
        with pytest.raises(ValueError): icon_data(unsafe, '#123456')
    with pytest.raises(ValueError): icon_data('clock', 'url(x)')


@pytest.mark.parametrize('background', ['#FFFFFF', '#0E1821'])
@pytest.mark.parametrize('ratio', [16/9, 4/3, 3/4])
def test_icons_geometry_and_all_exports(background, ratio):
    meta={'ratio':ratio,'accent':'#5864FF','background':background,'colors':{'lt1':'#FFFFFF','dk1':'#0E1821'}}
    slide=Slide(title='Забота и развитие', kind='icons', bullets=['Гигиена','Команда','Обучение','Проверка','Рост','Сроки'],
                icon_names=['heart','people','education','check','growth','clock'])
    deck=DeckContent(title='Тест', slides=[slide]); template={'metadata':meta,'path':None}
    for variant in 'abc':
        scene=build_scene(slide,meta,variant,0)
        pictures=[o for o in scene['objects'] if o.get('icon_name')]
        assert len(pictures)==6
        for obj in scene['objects']:
            assert obj['x']>=0 and obj['y']>=0
            assert obj['x']+obj['w']<=scene['width']+.1
            assert obj['y']+obj['h']<=scene['height']+.1
            if obj['type']=='text': assert obj['used_height']<=obj['h']+.1
    prs=Presentation(BytesIO(export_pptx(deck,template,'a')))
    assert sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in prs.slides[0].shapes)==6
    assert ATTRIBUTION in prs.slides[0].notes_slide.notes_text_frame.text
    assert ATTRIBUTION in PdfReader(BytesIO(export_pdf(deck,template,'a'))).metadata.subject
    assert ATTRIBUTION in export_html(deck,template,'a').decode()
    title=next(o for o in scene['objects'] if o['id']=='title')
    assert contrast(title['color'],roles(meta)['background'])>=3


def test_no_invented_bullets_or_unrelated_icons():
    meta={'ratio':16/9, 'accent':'#315A88'}
    empty=build_scene(Slide(title='Empty',kind='icons'),meta,'a',0)
    assert not any(o['id'].startswith('icon-') for o in empty['objects'])
    slide=Slide(title='Specific',kind='icons',bullets=['Гигиена рук','Квантовая запутанность'])
    scene=build_scene(slide,meta,'a',0)
    assert [o['icon_name'] for o in scene['objects'] if o.get('icon_name')]==['heart']


def test_prebuilt_scenes_reused_without_changes(monkeypatch):
    from app.services import audit, export
    template={'metadata':{'ratio':16/9,'accent':'#284B77'}, 'path':None}
    deck=DeckContent(title='Reuse',slides=[Slide(title='Icons',kind='icons',bullets=['Здоровье'],icon_names=['heart'])])
    scenes=[build_scene(deck.slides[0],template['metadata'],'a',0)]
    before=deepcopy(scenes)
    expected=audit.audit_deck(deck,template,'a','')
    def no_rebuild(*args): raise AssertionError('Repeated layout work')
    monkeypatch.setattr(audit,'build_scene',no_rebuild)
    monkeypatch.setattr(export,'build_scene',no_rebuild)
    assert audit.audit_deck(deck,template,'a','',scenes=scenes)==expected
    assert export.export_pptx(deck,template,'a',scenes=scenes).startswith(b'PK')
    assert scenes==before
    with pytest.raises(ValueError): audit.audit_deck(deck,template,'a','',scenes=[])
    with pytest.raises(ValueError): export.export_pptx(deck,template,'a',scenes=[])


def test_outline_model_can_select_library_icons(monkeypatch):
    from app.services import llm
    from app.config import settings
    monkeypatch.setattr(settings,'mode','live')
    async def model(system,payload,template):
        assert payload['icon_library']['icons']['heart']
        assert 'icon_names' in system
        return DeckContent(title='Здоровье',slides=[
            Slide(title='Здоровье',kind='title'),
            Slide(title='Принципы',kind='icons',bullets=['Здоровье','Обучение'],icon_names=['heart','education']),
            Slide(title='Следующий шаг')]).model_dump()
    monkeypatch.setattr(llm,'complete_with_template',model)
    result=asyncio.run(llm.make_outline(OutlineRequest(template_id='test',prompt='Гигиена',count=3),
                                     {'id':'test','metadata':{'ratio':16/9},'name':'test'}))
    assert result.slides[1].icon_names==['heart','education']
