"""Free AI canvas edits and literal unitless charts survive review and exports."""
import asyncio
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
import pytest

from app.models import DeckContent, Slide, SlideDesign, SlidePatch
from app.services import llm, images
from app.services.assistant import apply_patch
from app.services.assistant_context import presentation_context
from app.services.export import export_html, export_pdf, export_pptx
from app.services.layout import build_scene
from app.services.visual_grounding import numeric_series, validate_visual
from test_assistant import client, request_for

TEMPLATE = {'name': 'Canvas', 'metadata': {'ratio': 16/9}}


def respond(monkeypatch, changes):
    async def answer(system, payload, template):
        return {'summary': 'Предлагаю свободную композицию.', 'changes': changes}
    monkeypatch.setattr(llm, 'complete_with_template', answer)


@pytest.mark.parametrize('series,values', [
    ('1, 2, 3, 4, 5', [1,2,3,4,5]),
    ('1,2,3,4,5', [1,2,3,4,5]),
    ('2, 2, 3, 3', [2,2,3,3]),
    ('-1.5, 0, 2.5', [-1.5,0,2.5]),
    ('1,5; 2,5; -3,0', [1.5,2.5,-3]),
])
def test_requested_unitless_column_chart_is_reviewable_and_native(client, monkeypatch, series, values):
    body = request_for(client, instruction='Сделай вместо текста столбчатую диаграмму на основе данных '+series,
                       preserve_facts=False)
    body['content']['slides'][0]['body'] = 'Предыдущий результат: 30 участников.'
    def chart_data():
        return {'labels': [str(i) for i in range(1,len(values)+1)], 'values': values, 'unit': '', 'chart_type': 'column'}
    async def answer(system, payload, template):
        assert payload['provided_numeric_series'][0]['values'] == values
        return {'summary': 'Столбчатая диаграмма по предоставленным значениям.',
                'changes': [{'slide':0, 'title':'Данные', 'body':'', 'kind':'chart', 'chart':chart_data()}]}
    monkeypatch.setattr(llm, 'complete_with_template', answer)
    response = client.post('/api/projects/ai-test/assistant', json=body)
    assert response.status_code == 200, response.text
    proposal = response.json()['proposal']
    assert proposal['content']['slides'][0]['chart']['values'] == values
    assert proposal['content']['slides'][0]['chart']['unit'] == ''
    applied = client.post(f'/api/projects/ai-test/assistant/proposals/{proposal["id"]}/apply', json=body)
    assert applied.status_code == 200, applied.text
    pptx = client.get('/api/projects/ai-test/export/pptx')
    chart = next(shape.chart for shape in Presentation(BytesIO(pptx.content)).slides[0].shapes if shape.has_chart)
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert list(chart.series[0].values) == values


def test_unitless_named_values_do_not_need_a_fabricated_unit():
    slide = Slide(title='Scores', kind='chart', chart={'labels':['Alpha','Beta'], 'values':[10,20], 'unit':''})
    validate_visual(slide, 'Alpha: 10; Beta: 20')
    with pytest.raises(ValueError): validate_visual(slide, 'Alpha: 10%; Beta: 20%')


@pytest.mark.parametrize('data', [
    {'labels':['1','2','3'], 'values':[1,9,3], 'unit':''},
    {'labels':['1','2','3'], 'values':[3,2,1], 'unit':''},
    {'labels':['Revenue','Cost','Profit'], 'values':[1,2,3], 'unit':''},
    {'labels':['1','2','3'], 'values':[1,2,3], 'unit':'руб.'},
])
def test_plain_series_cannot_authorize_invented_values_labels_or_units(data):
    with pytest.raises(ValueError):
        validate_visual(Slide(title='Chart', kind='chart', chart=data), 'Данные 1, 2, 3')


def test_move_paint_add_hide_review_apply_and_undo(client, monkeypatch):
    body = request_for(client, instruction='Сделай фон синим, заголовок красным сверху, добавь круг и подпись',
                       preserve_facts=False)
    respond(monkeypatch, [{'slide':0, 'title':'Композиция', 'design': {
        'background_color':'#123456', 'objects':{'title':{'x':80,'y':40,'w':800,'h':100,'color':'#FF0000','font_size':48},
                                                  'body':{'hidden':True}}},
        'elements':[
            {'id':'custom-circle','type':'ellipse','x':900,'y':100,'w':160,'h':160,'fill':'#FF9900'},
            {'id':'custom-note','type':'text','x':100,'y':250,'w':500,'h':100,'text':'Новая идея','color':'#FFFFFF'},
        ]}])
    response = client.post('/api/projects/ai-test/assistant', json=body)
    assert response.status_code == 200, response.text
    proposal = response.json()['proposal']
    preview = client.post('/api/projects/ai-test/preview', json={'content':proposal['content'],'variant':'a'})
    scene = preview.json()['scenes'][0]
    title = next(obj for obj in scene['objects'] if obj['id']=='title')
    assert (title['x'],title['y'],title['color'],title['font_size']) == (80,40,'#FF0000',48)
    assert scene['objects'][0]['fill'] == '#123456'
    assert not any(obj['id']=='body' for obj in scene['render_objects'])
    assert next(obj for obj in scene['objects'] if obj['id']=='custom-circle')['fill'] == '#FF9900'
    applied = client.post(f'/api/projects/ai-test/assistant/proposals/{proposal["id"]}/apply', json=body)
    assert applied.status_code == 200, applied.text
    content = DeckContent.model_validate(applied.json()['project']['content'])
    native = Presentation(BytesIO(export_pptx(content,TEMPLATE,'a')))
    shape = next(shape for shape in native.slides[0].shapes if shape.has_text_frame and shape.text=='Композиция')
    assert shape.left == int(80 * native.slide_width/1280)
    assert str(shape.text_frame.paragraphs[0].font.color.rgb) == 'FF0000'
    html = export_html(content,TEMPLATE,'a').decode()
    assert '#FF9900' in html and 'Новая идея' in html and '#123456' in html
    assert export_pdf(content,TEMPLATE,'a').startswith(b'%PDF')
    undo = client.post('/api/projects/ai-test/undo',json={'base_updated_at':applied.json()['project']['updated_at']})
    assert undo.json()['content'] == body['content']


@pytest.mark.parametrize('design,options', [
    ({'objects':{'title':{'x':1250,'w':800}}}, {}),
    ({'objects':{'unknown-object':{'x':80}}}, {}),
    ({'objects':{'title':{'hidden':True}}}, {'locked_fields':[{'slide':0,'field':'title'}]}),
    ({'background_color':'url(https://example.test)'}, {}),
])
def test_invalid_canvas_edits_never_publish(client, monkeypatch, design, options):
    body = request_for(client, **options)
    respond(monkeypatch,[{'slide':0,'design':design}])
    response = client.post('/api/projects/ai-test/assistant',json=body)
    assert response.status_code == 502, response.text
    assert client.get('/api/projects/ai-test').json()['content'] == body['content']


def test_partial_style_keeps_previous_position_and_color():
    previous = Slide(title='Title', design={'composition':'grid','objects':{'title':{'x':80,'color':'#FF0000'}}})
    updated = apply_patch(previous,SlidePatch(slide=0,design=SlideDesign(objects={'title':{'y':120}})),'a')
    style = updated.designs['a'].objects['title']
    assert (style.x,style.y,style.color) == (80,120,'#FF0000')
    assert updated.designs['a'].composition == 'grid'


def test_edit_individual_chart_mark_exports_editable_shapes():
    slide = Slide(title='Chart',kind='chart',chart={'labels':['1','2','3'],'values':[1,2,3],'unit':'','chart_type':'column'})
    plain = build_scene(slide,TEMPLATE['metadata'],'a',0)
    chart_color = next(obj['fill'] for obj in plain['objects'] if obj['id']=='chart')
    mark = next(obj for obj in plain['render_objects'] if obj['type']=='rect' and obj['id'].startswith('chart-') and obj['fill']==chart_color)
    slide.design = SlideDesign(objects={mark['id']:{'fill':'#FF00FF'}})
    scene = build_scene(slide,TEMPLATE['metadata'],'a',0)
    assert next(obj for obj in scene['render_objects'] if obj['id']==mark['id'])['fill']=='#FF00FF'
    assert any(obj['id']==mark['id'] and obj['fill']=='#FF00FF' for obj in scene['objects'])
    native = Presentation(BytesIO(export_pptx(DeckContent(title='Deck',slides=[slide]),TEMPLATE,'a')))
    assert any(shape.fill.type == MSO_FILL_TYPE.SOLID and str(shape.fill.fore_color.rgb)=='FF00FF' for shape in native.slides[0].shapes if not shape.has_text_frame or not shape.text)


@pytest.mark.parametrize('source', ['Данные 1, 2, 3 кг', 'Значения (кг): 1, 2, 3'])
def test_literal_series_keeps_provided_unit(source):
    slide = Slide(title='Chart',kind='chart',chart={'labels':['1','2','3'],'values':[1,2,3],'unit':'кг'})
    validate_visual(slide,source)
    with pytest.raises(ValueError):
        validate_visual(slide.model_copy(update={'chart':slide.chart.model_copy(update={'unit':''})}),source)


def test_csv_unit_cannot_be_dropped():
    slide = Slide(title='Chart',kind='chart',chart={'labels':['Alpha','Beta'],'values':[10,20],'unit':''})
    with pytest.raises(ValueError): validate_visual(slide,'Команда,Доля %\nAlpha,10\nBeta,20')


def test_generate_canvas_image_keeps_binary_out_of_model_context(monkeypatch):
    import base64
    buffer = BytesIO();Image.new('RGB',(16,16),'red').save(buffer,format='PNG')
    png = 'data:image/png;base64,'+base64.b64encode(buffer.getvalue()).decode()
    async def generate(prompt): return {'image_data':png}
    monkeypatch.setattr(images,'generate_image',generate)
    content = DeckContent(title='Deck',slides=[Slide(title='Image',elements=[
        {'id':'custom-picture','type':'image','x':80,'y':200,'w':200,'h':200,'image_prompt':'Красная композиция'}])])
    asyncio.run(images.fill_missing_images(content))
    assert content.slides[0].element_images['custom-picture']==png
    assert png not in str(presentation_context(content,[0]))
    scene = build_scene(content.slides[0],TEMPLATE['metadata'],'a',0)
    assert next(obj for obj in scene['objects'] if obj['id']=='custom-picture')['src']==png


def test_qualitative_table_can_be_created_in_free_mode(client,monkeypatch):
    body = request_for(client,instruction='Добавь таблицу преимуществ и рисков',preserve_facts=False)
    respond(monkeypatch,[{'slide':0,'title':'Сравнение','kind':'table','body':'',
                         'table':[['Преимущество','Риск'],['Простота','Зависимость от контекста']]}])
    response = client.post('/api/projects/ai-test/assistant',json=body)
    assert response.status_code==200,response.text
