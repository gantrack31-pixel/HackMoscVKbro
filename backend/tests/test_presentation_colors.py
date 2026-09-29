from io import BytesIO
from copy import deepcopy
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from app.models import DeckContent,Slide,SlideDesign
from app.services.templates import analyze_template
from app.services.layout import build_scene
from app.services.export import export_pptx
from app.services.palette import contrast,surface_at,design_metadata
from app.services.audit import audit_deck
from test_workflow import client,register

def make_template(path,background):
    prs=Presentation()
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb=RGBColor.from_string('222222')
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(background)
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text='Unknown template title'
    slide.shapes.add_textbox(Inches(1),Inches(2.5),Inches(8),Inches(3)).text='Unknown template body'
    prs.save(path)
    return {'id':'unknown','name':'Unknown','path':str(path),'metadata':analyze_template(path)}

@pytest.mark.parametrize('background',['FFFFFF','151515','377FBE'])
@pytest.mark.parametrize('variant',['a','b','c'])
def test_unknown_template_uses_slide_background_not_other_master(tmp_path,background,variant):
    template=make_template(tmp_path/'unknown.pptx',background)
    assert template['metadata']['background']=='#'+background
    slide=Slide(title='Читаемый заголовок',body='Важный текст',bullets=['Пункт один','Пункт два'])
    scene=build_scene(slide,template['metadata'],variant,0)
    assert scene['objects'][0]['fill']=='#'+background
    prs=Presentation(BytesIO(export_pptx(DeckContent(title='Тест',slides=[slide]),template,variant)))
    assert str(prs.slides[0].background.fill.fore_color.rgb)==background
    assert_readable(scene)

def assert_readable(scene):
    objects=scene['render_objects']
    for i,obj in enumerate(objects):
        if obj['type']=='text' or obj.get('icon_color'):
            color=obj.get('icon_color') or obj['color']
            assert contrast(color,surface_at(obj,objects[:i],scene['objects'][0]['fill']))>=4.5,(obj['id'],color)

@pytest.mark.parametrize('kind',['text','steps','chart','table','diagram','icons'])
@pytest.mark.parametrize('background',['#FFFFFF','#111111'])
def test_bad_model_foregrounds_fixed_for_all_visuals_without_content_changes(kind,background):
    metadata={'ratio':16/9,'background':background,'text_color':background,'accent':background,
              'colors':{'lt1':'#FFFFFF','dk1':'#111111','accent1':'#4477AA'},'font':'Arial','layouts':[]}
    slide=Slide(title='Заголовок',kind=kind,body='Текст',bullets=['Первый','Второй'],
                chart={'labels':['А','Б'],'values':[10,20],'unit':'%'},table=[['Имя','Доля'],['А','10']],
                design=SlideDesign(composition='grid',foreground_role='inverse'))
    before=slide.model_dump()
    for variant in 'abc':
        scene=build_scene(slide,metadata,variant,0)
        assert_readable(scene)
        assert slide.model_dump()==before
        assert not any(i.code.startswith('contrast') for i in audit_deck(DeckContent(title='Тест',slides=[slide]),{'metadata':metadata},variant,''))

def test_intentional_background_change_recomputes_foregrounds_and_native_chart(tmp_path):
    template=make_template(tmp_path/'unknown.pptx','FFFFFF')
    slide=Slide(title='График',kind='chart',chart={'labels':['А'],'values':[12],'unit':'шт.'},design=SlideDesign(composition='grid',background_role='dark'))
    scene=build_scene(slide,template['metadata'],'a',0)
    assert scene['objects'][0]['fill']!='#FFFFFF' and scene['objects'][0]['override_template']
    assert_readable(scene)
    prs=Presentation(BytesIO(export_pptx(DeckContent(title='Тест',slides=[slide]),template,'a')))
    chart=next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert contrast('#'+str(chart.font.color.rgb),scene['objects'][0]['fill'])>=4.5

def test_complex_background_not_replaced_by_model():
    metadata={'source':'pptx','background':'#EFEFEE','background_info':{'kind':'complex'},'colors':{'dk1':'#000000','lt1':'#FFFFFF'}}
    assert design_metadata(metadata,SlideDesign(composition='grid',background_role='dark'))['background']=='#EFEFEE'


def test_audit_checks_nested_chart_table_labels_and_footer():
    metadata={'ratio':16/9,'background':'#FFFFFF','accent':'#0077FF','colors':{'lt1':'#FFFFFF','dk1':'#000000'}}
    slide=Slide(title='Таблица',kind='table',table=[['Колонка'],['Ячейка']])
    scene=build_scene(slide,metadata,'a',0)
    for obj in scene['render_objects']:
        if obj['type']=='text': obj['color']=surface_at(obj,scene['render_objects'][:scene['render_objects'].index(obj)],'#FFFFFF')
    issues=audit_deck(DeckContent(title='Тест',slides=[slide]),{'metadata':metadata},'a','',scenes=[scene])
    objects={i.object_id for i in issues if i.code=='contrast_visual'}
    assert 'footer' in objects and any(o.startswith('table-') for o in objects)

def test_preview_failure_does_not_block_valid_upload(client,tmp_path,monkeypatch):
    register(client)
    path=tmp_path/'new.pptx';make_template(path,'FFFFFF')
    def fail(*args): raise RuntimeError('private converter detail')
    monkeypatch.setattr('app.main.thumbnail',fail)
    response=client.post('/api/templates',files={'file':('new.pptx',path.read_bytes(),'application/vnd.openxmlformats-officedocument.presentationml.presentation')})
    assert response.status_code==201 and response.json()['metadata']['source']=='pptx'
    assert 'private converter' not in response.text
