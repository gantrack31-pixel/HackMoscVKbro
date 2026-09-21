"""Профиль и Яндекс проверяются на изолированной БД и подменённом провайдере."""
from io import BytesIO
from urllib.parse import parse_qs, urlparse
import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app.main import app
from app.config import settings
from app import database as db
from app.models import DeckContent, Slide
from app.services.layout import build_scene
from app.services.export import export_pptx
from app.services.theme_catalog import THEMES

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'database', tmp_path/'accounts.sqlite3')
    monkeypatch.setattr(settings, 'storage', tmp_path/'storage')
    monkeypatch.setattr(settings, 'mode', 'demo')
    with TestClient(app) as client:
        yield client

def register(client, email='owner@example.test'):
    result=client.post('/api/auth/register', json={'email':email,'password':'Example-password-387','first_name':'Анна','last_name':'Тестовая'})
    assert result.status_code==201
    client.headers['X-CSRF-Token']=result.json()['csrf']
    return result.json()['user']

def mock_yandex(monkeypatch, email='owner@example.test', identifier='yandex-567'):
    monkeypatch.setattr(settings,'yandex_id','public-client-id')
    monkeypatch.setattr(settings,'yandex_secret','')
    original=httpx.AsyncClient
    def handle(request):
        if request.url.host=='oauth.yandex.ru':
            data=parse_qs(request.content.decode())
            assert 'client_secret' not in data and len(data['code_verifier'][0])>=43
            return httpx.Response(200,json={'access_token':'mock-provider-token'})
        assert request.headers['Authorization']=='OAuth mock-provider-token'
        return httpx.Response(200,json={'id':identifier,'default_email':email,'first_name':'Яндекс','last_name':'Тест'})
    monkeypatch.setattr('app.auth.httpx.AsyncClient',lambda **kwargs:original(transport=httpx.MockTransport(handle),**kwargs))

def finish(client, url):
    state=parse_qs(urlparse(url).query)['state'][0]
    return client.get('/api/auth/yandex/callback',params={'state':state,'code':'one-use-code'},follow_redirects=False)

def test_profile_persistence_validation_and_isolation(client):
    first=register(client)
    data={'first_name':'  Мария  ','last_name':'Иванова','avatar_color':'#29745F'}
    assert client.put('/api/auth/profile',json=data,headers={'X-CSRF-Token':'bad'}).status_code==403
    updated=client.put('/api/auth/profile',json=data)
    assert updated.status_code==200
    assert updated.json()['first_name']=='Мария' and updated.json()['password_enabled']
    assert client.get('/api/auth/me').json()['user']['avatar_color']=='#29745F'
    assert client.put('/api/auth/profile',json={**data,'first_name':'   '}).status_code==422
    assert client.put('/api/auth/profile',json={**data,'avatar_color':'red'}).status_code==422
    client.put('/api/favorites/strategy')
    assert client.get('/api/auth/profile/stats').json()=={'projects':0,'favorites':1,'templates':0}
    register(client,'other@example.test')
    assert client.get('/api/auth/profile/stats').json()['favorites']==0
    with db.connection() as connection:
        assert connection.execute('SELECT first_name FROM users WHERE id=?',(first['id'],)).fetchone()[0]=='Мария'

def test_yandex_pkce_without_secret_and_persistent_login(client,monkeypatch):
    mock_yandex(monkeypatch)
    assert client.get('/api/health').json()['yandex_enabled'] is True
    start=client.get('/api/auth/yandex/start',follow_redirects=False)
    assert finish(client,start.headers['location']).headers['location'].endswith('#home')
    user=client.get('/api/auth/me').json()['user']
    assert user['yandex_connected'] and not user['password_enabled']
    client.cookies.clear()
    start=client.get('/api/auth/yandex/start',follow_redirects=False)
    finish(client,start.headers['location'])
    assert client.get('/api/auth/me').json()['user']['id']==user['id']

def test_yandex_email_collision_requires_explicit_link(client,monkeypatch):
    user=register(client)
    mock_yandex(monkeypatch)
    start=client.get('/api/auth/yandex/start',follow_redirects=False)
    assert 'auth-error=exists' in finish(client,start.headers['location']).headers['location']
    assert not client.get('/api/auth/me').json()['user']['yandex_connected']
    assert client.post('/api/auth/yandex/link',headers={'X-CSRF-Token':'wrong'}).status_code==403
    link=client.post('/api/auth/yandex/link')
    assert link.status_code==200
    assert finish(client,link.json()['url']).headers['location'].endswith('#profile?yandex=connected')
    result=client.get('/api/auth/me').json()['user']
    assert result['id']==user['id'] and result['yandex_connected'] and result['password_enabled']

def test_yandex_link_rejects_changed_session(client,monkeypatch):
    user=register(client)
    mock_yandex(monkeypatch)
    link=client.post('/api/auth/yandex/link').json()['url']
    client.post('/api/auth/login',json={'email':user['email'],'password':'Example-password-387'})
    assert 'yandex_error=state' in finish(client,link).headers['location']
    assert not client.get('/api/auth/me').json()['user']['yandex_connected']

def test_yandex_cannot_link_another_users_identity(client,monkeypatch):
    mock_yandex(monkeypatch,email='oauth@example.test')
    start=client.get('/api/auth/yandex/start',follow_redirects=False)
    finish(client,start.headers['location'])
    register(client,'another@example.test')
    link=client.post('/api/auth/yandex/link').json()['url']
    assert 'yandex_error=linked' in finish(client,link).headers['location']
    assert not client.get('/api/auth/me').json()['user']['yandex_connected']

@pytest.mark.parametrize('tid',[theme[0] for theme in THEMES])
def test_new_theme_is_real_editable_export(client,tid):
    register(client)
    catalog=client.get('/api/templates').json()
    assert len(catalog)==18
    public=next(t for t in catalog if t['id']==tid)
    assert public['metadata']['source']=='starter' and 'path' not in public
    template=db.template_get(tid)
    slide=Slide(title='Проверка новой темы',body='Редактируемое содержание')
    scene=build_scene(slide,template['metadata'],'a',0)
    assert scene['objects'][0]['fill']==template['metadata']['background']
    assert any(item['id'].startswith('theme-') for item in scene['objects'])
    content=DeckContent(title='Пример',slides=[slide])
    for variant in ['a','b','c']:
        deck=Presentation(BytesIO(export_pptx(content,template,variant)))
        assert len(deck.slides)==1
        assert any(s.has_text_frame and 'Редактируемое содержание' in s.text for s in deck.slides[0].shapes)
