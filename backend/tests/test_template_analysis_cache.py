from pptx import Presentation

from app.services import templates


def test_unchanged_pptx_analysis_is_reused_without_sharing_mutable_metadata(tmp_path, monkeypatch):
    path = tmp_path / 'cached.pptx'
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    calls = []
    original = templates.extract_examples

    def extract(*args):
        calls.append(True)
        return original(*args)

    monkeypatch.setattr(templates, 'extract_examples', extract)
    first = templates.analyze_template(path)
    first['colors']['accent1'] = '#ABCDEF'
    first['assets'].append({'id': 'caller-only'})
    second = templates.analyze_template(path)

    assert len(calls) == 1
    assert second['colors']['accent1'] != '#ABCDEF'
    assert second['assets'] == []

    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    assert templates.analyze_template(path)['count'] == 2
    assert len(calls) == 2
