"""Acceptance checks mapped to the supplied competition requirements."""
import asyncio
import base64
from io import BytesIO
from time import monotonic
from zipfile import ZipFile
from uuid import uuid4
import httpx
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient
import pytest
from app import database as db
from app.config import settings
from app.main import app
from app.models import DeckContent, Slide, ChartData
from app.services import workflow, images
from app.services.audit import audit_deck
from app.services.export import export_pptx, export_pdf, export_html
from app.services.layout import build_scene

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings,'database',tmp_path/'requirements.sqlite3')
    monkeypatch.setattr(settings,'storage',tmp_path/'data')
    monkeypatch.setattr(settings,'mode','demo')
    monkeypatch.setattr(settings,'email_verification_required',False)
    monkeypatch.setattr(settings,'image_base_url','')
    with TestClient(app) as c:
        r=c.post('/api/auth/register',json={'email':'requirements@example.test','password':'Requirements-2026!','first_name':'Test','last_name':'Requirements'})
        c.headers['X-CSRF-Token']=r.json()['csrf']
        yield c

def png():
    image=BytesIO();Image.new('RGB',(64,64),'#0077ff').save(image,format='PNG')
    return image.getvalue()

def acceptance_deck():
    slides=[Slide(title='Результаты пилота',body='Синтетический материал для проверки',kind='title')]
    for kind in ('bar','column','line'):
        slides.append(Slide(title='Динамика показателей',kind='chart',chart=ChartData(labels=['До','После'],values=[8,12],unit='шт.',chart_type=kind)))
    slides.append(Slide(title='Сравнение этапов',kind='table',table=[['Этап','Результат'],['Пилот','12'],['Запуск','18']]))
    for kind in ('process','cycle','hierarchy'):
        slides.append(Slide(title='Этапы работы',kind='diagram',diagram_type=kind,bullets=['Задача','Решение','Результат']))
    slides.append(Slide(title='Принципы команды',kind='icons',bullets=['Понятная цель','Одна команда','Надёжность'],icon_names=['target','people','shield']))
    slides.append(Slide(title='Иллюстрация',body='Изображение — один объект на редактируемом слайде.',kind='image',image_data='data:image/png;base64,'+base64.b64encode(png()).decode()))
    return DeckContent(title='Приёмочная проверка',slides=slides)

def unknown_template(ratio):
    prs=Presentation()
    prs.slide_width=Inches(10);prs.slide_height=Inches(10/ratio)
    slide=prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text='Unseen template '+str(uuid4())
    slide.shapes.title.text_frame.paragraphs[0].font.name='Arial'
    slide.shapes.title.text_frame.paragraphs[0].font.size=Pt(36)
    slide.placeholders[1].text='Content'
    out=BytesIO();prs.save(out)
    return out.getvalue()

@pytest.mark.parametrize('ratio',[16/9,4/3,3/4])
def test_unknown_templates_native_exports_and_ten_slides(client,ratio):
    started=monotonic()
    uploaded=client.post('/api/templates',files={'file':('unseen-'+str(uuid4())+'.pptx',unknown_template(ratio),'application/octet-stream')})
    assert uploaded.status_code==201,uploaded.text
    t=uploaded.json()
    assert t['metadata']['ratio']==pytest.approx(ratio,rel=1e-5)
    assert t['metadata']['normalization_version']=='3'
    assert t['metadata']['layouts'] and t['metadata']['colors']
    deck=acceptance_deck()
    response=client.post('/api/generate',json={'template_id':t['id'],'content':deck.model_dump(),'source_text':'Синтетический материал'})
    job=client.get('/api/jobs/'+response.json()['job_id']).json()
    assert job['state']=='complete',job
    p=client.get('/api/projects/'+job['project_id']).json()
    assert p['provenance']['within_budget'] and p['provenance']['slides']==10
    assert p['provenance']['recipe']['snippets'][0]['version']
    assert p['variants']['a']!=p['variants']['b']!=p['variants']['c']
    for variant in ('a','b','c'):
        for ext in ('pptx','pdf','html'):
            r=client.get(f'/api/projects/{p["id"]}/export/{ext}?variant={variant}')
            assert r.status_code==200 and len(r.content)>1000
            if ext=='pptx':
                result=Presentation(BytesIO(r.content))
                assert len(result.slides)==10
                assert result.slide_width/result.slide_height==pytest.approx(ratio,rel=1e-5)
                assert result.slides[1].shapes.has_chart if hasattr(result.slides[1].shapes,'has_chart') else any(s.has_chart for s in result.slides[1].shapes)
                assert any(s.has_table for s in result.slides[4].shapes)
                assert any(s.shape_type==MSO_SHAPE_TYPE.LINE for s in result.slides[5].shapes)
                assert any(s.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE for s in result.slides[8].shapes)
                assert any(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in result.slides[9].shapes)
                assert all(any(s.has_text_frame for s in slide.shapes) for slide in result.slides)
        issues=client.get(f'/api/projects/{p["id"]}/audit?variant={variant}').json()['issues']
        assert all(i['deterministic'] for i in issues)
    assert monotonic()-started<300

def test_packet_normalizes_multiple_formats_and_keeps_sources(client):
    doc=BytesIO()
    with ZipFile(doc,'w') as z:
        z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Текст из Word</w:t></w:r></w:p></w:body></w:document>')
    pdf=BytesIO();c=canvas.Canvas(pdf);c.drawString(40,700,'PDF source');c.save()
    response=client.post('/api/materials/import',files=[
        ('files',('brief.txt','  Бриф  \r\nЦель '.encode(),'text/plain')),
        ('files',('data.csv','Период;Значение\nПилот;12'.encode('cp1251'),'text/csv')),
        ('files',('notes.docx',doc.getvalue(),'application/octet-stream')),
        ('files',('report.pdf',pdf.getvalue(),'application/pdf')),
    ])
    assert response.status_code==200,response.text
    packet=response.json()
    assert len(packet['documents'])==4
    assert 'Текст из Word' in packet['text'] and 'PDF source' in packet['text']
    assert '\r' not in packet['text'] and packet['documents'][1]['warnings']
    assert all(len(d['sha256'])==64 for d in packet['documents'])
    assert client.post('/api/materials/import',files={'files':('secret.exe',b'fake','text/plain')}).status_code==422
    assert client.post('/api/materials/import',files={'files':('huge.txt',b'x'*50001,'text/plain')}).status_code==413
    assert client.post('/api/materials/import',files={'files':('t.txt',b'text','text/plain')},headers={'X-CSRF-Token':'bad'}).status_code==403

def test_deadline_rejects_partial_result_and_keeps_retry(client,monkeypatch):
    monkeypatch.setattr(workflow,'BUDGET_SECONDS',.04)
    async def slow(*args):
        await asyncio.sleep(2)
    monkeypatch.setattr('app.main.create_design_variants',slow)
    response=client.post('/api/generate',json={'template_id':'tech','content':{'title':'Deadline','slides':[{'title':'T'}]}})
    job=client.get('/api/jobs/'+response.json()['job_id']).json()
    assert job['state']=='failed' and job['can_retry'] and not job['project_id']
    assert '5 минут' in job['error']
    assert client.get('/api/projects').json()==[]

def test_image_provider_is_optional_and_returns_only_bounded_embedded_pixels(client,monkeypatch):
    assert client.post('/api/images/generate',json={'prompt':'Illustration'}).status_code==409
    monkeypatch.setattr(settings,'image_base_url','http://127.0.0.1:8002')
    def handler(request):
        assert request.url.path=='/generate'
        return httpx.Response(200,content=png(),headers={'content-type':'image/png'})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as transport:
            result=await images.generate_image('Example',transport)
            assert result['parameters_billion']==12
            assert result['image_data'].startswith('data:image/png;base64,')
    asyncio.run(run())

def test_recipe_is_versioned_and_manifest_is_immutable(client,monkeypatch):
    original=workflow.recipe()
    assert len(original['snippets'])==4 and len(original['stages'])==8
    assert all(len(s['version'])==12 for s in original['snippets'])
    r=client.post('/api/generate',json={'template_id':'tech','content':acceptance_deck().model_dump()})
    job=client.get('/api/jobs/'+r.json()['job_id']).json()
    p=client.get('/api/projects/'+job['project_id']).json()
    before=p['provenance']
    monkeypatch.setattr(workflow,'RECIPE_VERSION','future-version')
    assert workflow.recipe()['version']=='future-version'
    assert client.get('/api/projects/'+p['id']).json()['provenance']==before


def test_docx_table_keeps_cells_rows_and_inline_breaks(client):
    document = BytesIO()
    with ZipFile(document, 'w') as archive:
        archive.writestr('word/document.xml', '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body><w:p><w:r><w:t>Бриф</w:t><w:br/><w:t>Пилот</w:t></w:r></w:p>
            <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Период</w:t></w:r></w:p></w:tc>
              <w:tc><w:p><w:r><w:t>Минуты</w:t></w:r></w:p></w:tc></w:tr>
              <w:tr><w:tc><w:p><w:r><w:t>После</w:t></w:r></w:p></w:tc>
              <w:tc><w:p><w:r><w:t>8</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
          </w:body></w:document>''')
    response = client.post('/api/materials/import', files={
        'files': ('brief.docx', document.getvalue(), 'application/octet-stream')})
    assert response.status_code == 200
    assert response.json()['text'] == 'Источник: brief.docx\nБриф\nПилот\nПериод | Минуты\nПосле | 8'


def test_audit_detects_nonadjacent_overlap_once(monkeypatch):
    def block(name, x, y):
        return {'id': name, 'type': 'text', 'x': x, 'y': y, 'w': 100, 'h': 80,
                'text': name, 'used_height': 30, 'font_size': 24, 'color': '#000000'}
    scene = {'width': 1280, 'height': 720, 'objects': [
        {'id': 'background', 'type': 'rect', 'x': 0, 'y': 0, 'w': 1280, 'h': 720, 'fill': '#FFFFFF'},
        block('title', 20, 20), block('distant', 500, 400), block('body', 30, 30),
        block('other', 35, 35)]}
    monkeypatch.setattr('app.services.audit.build_scene', lambda *args: scene)
    issues = audit_deck(DeckContent(title='Test', slides=[Slide(title='Test')]), {'metadata': {}}, 'a', '')
    overlaps = [issue for issue in issues if issue.code == 'overlap']
    assert {issue.object_id for issue in overlaps} == {'body', 'other'}
    assert len({issue.id for issue in overlaps}) == len(overlaps) == 2


def test_deadline_includes_waiting_for_generation_slot(client, monkeypatch):
    monkeypatch.setattr(workflow, 'ASSEMBLY_BUDGET_SECONDS', .02)
    class OccupiedSlots:
        async def __aenter__(self):
            await asyncio.sleep(1)
        async def __aexit__(self, *args):
            pass
    monkeypatch.setattr('app.main.generation_slots', OccupiedSlots())
    response = client.post('/api/generate', json={
        'template_id': 'tech', 'content': {'title': 'Queue deadline', 'slides': [{'title': 'Test'}]}})
    job = client.get('/api/jobs/' + response.json()['job_id']).json()
    assert job['state'] == 'failed' and job['can_retry'] and not job['project_id']
    assert client.get('/api/projects').json() == []
