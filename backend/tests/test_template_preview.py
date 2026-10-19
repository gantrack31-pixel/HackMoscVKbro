from io import BytesIO
from zipfile import ZipFile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest
from app.config import settings
from app.services import template_preview as preview

def fixture_pptx(path,color):
    prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(path)
    image=BytesIO();Image.new('RGB',(320,180),color).save(image,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('docProps/thumbnail.png',image.getvalue())

def test_embedded_real_preview_is_cached_and_content_hash_invalidates(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'storage',tmp_path/'storage')
    monkeypatch.setattr(preview.shutil,'which',lambda _:None)
    path=tmp_path/'template.pptx';Presentation().save(path)
    image=BytesIO();Image.new('RGB',(1200,900),'red').save(image,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('docProps/thumbnail.png',image.getvalue())
    first=preview.thumbnail(path)
    assert Image.open(first).getpixel((20,20))==(255,0,0)
    original=preview.render_first_slide
    def unexpected(_): raise AssertionError('Cache must avoid reconversion')
    monkeypatch.setattr(preview,'render_first_slide',unexpected)
    assert preview.thumbnail(path)==first
    monkeypatch.setattr(preview,'render_first_slide',original)
    Presentation().save(path)
    image=BytesIO();Image.new('RGB',(1200,900),'blue').save(image,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('docProps/thumbnail.png',image.getvalue())
    second=preview.thumbnail(path)
    assert second!=first and Image.open(second).getpixel((20,20))==(0,0,255)

def test_preview_cache_version_tracks_higher_resolution_assets():
    assert preview.PREVIEW_VERSION=='first-slide-4-faithful'

def test_clean_preview_keeps_large_slides_sharp_and_bounds_oversized_images():
    large=BytesIO();Image.new('RGB',(2560,1440),'navy').save(large,format='PNG')
    result=Image.open(BytesIO(preview.clean_png(large.getvalue())))
    assert result.size==(2560,1440)
    tiny=BytesIO();Image.new('RGB',(320,180),'navy').save(tiny,format='PNG')
    result=Image.open(BytesIO(preview.clean_png(tiny.getvalue())))
    assert result.size==(320,180)

def test_unavailable_preview_is_explicit_and_briefly_negative_cached(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'storage',tmp_path/'storage')
    path=tmp_path/'template.pptx';Presentation().save(path)
    calls=[]
    monkeypatch.setattr(preview,'render_first_slide',lambda p:calls.append(p))
    assert preview.thumbnail(path) is None and preview.thumbnail(path) is None
    assert len(calls)==1 and preview.thumbnail(tmp_path/'missing.pptx') is None

def test_low_resolution_embedded_thumbnail_is_not_upscaled(tmp_path,monkeypatch):
    monkeypatch.setattr(preview.shutil,'which',lambda _:None)
    path=tmp_path/'low-resolution.pptx';fixture_pptx(path,'purple')
    assert preview.render_first_slide(path) is None

def test_high_resolution_full_slide_image_is_used_when_renderer_is_missing(tmp_path,monkeypatch):
    monkeypatch.setattr(preview.shutil,'which',lambda _:None)
    path=tmp_path/'full-slide-image.pptx'
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    image_path=tmp_path/'full-slide.png';Image.new('RGB',(1920,1080),'navy').save(image_path,format='PNG')
    slide.shapes.add_picture(str(image_path),0,0,width=prs.slide_width,height=prs.slide_height)
    prs.save(path)
    result=preview.render_first_slide(path)
    assert result is not None
    assert Image.open(BytesIO(result)).size==(1920,1080)

def test_unreferenced_package_media_is_never_presented_as_the_slide(tmp_path,monkeypatch):
    monkeypatch.setattr(preview.shutil,'which',lambda _:None)
    path=tmp_path/'slide-background.pptx'
    prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(path)
    image=BytesIO();Image.new('RGB',(2048,1152),'darkblue').save(image,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('ppt/media/full-slide-background.png',image.getvalue())
    result=preview.render_first_slide(path)
    assert result is None


def picture_fixture(path, image_path, *, overlay=False):
    prs=Presentation();prs.slide_width=Inches(12);prs.slide_height=Inches(6.75)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    picture=slide.shapes.add_picture(str(image_path),0,0,width=prs.slide_width,height=prs.slide_height)
    if overlay:
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text='Заголовок поверх фона'
    return prs,picture


@pytest.mark.parametrize('change',['overlay','crop','rotate','partial','transparent','border','rounded'])
def test_incomplete_picture_fallback_never_claims_to_render_full_slide(tmp_path,monkeypatch,change):
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    monkeypatch.setattr(preview,'office_preview',lambda _:None)
    path=tmp_path/'overlay.pptx';image_path=tmp_path/'image.png'
    Image.new('RGBA',(1920,1080),(30,60,90,128 if change=='transparent' else 255)).save(image_path)
    prs,picture=picture_fixture(path,image_path,overlay=change=='overlay')
    if change=='crop': picture.crop_left=.1
    if change=='rotate': picture.rotation=15
    if change=='partial': picture.left=Inches(.5)
    if change=='border': picture.line.width=Inches(.03)
    if change=='rounded': picture.auto_shape_type=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
    prs.save(path)
    assert preview.render_first_slide(path) is None


def test_stretched_raster_fallback_matches_slide_ratio_without_upscaling(tmp_path,monkeypatch):
    monkeypatch.setattr(preview,'office_preview',lambda _:None)
    path=tmp_path/'stretch.pptx';image_path=tmp_path/'image.png'
    Image.new('RGB',(1920,1080),'navy').save(image_path)
    prs,picture=picture_fixture(path,image_path)
    prs.slide_width=Inches(10);prs.slide_height=Inches(7.5)
    picture.width=prs.slide_width;picture.height=prs.slide_height
    prs.save(path)
    result=Image.open(BytesIO(preview.render_first_slide(path)))
    assert result.size==(1440,1080)


def test_renderer_failure_uses_faithful_embedded_cover(tmp_path,monkeypatch):
    def unavailable(_): raise OSError('Renderer unavailable')
    monkeypatch.setattr(preview,'office_preview',unavailable)
    path=tmp_path/'embedded.pptx';image_path=tmp_path/'image.png'
    Image.new('RGB',(1920,1080),'navy').save(image_path)
    prs,_=picture_fixture(path,image_path,overlay=True);prs.save(path)
    thumbnail=BytesIO();Image.new('RGB',(1600,900),'orange').save(thumbnail,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('docProps/thumbnail.png',thumbnail.getvalue())
    result=Image.open(BytesIO(preview.render_first_slide(path)))
    assert result.size==(1600,900) and result.getpixel((1,1))==(255,165,0)


def test_wrong_ratio_embedded_thumbnail_is_not_a_slide_cover(tmp_path,monkeypatch):
    monkeypatch.setattr(preview,'office_preview',lambda _:None)
    path=tmp_path/'square-thumbnail.pptx';Presentation().save(path)
    image=BytesIO();Image.new('RGB',(1200,1200),'blue').save(image,format='PNG')
    with ZipFile(path,'a') as archive: archive.writestr('docProps/thumbnail.png',image.getvalue())
    assert preview.render_first_slide(path) is None

def test_libreoffice_conversion_uses_first_page_isolated_profile_and_no_external_links(tmp_path,monkeypatch):
    monkeypatch.setattr(preview.shutil,'which',lambda cmd:cmd)
    path=tmp_path/'template.pptx';Presentation().save(path)
    with ZipFile(path,'a') as archive:
        archive.writestr('ppt/external.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" TargetMode="External" Target="https://example.test/image" Type="image"/></Relationships>')
    calls=[]
    def run(args,**kwargs):
        calls.append(args)
        assert kwargs['timeout']<=45 and kwargs['check']
        if args[0]=='libreoffice':
            assert any('PageRange' in a for a in args) and any(a.startswith('-env:UserInstallation=') for a in args)
            with ZipFile(args[-1]) as safe: assert b'example.test' not in safe.read('ppt/external.rels')
        else:
            assert args[1:4]==['-f','1','-singlefile']
            Image.new('RGB',(320,180),'green').save(args[-1]+'.png')
    monkeypatch.setattr(preview.subprocess,'run',run)
    data=preview.render_first_slide(path)
    assert Image.open(BytesIO(data)).size==(320,180) and len(calls)==2
