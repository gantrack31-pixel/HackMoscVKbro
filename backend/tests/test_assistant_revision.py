"""Regression tests for the preview/edit/export audit. No live model or credentials."""
import asyncio
import base64
from copy import deepcopy
from io import BytesIO
import json
import pytest
from PIL import Image
from pptx import Presentation
from app import database as db
from app.models import AssistantPlan, AssistantRequest, DeckContent, Slide, ChartData
from app.services import llm
from app.services.assistant import apply_patch, edit_presentation, enforce_constraints
from app.services.assistant_context import presentation_context, relevant_source
from app.services.scene_snapshot import scene_image
from test_assistant import client, request_for  # noqa: F401 - imported pytest fixture
from test_analytics_integration import upload_choice

TEMPLATE = {'name': 'Test', 'metadata': {'ratio': 16/9, 'assets': []}}


def reply(monkeypatch, changes, **extra):
    async def answer(*args, **kwargs):
        return {'summary': 'Предложение для просмотра.', 'changes': changes, **extra}
    monkeypatch.setattr(llm, 'complete_with_template', answer)
    monkeypatch.setattr(llm, 'complete_json', answer)


def run(slide, **kwargs):
    request = AssistantRequest(content=DeckContent(title='Проверка', slides=[slide]), base_updated_at='v1',
                               instruction='Исправь', slide=0, **kwargs)
    return asyncio.run(edit_presentation(request, TEMPLATE, ''))


def test_null_patch_really_clears_nullable_field_and_kind_cleans_visuals():
    slide = Slide(title='A', kind='chart', chart=ChartData(labels=['A'], values=[30]))
    patch = AssistantPlan.model_validate({'summary': 'Очистить', 'changes': [{'slide': 0, 'kind': 'text', 'chart': None}]}).changes[0]
    updated = apply_patch(slide, patch, 'a')
    assert updated.kind == 'text' and updated.chart is None
    with pytest.raises(ValueError):
        apply_patch(slide, patch.model_copy(update={'kind': 'chart'}), 'a')


def test_empty_repair_never_claims_to_fix_overflow(monkeypatch):
    reply(monkeypatch, [])
    content, report = run(Slide(title='Заголовок', body='Очень длинный подробный текст. '*70), action='repair')
    assert report['changed_slides'] == []
    assert content.slides[0].fixed == []
    assert any(i['code'] == 'overflow' and not i['fixable'] for i in report['issues'])


def test_ordinary_edit_rejects_new_overflow_and_explains_reason(monkeypatch):
    reply(monkeypatch, [{'slide': 0, 'body': 'Очень длинный подробный текст. '*70}])
    with pytest.raises(llm.LLMError, match='новые ошибки'):
        run(Slide(title='Заголовок', body='Коротко.'))


def test_number_replacement_requires_explicit_permission_and_evidence(monkeypatch):
    correction = {'before': 'Участников: 30.', 'after': 'Участников: 35.', 'evidence': 'Участников: 35.'}
    reply(monkeypatch, [{'slide': 0, 'body': 'Участников: 35.'}], fact_corrections=[correction])
    with pytest.raises(llm.LLMError, match='разрешите'):
        run(Slide(title='Пилот', body='Участников: 30.'), material='Участников: 35.')
    result, report = run(Slide(title='Пилот', body='Участников: 30.'), material='Участников: 35.', allow_fact_corrections=True)
    assert result.slides[0].body == 'Участников: 35.' and not result.slides[0].notes
    assert report['fact_corrections'] == [correction]
    assert result.evidence_sources[0].text == 'Участников: 35.'


def test_fabricated_fact_correction_evidence_is_rejected(monkeypatch):
    reply(monkeypatch, [{'slide': 0, 'body': 'Участников: 35.'}], fact_corrections=[
        {'before': 'Участников: 30.', 'after': 'Участников: 35.', 'evidence': 'Участников: 35.'}])
    with pytest.raises(llm.LLMError, match='подтверждающего'):
        run(Slide(title='Пилот', body='Участников: 30.'), allow_fact_corrections=True, material='Участников: 40.')


def test_swapped_label_values_rejected_even_when_number_set_unchanged():
    before = Slide(title='Метрики', body='Выручка: 30 млн. Затраты: 10 млн.')
    after = before.model_copy(update={'body': 'Выручка: 10 млн. Затраты: 30 млн.'})
    request = AssistantRequest(content=DeckContent(title='Метрики', slides=[before]), base_updated_at='v1')
    with pytest.raises(ValueError, match='показателю'):
        enforce_constraints(before, after, request, 0, '')


def test_supplementary_quote_survives_apply_audit_and_undo(client, monkeypatch):
    body = request_for(client, material='Пилот завершён успешно.')
    reply(monkeypatch, [{'slide': 0, 'title': 'Пилот', 'body': 'Пилот завершён успешно.', 'source_quote': 'Пилот завершён успешно.'}])
    response = client.post('/api/projects/ai-test/assistant', json=body)
    assert response.status_code == 200, response.text
    proposal = response.json()['proposal']
    assert proposal and client.get('/api/projects/ai-test').json()['content'] == body['content']
    applied = client.post(f'/api/projects/ai-test/assistant/proposals/{proposal["id"]}/apply', json=body)
    assert applied.status_code == 200, applied.text
    project = applied.json()['project']
    assert project['content']['evidence_sources'][0]['text'] == body['material']
    assert project['sources'][0]['status'] == 'matched'
    audited = client.get('/api/projects/ai-test/audit?variant=a').json()
    assert not any(i['code'] == 'source_not_found' and i['slide'] == 0 for i in audited['issues'])
    undo = client.post('/api/projects/ai-test/undo', json={'base_updated_at': project['updated_at']}).json()
    assert not undo['content'].get('evidence_sources')


def test_compact_context_removes_binary_and_keeps_absolute_indices():
    content = DeckContent(title='Deck', slides=[Slide(title='Other', body='Long '*100), Slide(title='Edited', body='Current')])
    context = presentation_context(content, [1])
    assert context['slides'][0] == {'slide': 0, 'title': 'Other', 'kind': 'text', 'context_only': True}
    assert context['slides'][1]['body'] == 'Current'
    assert 'image_data' not in context['slides'][1] and 'resource_icon_data' not in context['slides'][1]
    source = ('Не относящиеся сведения. '*90 + '\n')*20 + 'Ключевойметрик: подтверждённый результат.'
    subset = relevant_source(source, 'Ключевойметрик', limit=3200)
    assert len(subset) <= 3200 and 'Ключевойметрик' in subset


def test_geometry_contains_coordinates_and_audits_only_edited_slide(monkeypatch):
    seen = []
    async def answer(system, payload, template):
        seen.append(payload)
        return {'summary': 'Без изменений.', 'changes': []}
    monkeypatch.setattr(llm, 'complete_with_template', answer)
    request = AssistantRequest(content=DeckContent(title='Test', slides=[Slide(title='First'), Slide(title='Second')]),
                               base_updated_at='v1', slide=1, instruction='Объясни')
    _, report = asyncio.run(edit_presentation(request, TEMPLATE, ''))
    assert report['audit_scope'] == [1]
    assert all(i['slide'] == 1 for i in report['issues'])
    assert seen[0]['geometry'][0]['slide'] == 1
    assert {'x', 'y', 'w', 'h'} <= seen[0]['geometry'][0]['objects'][0].keys()


def test_vision_uses_current_slide_not_template_assets(monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, 'vision', True)
    async def answer(system, payload, **kwargs):
        assert len(kwargs['images']) == 1
        with Image.open(BytesIO(base64.b64decode(kwargs['images'][0].split(',', 1)[1]))) as image:
            assert image.size == (1280, 720)
        return {'summary': 'Без изменений.', 'changes': []}
    monkeypatch.setattr(llm, 'complete_json', answer)
    run(Slide(title='Текущий слайд', body='Текст, не картинка шаблона.'))


def test_snapshot_is_valid_png():
    from app.services.layout import build_scene
    scene = build_scene(Slide(title='Проверка', body='Русский текст'), TEMPLATE['metadata'], 'a', 0)
    with Image.open(BytesIO(base64.b64decode(scene_image(scene).split(',', 1)[1]))) as image:
        image.verify()


def test_changed_analytics_cannot_be_saved_or_exported(client, monkeypatch):
    _, slides = upload_choice(client, 'Дата,Revenue\n2025-01-01,100\n2025-12-01,123',
                              {'chart_type': 'line', 'dimension': 'Дата', 'measures': ['Revenue']})
    analytical = Slide.model_validate(slides[0])
    request = AssistantRequest(content=DeckContent(title='Analytics', slides=[analytical]), base_updated_at='v1')
    with pytest.raises(ValueError, match='проверенный вывод'):
        enforce_constraints(analytical, analytical.model_copy(update={'title': 'Непроверенный вывод'}), request, 0, '')
    owner = db.project_get('ai-test')['user_id']
    db.project_create('analytics-review', 'tech', DeckContent(title='Analytics', slides=slides).model_dump(), '', owner)
    project = client.get('/api/projects/analytics-review').json()
    changed = deepcopy(project['content'])
    changed['slides'][0]['analytics']['series'][0]['values'][1] = 999
    response = client.put('/api/projects/analytics-review', json={'content': changed, 'variant': 'a', 'base_updated_at': project['updated_at']})
    assert response.status_code == 409 and response.json()['detail']['code'] == 'ANALYTICS_CHANGED'
    assert client.get('/api/projects/analytics-review').json()['content'] == project['content']
    # Even historical corrupted DB snapshots are not silently exported as verified.
    with db.connection() as conn:
        conn.execute('UPDATE projects SET content=? WHERE id=?', (json.dumps(changed), 'analytics-review'))
    assert client.get('/api/projects/analytics-review/export/pptx?variant=a').status_code == 409
    from app import main
    monkeypatch.setattr(main.yandex_disk, 'token_for', lambda user_id: 'test-only')
    assert client.post('/api/projects/analytics-review/disk').status_code == 409


def test_caller_cannot_forge_fingerprint_to_authorize_new_evidence(client):
    from app.services.analytics_plan import visual_fingerprint
    _, slides = upload_choice(client, 'Region,Revenue\nNorth,100\nSouth,123',
                              {'chart_type': 'bar', 'dimension': 'Region', 'measures': ['Revenue']})
    project = client.get('/api/projects/ai-test').json()
    changed = deepcopy(project['content'])
    visual = slides[0]['analytics']
    visual['series'][0]['values'][0] = 999
    visual['metadata']['verified_visual_fingerprint'] = visual_fingerprint(visual)
    changed['slides'].append(slides[0])
    response = client.put('/api/projects/ai-test', json={'content': changed, 'variant': 'a', 'base_updated_at': project['updated_at']})
    assert response.status_code == 409


def test_portable_font_is_explicit_and_consistent():
    from app.services.layout import build_scene
    from app.services.export import export_pptx, export_html
    template = {'name': 'Corporate', 'metadata': {'ratio': 16/9, 'heading_font': 'Unavailable Corporate'}}
    content = DeckContent(title='Deck', slides=[Slide(title='Заголовок', body='Текст')])
    scene = build_scene(content.slides[0], template['metadata'], 'a', 0)
    title = next(o for o in scene['objects'] if o['id'] == 'title')
    assert title['font_family'] == 'Unavailable Corporate'
    assert title['effective_font_family'] == 'Manrope' and scene['font_warnings']
    deck = Presentation(BytesIO(export_pptx(content, template, 'a')))
    assert all(p.font.name == 'Manrope' for shape in deck.slides[0].shapes if shape.has_text_frame
               for p in shape.text_frame.paragraphs if p.text)
    assert 'font-family="Manrope"' in export_html(content, template, 'a').decode()
