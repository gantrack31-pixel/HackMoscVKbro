"""Regression tests for portable brand fonts and SVG logos."""
from base64 import b64decode
from io import BytesIO
import json
import struct
from pathlib import Path
from zipfile import ZipFile
import pytest
import reportlab
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from app.models import DeckContent, Slide
from app.services.brand_parser import parse_brand, BrandError
from app.services.brand_style import resolve_brand, parse_isolated
from app.services.export import export_html, export_pdf, export_pptx
from app.services.layout import build_scene
from test_brand_style import package
from test_analytical_rendering import analytical_slide
from test_workflow import client, register

SVG = b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 400 100"><rect width="400" height="100" fill="#008877"/></svg>'
FONT_DIR = Path(reportlab.__file__).parent / 'fonts'
FAMILY = 'Bitstream Vera Sans'


def editable_font(filename):
    # Synthetic editable fixture; the unchanged ReportLab font is tested as restricted below.
    data = bytearray((FONT_DIR / filename).read_bytes())
    count = struct.unpack_from('>H', data, 4)[0]
    head = None
    def checksum(value):
        padded = value + bytes((-len(value)) % 4)
        return sum(struct.unpack('>'+str(len(padded)//4)+'I', padded)) & 0xFFFFFFFF
    for i in range(count):
        record = 12+16*i
        tag, _, offset, size = struct.unpack_from('>4sIII', data, record)
        if tag == b'OS/2':
            struct.pack_into('>H', data, offset+8, 0)
            struct.pack_into('>I', data, record+4, checksum(bytes(data[offset:offset+size])))
        if tag == b'head': head = offset
    struct.pack_into('>I', data, head+8, 0)
    struct.pack_into('>I', data, head+8, (0xB1B0AFBA-checksum(bytes(data))) & 0xFFFFFFFF)
    return bytes(data)


def branded_template():
    data = package([('style.json', json.dumps({'typography': {'heading': FAMILY, 'body': FAMILY}}).encode()),
                    ('Regular.ttf', editable_font('Vera.ttf')),
                    ('Bold.ttf', editable_font('VeraBd.ttf')), ('logo.svg', SVG)])
    brand = parse_brand(data, 'brand.zip'); brand['id'] = 'brand-assets'
    return resolve_brand({'metadata': {'ratio': 16 / 9, 'accent': '#008877', 'background': '#FFFFFF'}}, brand)


def test_svg_logo_is_parsed_in_worker_with_colors_and_aspect():
    logo = parse_isolated(SVG, 'logo.svg')['logos'][0]
    image = Image.open(BytesIO(b64decode(logo['src'].split(',')[1])))
    assert image.width / image.height == pytest.approx(4)
    assert image.convert('RGBA').getpixel((20, 20)) == (0, 136, 119, 255)


def test_svg_local_gradient_and_fractional_aspect_are_preserved():
    svg=b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100.5 30"><defs><linearGradient id="ink"><stop stop-color="#008877"/><stop offset="1" stop-color="#008877"/></linearGradient></defs><rect width="100.5" height="30" fill="url(#ink)"/></svg>'
    logo=parse_brand(svg,'logo.svg')['logos'][0]
    assert logo['width']/logo['height']==pytest.approx(100.5/30)
    assert '#008877' in logo['visible_colors']


def test_svg_live_text_is_visible_and_has_an_outline_warning():
    brand=parse_brand(b'<svg xmlns="http://www.w3.org/2000/svg" width="160" height="40"><text x="5" y="30" font-size="24">North</text></svg>','logo.svg')
    image=Image.open(BytesIO(b64decode(brand['logos'][0]['src'].split(',')[1])))
    assert image.getbbox() is not None
    assert any('в кривые' in warning for warning in brand['warnings'])


def test_project_manifest_preserves_actual_font_files_for_reload_and_export():
    from app.services.workflow import manifest
    template={**branded_template(),'id':'tech','name':'Tech'}
    result=manifest(template,DeckContent(title='Brand',slides=[Slide(title='Brand title')]),'',{},1,'job')
    assert result['effective_style']['metadata']['font_assets']==template['metadata']['font_assets']


@pytest.mark.parametrize('fragment', [
    '<script>alert(1)</script>', '<foreignObject><p>x</p></foreignObject>',
    '<image href="https://example.com/private.png"/>', '<image href="file:///etc/passwd"/>',
    '<rect onclick="alert(1)"/>', '<style>@import "https://example.com/style.css";</style>',
    '<rect fill="url(https://example.com/fill.svg)"/>',
])
def test_svg_rejects_active_content_and_external_resources(fragment):
    data = f'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">{fragment}</svg>'.encode()
    with pytest.raises(BrandError): parse_brand(data, 'logo.svg')


@pytest.mark.parametrize('data', [b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg>&x;</svg>',
    b'<svg xmlns="http://www.w3.org/2000/svg" width="10000000" height="10000000"/>'])
def test_svg_rejects_entities_and_unbounded_geometry(data):
    with pytest.raises(BrandError): parse_brand(data, 'logo.svg')


@pytest.mark.parametrize('slide', [Slide(title='Brand title', body='Measured body text'), analytical_slide('grouped')], ids=['ordinary', 'analytical'])
def test_uploaded_font_is_used_in_scene_pdf_html_and_embedded_powerpoint(slide):
    slide=slide.model_copy(update={'title':'Brand title'})
    template = branded_template()
    assert template['metadata']['heading_font'] == FAMILY
    scene = build_scene(slide, template['metadata'], 'a', 0)
    assert next(node['effective_font_family'] for node in scene['objects'] if node['id']=='title') == FAMILY
    assert scene['font_faces']
    deck = DeckContent(title='Brand', slides=[slide])
    pdf = PdfReader(BytesIO(export_pdf(deck, template, 'a')))
    assert any('DecklyFont' in str(font.get_object().get('/BaseFont')) for font in pdf.pages[0]['/Resources']['/Font'].values())
    html = export_html(deck, template, 'a').decode()
    assert 'data:font/ttf;base64,' in html and scene['font_faces'][0]['css_family'] in html
    output = export_pptx(deck, template, 'a')
    prs = Presentation(BytesIO(output))
    assert any(p.font.name == FAMILY for shape in prs.slides[0].shapes if shape.has_text_frame for p in shape.text_frame.paragraphs)
    with ZipFile(BytesIO(output)) as archive:
        assert FAMILY.encode() in archive.read('ppt/presentation.xml')
        assert any(editable_font('Vera.ttf') in archive.read(name) for name in archive.namelist() if name.endswith('.fntdata'))


def test_font_change_changes_real_line_measurements_without_changing_text():
    slide = Slide(title='Widest word', body='Wide words ' * 40)
    branded = build_scene(slide, branded_template()['metadata'], 'a', 0)
    plain = build_scene(slide, {'ratio': 16 / 9}, 'a', 0)
    body = next(node for node in branded['objects'] if node['id'] == 'body')
    original = next(node for node in plain['objects'] if node['id'] == 'body')
    assert body['text'] == original['text'] == slide.body
    assert body['lines'] != original['lines']


def test_corrupt_font_is_rejected():
    with pytest.raises(BrandError): parse_brand(b'not a font', 'broken.ttf')


def test_restricted_font_is_rejected():
    with pytest.raises(BrandError,match='Лицензионные'):
        parse_brand((FONT_DIR / 'Vera.ttf').read_bytes(), 'restricted.ttf')


def test_missing_glyphs_fall_back_consistently_instead_of_losing_cyrillic_text():
    slide=Slide(title='Русский заголовок',body='Текст должен оставаться читаемым')
    scene=build_scene(slide,branded_template()['metadata'],'a',0)
    title=next(node for node in scene['objects'] if node['id']=='title')
    assert title['effective_font_family']=='Manrope'
    assert title['text']==slide.title and scene['font_warnings']


def test_standalone_font_upload_and_brand_preview_use_its_actual_family(client):
    register(client)
    uploaded=client.post('/api/brands/upload',files={'file':('Regular.ttf',editable_font('Vera.ttf'),'font/ttf')})
    assert uploaded.status_code==201,uploaded.text
    preview=client.post(f'/api/brands/{uploaded.json()["id"]}/preview',json={'template_id':'tech'})
    assert preview.status_code==200,preview.text
    assert preview.json()['effective']['heading_font']==FAMILY
    assert not preview.json()['brand_style']['font_mapping'][0]['substituted']


def test_svg_upload_persists_original_and_generates_a_logo_preview(client):
    register(client)
    uploaded=client.post('/api/brands/upload',files={'file':('logo.svg',SVG,'image/svg+xml')})
    assert uploaded.status_code==201,uploaded.text
    brand_id=uploaded.json()['id']
    assert client.get(f'/api/brands/{brand_id}/source').content==SVG
    preview=client.post(f'/api/brands/{brand_id}/preview',json={'template_id':'tech'})
    assert preview.status_code==200
    assert preview.json()['logo']['src'].startswith('data:image/png;base64,')


def test_parallel_slide_builds_keep_brand_font_context_isolated():
    from concurrent.futures import ThreadPoolExecutor
    slide=Slide(title='Brand title',body='Body')
    metadata=branded_template()['metadata']
    def render(index):
        scene=build_scene(slide,metadata if index%2 else {'ratio':16/9},'a',0)
        return next(node['effective_font_family'] for node in scene['objects'] if node['id']=='title')
    with ThreadPoolExecutor(max_workers=4) as executor: actual=list(executor.map(render,range(12)))
    assert actual==[FAMILY if i%2 else 'Manrope' for i in range(12)]
