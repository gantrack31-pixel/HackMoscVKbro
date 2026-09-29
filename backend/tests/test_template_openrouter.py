import asyncio
from io import BytesIO
from typing import get_args
from zipfile import ZipFile
import httpx
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pydantic import ValidationError
from app import database as db
from app.config import Settings, settings
from app.main import get_template
from app.models import DeckContent, Slide, SmartArt, DiagramType
from app.services import audit, llm
from app.services.layout import build_scene
from app.services.export import export_pptx, export_pdf, export_html
from app.services.palette import allowed_color, blend, roles
from app.services.sources import slide_binding, validate_bindings
from app.services.templates import analyze_template
from app.services.template_context import model_template
from test_assistant import client


def template_file(path, color):
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color)
    shape.text='Corporate template'
    shape.text_frame.paragraphs[0].font.name='Corporate Font'
    prs.save(path)
    return prs


def test_template_identity_refresh_and_full_primitive_manifest(client,tmp_path):
    path=tmp_path/'active.pptx';template_file(path,'E47123')
    metadata=analyze_template(path)
    assert '#E47123' in metadata['palette']
    assert 'Corporate Font' in metadata['manifest']['fonts']
    assert any(e['part']=='ppt/slides/slide1.xml' and e['properties'] for e in metadata['manifest']['elements'])
    db.template_save('active','Active',str(path),{**metadata,'composition':'split'})
    old=metadata['file_fingerprint']
    template_file(path,'284A6C')
    active=get_template('active')
    assert active['metadata']['file_fingerprint']!=old
    assert '#284A6C' in active['metadata']['palette'] and '#E47123' not in active['metadata']['palette']
    assert 'composition' not in active['metadata']  # no stale preset metadata
    context=model_template(active)
    assert context['fingerprint']==active['metadata']['file_fingerprint']
    assert context['id']=='active' and '#284A6C' in context['palette'].values()


def test_palette_tolerance_and_grouping(monkeypatch):
    meta={'ratio':16/9,'accent':'#204060','colors':{'accent1':'#204060'}}
    assert allowed_color('#214161',meta)
    assert allowed_color(blend('#204060','#FFFFFF',.5),meta)
    assert not allowed_color('#FA0090',meta)
    scene=build_scene(Slide(title='Title'),meta,'a',0)
    scene['render_objects'] += [{'id':f'off-{i}','type':'rect','fill':'#FA0090'} for i in range(30)]
    monkeypatch.setattr(audit,'build_scene',lambda *args:scene)
    issues=audit.audit_deck(DeckContent(title='Deck',slides=[Slide(title='Title')]),{'metadata':meta},'a','')
    colors=[i for i in issues if i.code.startswith('color_off_palette')]
    assert len(colors)==1 and '30' in colors[0].detail


@pytest.mark.parametrize('kind',get_args(DiagramType))
@pytest.mark.parametrize('ratio',[16/9,4/3,3/4])
def test_diagrams_have_palette_geometry_and_editable_exports(kind,ratio):
    meta={'ratio':ratio,'accent':'#652B91','background':'#151025',
          'colors':{'accent1':'#652B91','lt1':'#FFFFFF','dk1':'#151025'}}
    count=4 if kind=='matrix' else 6
    slide=Slide(title='Этапы',kind='diagram',diagram_type=kind,bullets=[f'Этап {i+1}' for i in range(count)])
    deck=DeckContent(title='Схемы',slides=[slide])
    for variant in 'abc':
        scene=build_scene(slide,meta,variant,0)
        assert len([o for o in scene['objects'] if o['id'].startswith('node-label')])==count
        for obj in scene['render_objects']:
            assert min(obj['x'],obj['y'],obj['w'],obj['h'])>=0, (kind,ratio,obj)
            assert obj['x']+obj['w']<=scene['width']+.1 and obj['y']+obj['h']<=scene['height']+.1
            assert all(allowed_color(obj[key],meta) for key in ('color','fill','stroke') if key in obj)
        errors=[i for i in audit.audit_deck(deck,{'metadata':meta},variant,'') if i.severity=='error']
        assert not errors, (kind,ratio,variant,errors)
    template={'metadata':meta,'path':None}
    output=export_pptx(deck,template,'a')
    prs=Presentation(BytesIO(output))
    assert sum(s.has_text_frame and s.text.startswith('Этап ') for s in prs.slides[0].shapes)==count
    assert export_pdf(deck,template,'a').startswith(b'%PDF')
    assert '<svg' in export_html(deck,template,'a').decode()
    if kind=='honeycomb':
        with ZipFile(BytesIO(output)) as z:assert b'hexagon' in z.read('ppt/slides/slide1.xml')


def test_structured_smartart_validates_bounds_edges_and_overlap():
    nodes=[{'bullet':0,'x':0,'y':0,'w':.4,'h':.5},{'bullet':1,'x':.6,'y':0,'w':.4,'h':.5}]
    art=SmartArt(nodes=nodes,edges=[{'source':0,'target':1}])
    assert len(art.nodes)==2
    for invalid in (
        {'nodes':[dict(nodes[0],x=.9)]},
        {'nodes':[nodes[0],dict(nodes[1],x=.2)]},
        {'nodes':nodes,'edges':[{'source':0,'target':5}]},
    ):
        with pytest.raises(ValidationError):SmartArt.model_validate(invalid)


def test_fuzzy_bindings_are_suggestions_and_numbers_must_match():
    source='В пилоте участвовали 30 студентов. Команда подготовила учебные материалы для школы.'
    slide=Slide(title='Пилот',body='В пилоте участвовало 30 студентов и команда готовила учебные материалы.')
    result=slide_binding(slide,source)
    assert result['status']=='related' and not result['semantic_verified']
    assert validate_bindings(result['candidates'],source)
    slide.body=slide.body.replace('30','90')
    assert not slide_binding(slide,source)['candidates']
    issues=audit.audit_deck(DeckContent(title='D',slides=[Slide(title='Введение',kind='title')]),{'metadata':{'ratio':16/9}},'a','')
    assert all(i.severity=='info' for i in issues if i.category=='source')


def test_openrouter_configuration_never_uses_legacy_credentials(monkeypatch):
    monkeypatch.delenv('LLM_BASE_URL',raising=False)
    monkeypatch.delenv('OPENROUTER_BASE_URL',raising=False)
    monkeypatch.delenv('OPENROUTER_MODEL',raising=False)
    monkeypatch.setenv('OPENROUTER_API_KEY','router-test-key')
    monkeypatch.setenv('LLM_API_KEY','legacy-key')
    monkeypatch.setenv('LLM_MODEL','gpt://legacy/model')
    config=Settings()
    assert config.base_url=='https://openrouter.ai/api/v1'
    assert config.model=='qwen/qwen3.8-27b' and config.api_key=='router-test-key'
    assert config.api_key not in repr(config)


def test_explicit_custom_provider_does_not_inherit_openrouter_key(monkeypatch):
    monkeypatch.setenv('LLM_BASE_URL','http://127.0.0.1:8001/v1')
    monkeypatch.setenv('LLM_MODEL','Qwen/Qwen3.8-27B')
    monkeypatch.delenv('LLM_API_KEY',raising=False)
    monkeypatch.setenv('OPENROUTER_API_KEY','router-key')
    config=Settings()
    assert config.base_url=='http://127.0.0.1:8001/v1' and config.model=='Qwen/Qwen3.8-27B'
    assert config.api_key==''


@pytest.mark.parametrize('status',[401,429,400])
def test_provider_errors_are_safe(status,monkeypatch):
    monkeypatch.setattr(settings,'api_key','secret-router-key')
    def respond(request):
        return httpx.Response(status,json={'error':{'message':'Rejected secret-router-key'}})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            with pytest.raises(llm.LLMError) as exc: await llm.complete_json('JSON',{},client)
            assert 'secret-router-key' not in str(exc.value)
    asyncio.run(run())
