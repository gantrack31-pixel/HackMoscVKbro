"""Browser cache and transfer regressions without touching application data."""
import gzip

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.main import StaticFiles
from test_security import client


@pytest.fixture
def static_client(tmp_path):
    assets = tmp_path / 'assets'
    assets.mkdir()
    payload = b'const example = "Deckly";\n' * 100
    (assets / 'entry-Abcd1234.js').write_bytes(payload)
    (assets / 'entry-Abcd1234.js.gz').write_bytes(gzip.compress(payload))
    (assets / 'plain.js').write_bytes(payload)
    (assets / 'Manrope-Abcd1234.woff2').write_bytes(b'font bytes')
    (tmp_path / 'index.html').write_text('<html>Deckly</html>')
    app = FastAPI()
    app.mount('/', StaticFiles(directory=tmp_path, html=True))
    with TestClient(app) as browser:
        yield browser


def test_hashed_assets_are_immutable_and_compressed(static_client):
    response = static_client.get('/assets/entry-Abcd1234.js', headers={'Accept-Encoding': 'gzip'})
    assert response.headers['cache-control'] == 'public, max-age=31536000, immutable'
    assert response.headers['content-encoding'] == 'gzip'
    assert 'accept-encoding' in response.headers['vary'].lower()
    assert response.headers['content-type'].startswith(('text/javascript', 'application/javascript'))
    assert response.content == b'const example = "Deckly";\n' * 100


def test_revalidation_preserves_variant_headers(static_client):
    first = static_client.get('/assets/entry-Abcd1234.js', headers={'Accept-Encoding': 'gzip'})
    response = static_client.get('/assets/entry-Abcd1234.js', headers={
        'Accept-Encoding': 'gzip', 'If-None-Match': first.headers['etag'],
    })
    assert response.status_code == 304
    assert response.content == b''
    assert response.headers['cache-control'].endswith('immutable')
    assert 'accept-encoding' in response.headers['vary'].lower()


def test_html_is_never_immutable_and_plain_assets_revalidate(static_client):
    assert static_client.get('/').headers['cache-control'] == 'no-store'
    assert static_client.get('/assets/plain.js').headers['cache-control'] == 'no-cache'


def test_encoding_opt_out_and_ranges_use_original_representation(static_client):
    response = static_client.get('/assets/entry-Abcd1234.js', headers={
        'Accept-Encoding': 'br;q=0, gzip;q=0, identity',
    })
    assert 'content-encoding' not in response.headers
    ranged = static_client.get('/assets/entry-Abcd1234.js', headers={
        'Accept-Encoding': 'gzip', 'Range': 'bytes=0-4',
    })
    assert ranged.status_code == 206
    assert ranged.content == b'const'
    assert 'content-encoding' not in ranged.headers


def test_precompressed_sidecars_are_not_direct_downloads(static_client):
    assert static_client.get('/assets/entry-Abcd1234.js.gz').status_code == 404


def test_thumbnail_revalidates_after_access_check(client, monkeypatch, tmp_path):
    preview = tmp_path / 'thumbnail.png'
    preview.write_bytes(b'cached image')
    from app import database as db
    from app.config import BASE
    db.template_save('tech', 'Tech', str(BASE / 'data/templates/tech.pptx'), {}, None)
    monkeypatch.setattr('app.main.thumbnail', lambda _: preview)
    first = client.get('/api/templates/tech/thumbnail')
    assert first.status_code == 200
    assert first.headers['cache-control'] == 'private, no-cache'
    response = client.get('/api/templates/tech/thumbnail', headers={'If-None-Match': first.headers['etag']})
    assert response.status_code == 304
    assert response.content == b''
    assert response.headers['cache-control'] == 'private, no-cache'
    assert client.get('/api/templates/missing/thumbnail').headers['cache-control'] == 'no-store'


def test_encoding_variants_have_distinct_validators(static_client):
    first = static_client.get('/assets/entry-Abcd1234.js', headers={'Accept-Encoding': 'gzip'})
    identity = static_client.get('/assets/entry-Abcd1234.js', headers={
        'Accept-Encoding': 'identity', 'If-None-Match': first.headers['etag'],
    })
    assert identity.status_code == 200
    assert identity.headers['etag'] != first.headers['etag']


def test_cached_private_thumbnail_is_not_returned_after_logout(client, monkeypatch, tmp_path):
    from io import BytesIO
    from pptx import Presentation
    from app.config import settings
    monkeypatch.setattr(settings, 'email_verification_required', False)
    auth = client.post('/api/auth/register', json={
        'email': 'cache@example.test', 'password': 'Thumbnail-Test-2037!',
        'first_name': 'Cache', 'last_name': 'Test',
    }).json()
    client.headers['X-CSRF-Token'] = auth['csrf']
    source = BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    uploaded = client.post('/api/templates', files={
        'file': ('private.pptx', source.getvalue(), 'application/vnd.openxmlformats-officedocument.presentationml.presentation'),
    })
    assert uploaded.status_code == 201
    preview = tmp_path / 'private.png'
    preview.write_bytes(b'private thumbnail')
    monkeypatch.setattr('app.main.thumbnail', lambda _: preview)
    url = f"/api/templates/{uploaded.json()['id']}/thumbnail"
    cached = client.get(url)
    assert cached.status_code == 200

    client.post('/api/auth/logout')
    response = client.get(url, headers={'If-None-Match': cached.headers['etag']})

    assert response.status_code == 404
    assert response.headers['cache-control'] == 'no-store'
    assert b'private thumbnail' not in response.content


def test_catalog_json_compression_preserves_data_and_private_cache_policy(client):
    original = client.get('/api/public/templates', headers={'Accept-Encoding': 'identity'})
    compressed = client.get('/api/public/templates', headers={'Accept-Encoding': 'gzip'})
    assert compressed.headers.get('content-encoding') == 'gzip'
    assert compressed.headers['cache-control'] == 'no-store'
    assert compressed.json() == original.json()
    assert int(compressed.headers['content-length']) < len(original.content)
    assert 'content-encoding' not in client.get('/api/public/templates', headers={
        'Accept-Encoding': 'gzip;q=0, identity',
    }).headers


def test_hashed_fonts_have_the_font_mime_type(static_client):
    response = static_client.get('/assets/Manrope-Abcd1234.woff2')
    assert response.headers['content-type'] == 'font/woff2'
    assert response.headers['cache-control'].endswith('immutable')


def test_old_sidecars_are_ignored_when_the_source_changes(static_client):
    import os
    from pathlib import Path
    source = Path(static_client.app.routes[-1].app.directory) / 'assets/entry-Abcd1234.js'
    source.write_bytes(b'new source content' * 100)
    os.utime(source, (source.stat().st_atime, source.stat().st_mtime + 5))
    response = static_client.get('/assets/entry-Abcd1234.js', headers={'Accept-Encoding': 'gzip'})
    assert 'content-encoding' not in response.headers
    assert response.content == b'new source content' * 100
