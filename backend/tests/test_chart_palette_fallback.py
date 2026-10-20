"""An unreadable chart accent falls back to another template accent before ink."""
from io import BytesIO

import pytest
from pptx import Presentation

from app.models import DeckContent, Slide
from app.services.export import export_pptx
from app.services.layout import build_scene
from app.services.palette import contrast


@pytest.mark.parametrize('background,accent,alternate', [
    ('#FFFFFF', '#FFFF99', '#2255AA'),
    ('#111111', '#222222', '#77CCEE'),
    ('#FFFFFF', '#FFFF99', '#475569'),
])
@pytest.mark.parametrize('variant', 'abc')
def test_chart_uses_readable_template_accent_in_scene_and_powerpoint(background, accent, alternate, variant):
    metadata = {'background': background, 'accent': accent, 'ratio': 16 / 9,
                'colors': {'dk1': '#000000', 'lt1': '#FFFFFF', 'accent1': accent, 'accent2': alternate}}
    slide = Slide(title='Показатели', kind='chart', chart={'labels': ['A', 'B'], 'values': [10, 20]})

    scene = build_scene(slide, metadata, variant, 0)
    chart = next(obj for obj in scene['objects'] if obj['id'] == 'chart')
    assert chart['fill'] == alternate
    assert contrast(chart['fill'], background) >= 3
    presentation = Presentation(BytesIO(export_pptx(DeckContent(title='Тест', slides=[slide]), {'metadata': metadata}, variant)))
    native = next(shape.chart for shape in presentation.slides[0].shapes if shape.has_chart)
    assert str(native.series[0].format.fill.fore_color.rgb) == alternate[1:]


def test_chart_preserves_a_readable_requested_accent():
    metadata = {'background': '#FFFFFF', 'accent': '#2255AA', 'colors': {'accent2': '#008877'}}
    slide = Slide(title='Показатели', kind='chart', chart={'labels': ['A'], 'values': [1]})
    scene = build_scene(slide, metadata, 'a', 0)
    assert next(obj['fill'] for obj in scene['objects'] if obj['id'] == 'chart') == '#2255AA'


def test_line_preview_keeps_the_same_readable_accent_as_native_chart():
    metadata = {'background': '#FFFFFF', 'accent': '#0077FF', 'colors': {'dk1': '#000000', 'accent1': '#0077FF'}}
    slide = Slide(title='Динамика', kind='chart', chart={'labels': ['A', 'B'], 'values': [1, 2], 'chart_type': 'line'})
    scene = build_scene(slide, metadata, 'a', 0)
    assert {node['stroke'] for node in scene['render_objects'] if node['type'] == 'line' and node['id'].startswith('chart-')} == {'#0077FF'}
