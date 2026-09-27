"""Smoke-test a real model on synthetic material; never print keys or user data."""
import argparse
import asyncio
from io import BytesIO
import json
from pathlib import Path
import sys
from time import monotonic

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.config import settings, BASE
from app.models import OutlineRequest
from app.services.llm import make_outline, create_design_variants
from app.services.templates import analyze_template
from app.services.audit import audit_deck
from app.services.export import export_pptx, export_pdf, export_html
from pptx import Presentation

async def validate():
    started = monotonic()
    # Unseen template: the same parser used by the upload endpoint.
    import tempfile
    with tempfile.TemporaryDirectory(prefix='deckly-live-check-') as directory:
        path = Path(directory) / 'unseen-template.pptx'
        original = Presentation()
        original.slide_width = 9144000
        original.slide_height = 6858000
        original.slides.add_slide(original.slide_layouts[1])
        original.save(path)
        template = {'id': 'acceptance-unseen', 'name': 'Новый шаблон 4:3',
                    'path': str(path), 'metadata': analyze_template(path)}
        material = ('Синтетический проект «Навигатор». Цель — сократить время поиска документов. '
                    'Пилот охватил 30 сотрудников. Среднее время поиска до пилота 12 минут, после 8 минут. '
                    'План: собрать материалы, подключить поиск, провести обучение, оценить результат. '
                    'Команда: аналитик, разработчик, координатор. Риски: неполные материалы и устаревшие инструкции. '
                    'Предложение: расширить пилот после проверки качества. Создай 10 слайдов, '
                    'включи диаграмму по указанным числам, таблицу ролей, процесс и пиктограммы; без изображений.')
        request = OutlineRequest(template_id=template['id'], prompt=material, count=10, purpose='project')
        async with asyncio.timeout(120):
            outline = await make_outline(request, template)
        outline_seconds = monotonic() - started
        async with asyncio.timeout(180):
            content = await create_design_variants(outline, template)
            results = {}
            for variant in 'abc':
                pptx = await asyncio.to_thread(export_pptx, content, template, variant)
                deck = Presentation(BytesIO(pptx))
                assert len(deck.slides) == 10
                assert all(any(s.has_text_frame for s in slide.shapes) for slide in deck.slides)
                pdf = await asyncio.to_thread(export_pdf, content, template, variant)
                html = await asyncio.to_thread(export_html, content, template, variant)
                issues = audit_deck(content, template, variant, material)
                results[variant] = {'pptx_bytes': len(pptx), 'pdf_bytes': len(pdf), 'html_bytes': len(html),
                                    'errors': sum(i.severity == 'error' for i in issues),
                                    'issue_codes': sorted(set(i.code for i in issues))}
        elapsed = monotonic() - started
        assert elapsed < 300
        return {'mode': settings.mode, 'slides': len(content.slides),
                'kinds': sorted(set(s.kind for s in content.slides)),
                'outline_seconds': round(outline_seconds, 2), 'total_seconds': round(elapsed, 2),
                'variants': results, 'image_service_configured': bool(settings.image_base_url)}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live', action='store_true', help='Send synthetic material to the configured model')
    args = parser.parse_args()
    if not args.live:
        parser.error('Use --live to explicitly run the external model check.')
    if settings.mode != 'live':
        parser.error('LLM_MODE=live is required; demo is not a model verification.')
    print(json.dumps(asyncio.run(validate()), ensure_ascii=False, indent=2))
